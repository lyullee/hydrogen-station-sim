"""Prospectively specified PRESLHY E3.1 blowdown validation utilities."""

from __future__ import annotations

from dataclasses import dataclass
from io import BytesIO
import math
from pathlib import Path, PurePosixPath
import re
from typing import Iterable
import zipfile

import numpy as np
from openpyxl import load_workbook
from scipy.integrate import solve_ivp

from .full_station import CascadeBank, CascadeBankParameters, CascadeBankState
from .risk.live import DynamicLeakModel, LeakScenario, LeakSourceState
from .tabulated import PropsSI


SENSOR_FULL_SCALE_STEP_BAR = 0.00125 * 250.0


@dataclass(frozen=True)
class PreslhyTrace:
    case_id: str
    source_package: str
    source_member: str
    nozzle_diameter_mm: float
    time_s: np.ndarray
    pressure_bar_abs: np.ndarray
    initial_temperature_k: float
    ambient_pressure_pa: float
    pressure_unit_interpretation: str
    temperature_substituted: bool
    ambient_pressure_substituted: bool

    @property
    def initial_pressure_pa(self) -> float:
        before = np.flatnonzero(self.time_s < 0.0)
        index = int(before[-1]) if before.size else 0
        return float(self.pressure_bar_abs[index] * 1.0e5)


@dataclass(frozen=True)
class PreslhyCaseResult:
    case_id: str
    nozzle_diameter_mm: float
    initial_pressure_bar_abs: float
    samples: int
    pressure_rmse_bar: float
    pressure_mae_bar: float
    pressure_nrmse_percent_initial_absolute_pressure: float
    experimental_time_to_50_percent_gauge_s: float
    predicted_time_to_50_percent_gauge_s: float
    time_to_50_percent_gauge_relative_error_percent: float
    pressure_screen_pass: bool
    half_time_screen_pass: bool
    joint_primary_screen_pass: bool
    peak_mass_flow_kg_s: float
    cumulative_released_mass_kg: float


def _normalise(value: object) -> str:
    return re.sub(r"[^a-z0-9]+", "", str(value or "").lower())


def _header_layout(sheet, max_rows: int = 15, max_columns: int = 80):
    rows = list(
        sheet.iter_rows(
            min_row=1,
            max_row=max_rows,
            max_col=min(sheet.max_column, max_columns),
            values_only=True,
        )
    )
    width = max((len(row) for row in rows), default=0)
    profiles = [
        [
            _normalise(row[column])
            for row in rows
            if column < len(row) and row[column] is not None
        ]
        for column in range(width)
    ]
    return rows, profiles


def _data_time_column(rows) -> tuple[int, int]:
    for row_index, row in enumerate(rows, start=1):
        for column, value in enumerate(row):
            normal = _normalise(value)
            if normal == "times" or normal == "synchronizedtimes":
                return row_index, column
    raise ValueError("synchronized Time [s] data column not found")


def _profile_column(profiles, *exact_tokens: str) -> int:
    for index, profile in enumerate(profiles):
        if any(
            value == token or value.startswith(token)
            for value in profile
            for token in exact_tokens
        ):
            return index
    raise ValueError(f"column profile containing {exact_tokens} not found")


def _profile_label(rows, column: int) -> str:
    return " | ".join(
        str(row[column])
        for row in rows
        if column < len(row) and row[column] not in (None, "")
    )


def _numeric_rows(sheet, header_row: int, columns: tuple[int, ...]):
    blank_run = 0
    for row in sheet.iter_rows(
        min_row=header_row + 1,
        min_col=1,
        max_col=max(columns) + 1,
        values_only=True,
    ):
        values = tuple(row[index] if index < len(row) else None for index in columns)
        if all(value is None for value in values):
            blank_run += 1
            if blank_run >= 50:
                return
            continue
        blank_run = 0
        try:
            converted = tuple(float(value) for value in values)
        except (TypeError, ValueError):
            continue
        if all(math.isfinite(value) for value in converted):
            yield converted


def _pressure_unit(header: object, pressure_bar: np.ndarray) -> tuple[np.ndarray, str]:
    label = str(header or "").lower()
    if "abs" in label or "bara" in _normalise(label):
        return pressure_bar, "header_absolute"
    if "gauge" in label or "barg" in _normalise(label):
        return pressure_bar + 1.01325, "header_gauge_plus_standard_ambient"
    tail = pressure_bar[max(0, len(pressure_bar) - max(20, len(pressure_bar) // 20)) :]
    terminal = float(np.median(tail))
    if -0.5 <= terminal <= 0.5:
        return pressure_bar + 1.01325, "terminal_near_zero_gauge_plus_standard_ambient"
    if 0.5 < terminal <= 1.5:
        return pressure_bar, "terminal_near_ambient_absolute"
    raise ValueError(f"ambiguous Pves pressure basis; terminal median={terminal:.3f} bar")


def _initial_temperature(workbook, case_prefix: str) -> tuple[float, bool]:
    candidates = [
        sheet for sheet in workbook.worksheets
        if "temp" in sheet.title.lower() and case_prefix in sheet.title
    ] or [sheet for sheet in workbook.worksheets if "temp" in sheet.title.lower()]
    for sheet in candidates:
        try:
            rows, profiles = _header_layout(sheet)
            header_row, time_column = _data_time_column(rows)
        except ValueError:
            continue
        temperature_columns = [
            index for index, profile in enumerate(profiles)
            if any(
                re.fullmatch(r"t[123](?:o)?(?:degc|c|kelvin)?", value)
                for value in profile
            )
        ]
        samples = []
        for row in sheet.iter_rows(
            min_row=header_row + 1,
            min_col=1,
            max_col=max([time_column, *temperature_columns], default=time_column) + 1,
            values_only=True,
        ):
            try:
                time_s = float(row[time_column])
            except (TypeError, ValueError, IndexError):
                continue
            if not -0.5 <= time_s < 0.0:
                continue
            for column in temperature_columns:
                try:
                    value = float(row[column])
                except (TypeError, ValueError, IndexError):
                    continue
                if math.isfinite(value):
                    samples.append(value)
        if samples:
            median = float(np.median(samples))
            return (median + 273.15 if median < 200.0 else median), False
    return 293.15, True


def _ambient_pressure(workbook, case_prefix: str) -> tuple[float, bool]:
    candidates = [
        sheet for sheet in workbook.worksheets
        if case_prefix in sheet.title and (
            "amb" in sheet.title.lower() or "ch2" in sheet.title.lower()
        )
    ]
    for sheet in candidates:
        try:
            rows, profiles = _header_layout(sheet)
            header_row, time_column = _data_time_column(rows)
        except ValueError:
            continue
        pressure_columns = [
            index for index, profile in enumerate(profiles)
            if any(
                (
                    ("ambient" in header or "pamb" in header or "airpressure" in header)
                    and ("pressure" in header or "pamb" in header)
                )
                for header in profile
            )
        ]
        for pressure_column in pressure_columns:
            samples = [
                pressure
                for time_s, pressure in _numeric_rows(
                    sheet, header_row, (time_column, pressure_column)
                )
                if -5.0 <= time_s < 0.0
            ]
            if not samples:
                continue
            value = float(np.median(samples))
            header = _profile_label(rows, pressure_column).lower()
            normal = _normalise(header)
            if "hpa" in normal or "mbar" in normal:
                pressure_pa = value * 100.0
            elif "kpa" in normal:
                pressure_pa = value * 1000.0
            elif "bar" in normal:
                pressure_pa = value * 1.0e5
            elif re.search(r"(?:^|[^a-z])pa(?:[^a-z]|$)", header):
                pressure_pa = value
            elif 800.0 <= value <= 1200.0:
                pressure_pa = value * 100.0
            elif 80.0 <= value <= 120.0:
                pressure_pa = value * 1000.0
            elif 0.8 <= value <= 1.2:
                pressure_pa = value * 1.0e5
            else:
                continue
            if 80_000.0 <= pressure_pa <= 120_000.0:
                return pressure_pa, False
    return 101_325.0, True


def read_preslhy_workbook(
    payload: bytes,
    *,
    source_package: str,
    source_member: str,
    nozzle_diameter_mm: float,
) -> PreslhyTrace:
    workbook = load_workbook(BytesIO(payload), read_only=True, data_only=True)
    try:
        pressure_sheets = [sheet for sheet in workbook.worksheets if "press" in sheet.title.lower()]
        if not pressure_sheets:
            raise ValueError("pressure worksheet not found")
        sheet = pressure_sheets[0]
        header_rows, profiles = _header_layout(sheet)
        header_row, time_column = _data_time_column(header_rows)
        pressure_column = _profile_column(profiles, "pves", "pvessel")
        samples = list(
            _numeric_rows(sheet, header_row, (time_column, pressure_column))
        )
        if len(samples) < 20:
            raise ValueError("fewer than 20 finite pressure samples")
        values = np.asarray(samples, dtype=float)
        order = np.argsort(values[:, 0], kind="stable")
        time_s = values[order, 0]
        pressure, unit = _pressure_unit(
            _profile_label(header_rows, pressure_column), values[order, 1]
        )
        unique = np.concatenate(([True], np.diff(time_s) > 0.0))
        time_s, pressure = time_s[unique], pressure[unique]
        if len(time_s) < 20:
            raise ValueError("fewer than 20 unique pressure timestamps")
        case_prefix = PurePosixPath(source_member).stem.split("-")[0]
        temperature_k, substituted = _initial_temperature(workbook, case_prefix)
        ambient_pressure_pa, ambient_substituted = _ambient_pressure(
            workbook, case_prefix
        )
        return PreslhyTrace(
            case_id=PurePosixPath(source_member).stem,
            source_package=source_package,
            source_member=source_member,
            nozzle_diameter_mm=float(nozzle_diameter_mm),
            time_s=time_s,
            pressure_bar_abs=pressure,
            initial_temperature_k=temperature_k,
            ambient_pressure_pa=ambient_pressure_pa,
            pressure_unit_interpretation=unit,
            temperature_substituted=substituted,
            ambient_pressure_substituted=ambient_substituted,
        )
    finally:
        workbook.close()


def iter_preslhy_traces(root: Path) -> Iterable[PreslhyTrace]:
    pattern = re.compile(r"PRE3P1A_KIT_D(05|1|2|4)_300K_DATA\.zip$", re.I)
    for package in sorted(root.glob("PRE3P1A_KIT_D*_300K_DATA.zip")):
        match = pattern.fullmatch(package.name)
        if not match:
            continue
        diameter_token = match.group(1)
        diameter_mm = 0.5 if diameter_token == "05" else float(diameter_token)
        with zipfile.ZipFile(package) as bundle:
            for member in sorted(bundle.namelist()):
                if member.lower().endswith(".xlsx") and not PurePosixPath(member).name.startswith("~$"):
                    yield read_preslhy_workbook(
                        bundle.read(member),
                        source_package=package.name,
                        source_member=member,
                        nozzle_diameter_mm=diameter_mm,
                    )


def eligible_pressure_window(trace: PreslhyTrace) -> tuple[np.ndarray, np.ndarray]:
    mask = trace.time_s >= 0.1
    time_s = trace.time_s[mask]
    pressure = trace.pressure_bar_abs[mask]
    if len(time_s) < 20:
        raise ValueError("fewer than 20 post-opening samples")
    increases = np.diff(pressure) > SENSOR_FULL_SCALE_STEP_BAR
    if increases.size and float(np.mean(increases)) > 0.10:
        raise ValueError("pressure trace has excessive significant post-release increases")
    threshold = 1.05 * trace.ambient_pressure_pa / 1.0e5
    reached = np.flatnonzero(pressure <= threshold)
    if reached.size:
        end = int(reached[0]) + 1
        time_s, pressure = time_s[:end], pressure[:end]
    return time_s, pressure


def simulate_preslhy_blowdown(
    trace: PreslhyTrace,
    evaluation_time_s: np.ndarray,
    *,
    discharge_coefficient: float = 0.8,
) -> tuple[np.ndarray, np.ndarray, np.ndarray]:
    initial_pressure = trace.initial_pressure_pa
    bank = CascadeBank(
        CascadeBankParameters(
            name="preslhy-vessel",
            internal_volume_m3=0.002815,
            target_pressure_pa=initial_pressure,
            wall_mass_kg=1.0,
            wall_specific_heat_j_kg_k=1.0,
            gas_wall_ua_w_k=0.0,
            wall_ambient_ua_w_k=0.0,
        )
    )
    initial = bank.initial_state(initial_pressure, trace.initial_temperature_k)
    leak = DynamicLeakModel()
    scenario = LeakScenario(
        release_id=trace.case_id,
        component_id="PRESLHY-DisCha",
        location="direct circular aperture",
        start_time_s=0.0,
        orifice_diameter_m=trace.nozzle_diameter_mm / 1000.0,
        discharge_coefficient=discharge_coefficient,
    )
    ambient = trace.ambient_pressure_pa

    def derivative(_time, vector):
        state = CascadeBankState.from_vector(vector)
        gas = bank.gas_state(state)
        if gas.pressure_pa <= ambient * 1.000001:
            return np.zeros(3)
        flow = leak.mass_flow_kg_s(
            scenario,
            LeakSourceState(gas.pressure_pa, gas.temperature_k, ambient),
        )
        flow = min(flow, state.hydrogen_mass_kg / 1.0e-3)
        return np.asarray([-flow, -flow * gas.specific_enthalpy_j_kg, 0.0])

    end_time = float(evaluation_time_s[-1])
    solution = solve_ivp(
        derivative,
        (0.0, end_time),
        initial.as_vector(),
        method="LSODA",
        rtol=1.0e-7,
        atol=(1.0e-10, 1.0e-2, 1.0e-8),
        max_step=max(0.001, min(0.1, end_time / 500.0)),
    )
    if not solution.success:
        raise RuntimeError(f"PRESLHY integration failed: {solution.message}")
    node_states = [
        bank.gas_state(CascadeBankState.from_vector(solution.y[:, index]))
        for index in range(solution.y.shape[1])
    ]
    pressure_nodes = np.asarray([state.pressure_pa for state in node_states])
    node_flows = [
        leak.mass_flow_kg_s(
            scenario,
            LeakSourceState(state.pressure_pa, state.temperature_k, ambient),
        )
        if state.pressure_pa > ambient * 1.000001
        else 0.0
        for state in node_states
    ]
    predicted = np.interp(evaluation_time_s, solution.t, pressure_nodes)
    mass = np.interp(evaluation_time_s, solution.t, solution.y[0])
    return predicted, mass, np.asarray([max(node_flows, default=0.0)])


def _crossing_time(time_s: np.ndarray, pressure_pa: np.ndarray, target_pa: float) -> float:
    reached = np.flatnonzero(pressure_pa <= target_pa)
    if not reached.size:
        return math.nan
    index = int(reached[0])
    if index == 0:
        return float(time_s[0])
    p0, p1 = pressure_pa[index - 1], pressure_pa[index]
    t0, t1 = time_s[index - 1], time_s[index]
    if p1 == p0:
        return float(t1)
    return float(t0 + (target_pa - p0) * (t1 - t0) / (p1 - p0))


def evaluate_preslhy_trace(trace: PreslhyTrace) -> PreslhyCaseResult:
    time_s, measured_bar = eligible_pressure_window(trace)
    measured_pa = measured_bar * 1.0e5
    predicted_pa, mass, peak_flow = simulate_preslhy_blowdown(trace, time_s)
    errors_bar = (predicted_pa - measured_pa) / 1.0e5
    initial_pressure = trace.initial_pressure_pa
    target = trace.ambient_pressure_pa + 0.5 * (
        initial_pressure - trace.ambient_pressure_pa
    )
    experimental_half = _crossing_time(time_s, measured_pa, target)
    predicted_half = _crossing_time(time_s, predicted_pa, target)
    if not math.isfinite(experimental_half) or experimental_half <= 0.0:
        half_error = math.inf
    elif not math.isfinite(predicted_half):
        half_error = math.inf
    else:
        half_error = abs(predicted_half - experimental_half) / experimental_half * 100.0
    nrmse = float(np.sqrt(np.mean(errors_bar**2)) * 1.0e5 / initial_pressure * 100.0)
    pressure_pass = nrmse <= 10.0
    half_pass = half_error <= 20.0
    return PreslhyCaseResult(
        case_id=trace.case_id,
        nozzle_diameter_mm=trace.nozzle_diameter_mm,
        initial_pressure_bar_abs=initial_pressure / 1.0e5,
        samples=len(time_s),
        pressure_rmse_bar=float(np.sqrt(np.mean(errors_bar**2))),
        pressure_mae_bar=float(np.mean(np.abs(errors_bar))),
        pressure_nrmse_percent_initial_absolute_pressure=nrmse,
        experimental_time_to_50_percent_gauge_s=experimental_half,
        predicted_time_to_50_percent_gauge_s=predicted_half,
        time_to_50_percent_gauge_relative_error_percent=half_error,
        pressure_screen_pass=pressure_pass,
        half_time_screen_pass=half_pass,
        joint_primary_screen_pass=pressure_pass and half_pass,
        peak_mass_flow_kg_s=float(peak_flow[0]),
        cumulative_released_mass_kg=float(
            PropsSI(
                "Dmass", "P", initial_pressure, "T", trace.initial_temperature_k,
                "Hydrogen",
            )
            * 0.002815
            - mass[-1]
        ),
    )
