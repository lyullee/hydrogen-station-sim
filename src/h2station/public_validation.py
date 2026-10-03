"""Readers and metrics for third-party public validation datasets.

No source workbook or archive is bundled with the project.  The companion
download script verifies source files against the committed research manifest.
"""

from __future__ import annotations

from dataclasses import asdict, dataclass, replace
import csv
from io import BytesIO
from io import StringIO
import math
import re
from pathlib import Path
from statistics import median
from typing import Iterable, Mapping, Sequence
import zipfile

import numpy as np

from .tabulated import PropsSI


@dataclass(frozen=True)
class H2ProtocolMetadata:
    fill_number: str
    lab_test_number: int | None
    test_code: str
    nominal_pressure_mpa: float | None
    tank_capacity_kg: float | None
    chamber_temperature_c: float | None
    scheduled_aprr_mpa_min: float | None
    notes: str


@dataclass(frozen=True)
class H2ProtocolTrace:
    case_id: str
    metadata: H2ProtocolMetadata
    source_archive: str
    source_member: str
    pressure_source: str
    soc_source: str
    time_s: np.ndarray
    pressure_mpa: np.ndarray
    soc_percent: np.ndarray
    mass_flow_g_s: np.ndarray
    tank_temperature_mean_c: np.ndarray
    tank_temperature_max_c: np.ndarray
    inlet_gas_temperature_c: np.ndarray
    chamber_temperature_c: np.ndarray
    source_pressure_1_mpa: np.ndarray | None = None
    source_pressure_3_mpa: np.ndarray | None = None
    protocol_pressure_time_s: np.ndarray | None = None
    protocol_pressure_mpa: np.ndarray | None = None

    def summary(self) -> dict:
        integrated_mass_kg = float(np.trapezoid(self.mass_flow_g_s / 1000.0, self.time_s))
        nominal_soc_mass_change_kg = (
            float(self.metadata.tank_capacity_kg)
            * float(self.soc_percent[-1] - self.soc_percent[0]) / 100.0
            if self.metadata.tank_capacity_kg is not None else None
        )
        summary = {
            "case_id": self.case_id,
            **asdict(self.metadata),
            "source_archive": self.source_archive,
            "source_member": self.source_member,
            "pressure_source": self.pressure_source,
            "soc_source": self.soc_source,
            "sample_count": int(len(self.time_s)),
            "sample_period_median_s": float(median(np.diff(self.time_s))),
            "active_duration_s": float(self.time_s[-1] - self.time_s[0]),
            "initial_pressure_mpa": float(self.pressure_mpa[0]),
            "final_pressure_mpa": float(self.pressure_mpa[-1]),
            "initial_soc_percent": float(self.soc_percent[0]),
            "final_soc_percent": float(self.soc_percent[-1]),
            "peak_tank_temperature_c": float(np.max(self.tank_temperature_max_c)),
            "median_inlet_gas_temperature_c": float(np.median(self.inlet_gas_temperature_c)),
            "peak_mass_flow_g_s": float(np.max(self.mass_flow_g_s)),
            "mean_mass_flow_g_s": float(np.mean(self.mass_flow_g_s)),
            "integrated_mass_flow_kg": integrated_mass_kg,
            "nominal_soc_mass_change_kg": nominal_soc_mass_change_kg,
            "mass_closure_ratio": (
                integrated_mass_kg / nominal_soc_mass_change_kg
                if nominal_soc_mass_change_kg not in (None, 0.0) else None
            ),
        }
        if (
            self.protocol_pressure_time_s is not None
            and self.protocol_pressure_mpa is not None
            and len(self.protocol_pressure_time_s) >= 2
        ):
            duration_s = float(
                self.protocol_pressure_time_s[-1] - self.protocol_pressure_time_s[0]
            )
            summary.update({
                "protocol_initial_pressure_mpa": float(self.protocol_pressure_mpa[0]),
                "protocol_target_pressure_mpa": float(self.protocol_pressure_mpa[-1]),
                "protocol_duration_s": duration_s,
                "protocol_effective_aprr_mpa_min": (
                    60.0
                    * float(self.protocol_pressure_mpa[-1] - self.protocol_pressure_mpa[0])
                    / duration_s
                    if duration_s > 0.0 else None
                ),
            })
        if self.source_pressure_1_mpa is not None:
            summary["source_pressure_1_initial_mpa"] = float(
                self.source_pressure_1_mpa[0]
            )
            summary["source_pressure_1_final_mpa"] = float(
                self.source_pressure_1_mpa[-1]
            )
        if self.source_pressure_3_mpa is not None:
            summary["source_pressure_3_initial_mpa"] = float(
                self.source_pressure_3_mpa[0]
            )
            summary["source_pressure_3_final_mpa"] = float(
                self.source_pressure_3_mpa[-1]
            )
        return summary


@dataclass(frozen=True)
class TraceAgreement:
    sample_count: int
    pressure_rmse_mpa: float
    pressure_mae_mpa: float
    pressure_final_error_mpa: float
    temperature_rmse_c: float
    temperature_mae_c: float
    temperature_peak_error_c: float
    soc_rmse_percentage_points: float
    soc_final_error_percentage_points: float

    def to_dict(self) -> dict:
        return asdict(self)


@dataclass(frozen=True)
class HIADCase:
    event_id: str
    quality: str
    title: str
    description: str
    initiating_system: str
    physical_effect: str
    consequence_nature: str
    root_cause_summary: str
    application: str
    sub_application: str
    supply_chain_stage: str
    operational_condition: str
    emergency_action: str
    lesson_learnt: str
    corrective_measures: str
    references: tuple[str, ...]

    def to_dict(self) -> dict:
        result = asdict(self)
        result["references"] = list(self.references)
        return result


@dataclass(frozen=True)
class DispersionExperiment:
    case_id: str
    article_test_id: str
    source_archive: str
    mean_mass_flow_g_s: float
    mean_filling_pressure_bar: float | None
    channel_temperature_c: float | None
    baseline_duration_s: float
    filling_duration_s: float
    sensor_ids: tuple[str, ...]
    sensor_time_s: np.ndarray
    concentrations_percent: np.ndarray

    def summary(self) -> dict:
        steady_start = self.baseline_duration_s + 2.0 * self.filling_duration_s / 3.0
        steady_end = self.baseline_duration_s + self.filling_duration_s
        steady = self.concentrations_percent[
            (self.sensor_time_s >= steady_start) & (self.sensor_time_s <= steady_end)
        ]
        if steady.size == 0:
            raise ValueError(f"{self.case_id} has no sensor samples in its steady interval")
        nonnegative = np.maximum(steady, 0.0)
        return {
            "case_id": self.case_id,
            "article_test_id": self.article_test_id,
            "source_archive": self.source_archive,
            "mean_mass_flow_g_s": self.mean_mass_flow_g_s,
            "mean_filling_pressure_bar": self.mean_filling_pressure_bar,
            "channel_temperature_c": self.channel_temperature_c,
            "baseline_duration_s": self.baseline_duration_s,
            "filling_duration_s": self.filling_duration_s,
            "sensor_count": len(self.sensor_ids),
            "steady_interval_start_s": steady_start,
            "steady_interval_end_s": steady_end,
            "steady_mean_concentration_percent": float(np.mean(nonnegative)),
            "steady_peak_concentration_percent": float(np.max(nonnegative)),
            "steady_lfl_exceedance_fraction": float(np.mean(nonnegative >= 4.0)),
        }


def _text(value: object) -> str:
    return "" if value is None else str(value).strip()


def _number(value: object) -> float | None:
    if isinstance(value, (int, float)) and math.isfinite(float(value)):
        return float(value)
    match = re.search(r"[-+]?\d+(?:\.\d+)?", _text(value))
    return float(match.group()) if match else None


def _capacity_kg(value: object) -> float | None:
    text = _text(value).lower().replace("×", "x")
    product = re.search(r"(\d+(?:\.\d+)?)\s*x\s*(\d+(?:\.\d+)?)\s*kg", text)
    if product:
        return float(product.group(1)) * float(product.group(2))
    return _number(value)


def _test_key(value: object) -> str:
    return re.sub(r"[^a-z0-9]", "", _text(value).lower())


def _require_openpyxl():
    try:
        from openpyxl import load_workbook
    except ImportError as exc:  # pragma: no cover - exercised by installation state
        raise RuntimeError(
            "Public workbook preparation requires the 'research' extra: "
            "pip install -e .[research]"
        ) from exc
    return load_workbook


def read_h2protocol_overview(archive: Path) -> dict[str, H2ProtocolMetadata]:
    """Read the test overview and index it by normalized test code."""

    load_workbook = _require_openpyxl()
    with zipfile.ZipFile(archive) as bundle:
        overview_name = next(
            (name for name in bundle.namelist() if "overview" in name.lower()), None
        )
        if overview_name is None:
            raise ValueError(f"No overview workbook found in {archive}")
        workbook = load_workbook(
            BytesIO(bundle.read(overview_name)), read_only=True, data_only=True
        )
    worksheet = workbook.worksheets[0]
    headers = [_text(value) for value in next(worksheet.iter_rows(values_only=True))]
    lookup = {header: index for index, header in enumerate(headers)}
    required = {"Fill Number", "Lab Test #", "Test #", "P", "Tank Size", "Chamber Temp", "Fill Rate", "Notes"}
    missing = required - set(headers)
    if missing:
        raise ValueError(f"Overview workbook is missing columns: {sorted(missing)}")

    result: dict[str, H2ProtocolMetadata] = {}
    for row in worksheet.iter_rows(min_row=2, values_only=True):
        test_code = _text(row[lookup["Test #"]])
        if not test_code:
            continue
        pressure = _number(row[lookup["P"]])
        tank = _capacity_kg(row[lookup["Tank Size"]])
        chamber = _number(row[lookup["Chamber Temp"]])
        fill_rate = _number(row[lookup["Fill Rate"]])
        lab_number = _number(row[lookup["Lab Test #"]])
        metadata = H2ProtocolMetadata(
            fill_number=_text(row[lookup["Fill Number"]]),
            lab_test_number=None if lab_number is None else int(lab_number),
            test_code=test_code,
            nominal_pressure_mpa=pressure,
            tank_capacity_kg=tank,
            chamber_temperature_c=chamber,
            scheduled_aprr_mpa_min=fill_rate,
            notes=_text(row[lookup["Notes"]]),
        )
        result[_test_key(test_code)] = metadata
    return result


def _sheet_test_code(title: str) -> str:
    match = re.search(r"\btest\s+(.+?)(?:,|$)", title, re.IGNORECASE)
    value = _text(match.group(1) if match else title)
    return re.sub(r"\s+\d{4}[A-Za-z]+\d{1,2}.*$", "", value).strip()


def _select_numeric_column(
    rows: Sequence[tuple], headers: Mapping[str, int], candidates: Sequence[str],
    *, min_finite_fraction: float = 0.95,
) -> tuple[str, np.ndarray]:
    for name in candidates:
        if name not in headers:
            continue
        values = np.asarray([
            float(row[headers[name]]) if isinstance(row[headers[name]], (int, float)) else np.nan
            for row in rows
        ])
        if np.count_nonzero(np.isfinite(values)) >= 3 and (
            np.isfinite(values).mean() >= min_finite_fraction
        ):
            return name, values
    raise ValueError(f"None of the candidate columns is sufficiently complete: {candidates}")


def read_h2protocol_workbook(
    content: bytes,
    *,
    source_archive: str,
    source_member: str,
    overview: Mapping[str, H2ProtocolMetadata],
    flow_threshold_g_s: float = 1.0,
) -> H2ProtocolTrace:
    """Normalize the active fueling portion of one Powertech workbook."""

    load_workbook = _require_openpyxl()
    workbook = load_workbook(BytesIO(content), read_only=True, data_only=True)
    worksheet = workbook.worksheets[0]
    header_values = next(worksheet.iter_rows(values_only=True))
    headers = {_text(value): index for index, value in enumerate(header_values) if _text(value)}
    time_candidates = ("Time (s)", "Elapsed")
    if not any(name in headers for name in time_candidates):
        raise ValueError(f"{source_member} is missing an elapsed-time column")
    if not any(name in headers for name in ("SOC (%)", "SOC %", "SOC")):
        raise ValueError(f"{source_member} is missing an SOC column")
    if not any(name in headers for name in ("FM(R)", "FM")):
        raise ValueError(f"{source_member} is missing a mass-flow column")

    # The source workbooks contain formatted trailing rows and occasional notes
    # below the measurement table.  Exclude them before assessing channel
    # completeness; otherwise a valid test can look mostly empty.
    time_header = next(name for name in time_candidates if name in headers)
    time_index = headers[time_header]
    rows = [
        row for row in worksheet.iter_rows(min_row=2, values_only=True)
        if isinstance(row[time_index], (int, float))
    ]
    if len(rows) < 3:
        raise ValueError(f"{source_member} has fewer than three timed samples")
    time_name, time_all = _select_numeric_column(rows, headers, (time_header,))
    flow_name, flow_all = _select_numeric_column(
        rows, headers, ("FM(R)", "FM"), min_finite_fraction=0.01
    )
    active_indices = np.flatnonzero(
        np.isfinite(time_all) & np.isfinite(flow_all) & (flow_all > flow_threshold_g_s)
    )
    if active_indices.size < 3:
        raise ValueError(f"{source_member} has no sustained active fill")
    active_rows = rows[int(active_indices[0]): int(active_indices[-1]) + 1]

    _, time = _select_numeric_column(active_rows, headers, (time_name,))
    pressure_source, pressure = _select_numeric_column(
        active_rows, headers, ("Ptank", "Pinlet", "Precep")
    )
    try:
        soc_source, soc = _select_numeric_column(
            active_rows, headers, ("SOC (%)", "SOC %", "SOC")
        )
    except ValueError:
        soc_source = "derived-density-from-pressure-temperature"
        soc = np.full(len(active_rows), np.nan)
    _, flow = _select_numeric_column(active_rows, headers, (flow_name,))
    _, inlet_temperature = _select_numeric_column(
        active_rows, headers, ("Tinlet_G", "Trecep_G", "Tpc_out")
    )
    _, chamber = _select_numeric_column(active_rows, headers, ("Tchamber",))
    tank_columns = []
    # Early tests expose one spatially averaged gas-temperature channel while
    # later tests expose three in-tank thermocouples.
    tank_names = tuple(
        name for name in headers if re.match(r"^Ttank\d*_G(?:\s|\(|$)", name)
    )
    for name in tank_names:
        if name in headers:
            values = np.asarray([
                float(row[headers[name]]) if isinstance(row[headers[name]], (int, float)) else np.nan
                for row in active_rows
            ])
            if np.isfinite(values).mean() >= 0.95:
                tank_columns.append(values)
    if not tank_columns:
        raise ValueError(f"{source_member} has no complete tank-temperature channel")
    tank_stack = np.vstack(tank_columns)
    if not np.any(np.isfinite(soc)):
        mean_tank_temperature_k = np.mean(tank_stack, axis=0) + 273.15
        reference_density = float(
            PropsSI("Dmass", "P", 70.0e6, "T", 288.15, "Hydrogen")
        )
        soc = np.asarray([
            100.0 * float(PropsSI(
                "Dmass", "P", pressure_mpa * 1.0e6, "T", temperature_k,
                "Hydrogen",
            )) / reference_density
            for pressure_mpa, temperature_k in zip(pressure, mean_tank_temperature_k)
        ])

    valid = (
        np.isfinite(time) & np.isfinite(pressure) & np.isfinite(soc)
        & np.isfinite(flow) & np.isfinite(inlet_temperature) & np.isfinite(chamber)
        & np.all(np.isfinite(tank_stack), axis=0)
    )
    time = time[valid]
    if len(time) < 3 or np.any(np.diff(time) <= 0.0):
        raise ValueError(f"{source_member} has invalid active-fill timestamps")
    tank_stack = tank_stack[:, valid]
    test_code = _sheet_test_code(worksheet.title)
    metadata = overview.get(_test_key(test_code))
    if metadata is None:
        metadata = H2ProtocolMetadata(
            fill_number="", lab_test_number=None, test_code=test_code,
            nominal_pressure_mpa=None, tank_capacity_kg=None,
            chamber_temperature_c=None, scheduled_aprr_mpa_min=None, notes="",
        )
    case_suffix = (
        f"L{metadata.lab_test_number:02d}"
        if metadata.lab_test_number is not None
        else _test_key(test_code)
    )
    return H2ProtocolTrace(
        case_id=f"H2P-{case_suffix}",
        metadata=metadata,
        source_archive=source_archive,
        source_member=source_member,
        pressure_source=pressure_source,
        soc_source=soc_source,
        time_s=time - time[0],
        pressure_mpa=pressure[valid],
        soc_percent=soc[valid],
        mass_flow_g_s=flow[valid],
        tank_temperature_mean_c=np.mean(tank_stack, axis=0),
        tank_temperature_max_c=np.max(tank_stack, axis=0),
        inlet_gas_temperature_c=inlet_temperature[valid],
        chamber_temperature_c=chamber[valid],
    )


def iter_h2protocol_traces(raw_directory: Path) -> Iterable[H2ProtocolTrace]:
    archives = sorted(raw_directory.glob("*.zip"))
    if not archives:
        raise FileNotFoundError(f"No H2Protocol ZIP archives found under {raw_directory}")
    overview_archive = next((path for path in archives if "Files 1" in path.name), archives[0])
    overview = read_h2protocol_overview(overview_archive)
    seen: dict[str, bytes] = {}
    used_lab_tests: set[int] = set()
    for archive in archives:
        with zipfile.ZipFile(archive) as bundle:
            for member in sorted(bundle.namelist()):
                if not member.lower().endswith(".xlsx") or "overview" in member.lower():
                    continue
                content = bundle.read(member)
                if member in seen:
                    if seen[member] != content:
                        raise ValueError(f"Conflicting duplicate workbook: {member}")
                    continue
                seen[member] = content
                trace = read_h2protocol_workbook(
                    content, source_archive=archive.name, source_member=member,
                    overview=overview,
                )
                metadata = trace.metadata
                if metadata.lab_test_number is None:
                    sheet_key = _test_key(metadata.test_code)
                    candidates = [
                        candidate for key, candidate in overview.items()
                        if key.startswith(sheet_key)
                        and candidate.lab_test_number not in used_lab_tests
                    ]
                    if not candidates and "fmvaliation" in sheet_key:
                        candidates = [
                            candidate for key, candidate in overview.items()
                            if key.startswith("fm")
                            and candidate.lab_test_number not in used_lab_tests
                        ]
                    if not candidates:
                        raise ValueError(
                            f"No unused overview row matches {member} sheet test "
                            f"code {metadata.test_code!r}"
                        )
                    metadata = candidates[0]
                    trace = replace(
                        trace,
                        metadata=metadata,
                        case_id=f"H2P-L{metadata.lab_test_number:02d}",
                    )
                if metadata.lab_test_number in used_lab_tests:
                    raise ValueError(
                        f"Overview lab test {metadata.lab_test_number} matched more than once"
                    )
                if metadata.lab_test_number is not None:
                    used_lab_tests.add(metadata.lab_test_number)
                yield trace


def _mc_default_case_code(source_member: str) -> tuple[str, float]:
    name = Path(source_member).name
    if "9.8kg" in name:
        return "9.8KG", 9.8
    match = re.search(r"Validation Test\s+(\d-[A-Z])", name, re.IGNORECASE)
    if match is None:
        raise ValueError(f"Cannot identify MC Default case from {source_member}")
    return match.group(1).upper(), 4.7


def _protocol_pressure_profile(workbook) -> tuple[np.ndarray, np.ndarray]:
    """Read the publisher-supplied pressure schedule from the small Sheet1 table."""

    worksheet = next(
        (sheet for sheet in workbook.worksheets if sheet.title.lower() == "sheet1"),
        None,
    )
    if worksheet is None:
        raise ValueError("MC Default workbook has no Sheet1 protocol schedule")
    points: list[tuple[float, float]] = []
    for row in worksheet.iter_rows(min_row=2, values_only=True):
        if (
            len(row) >= 2
            and isinstance(row[0], (int, float))
            and isinstance(row[1], (int, float))
        ):
            points.append((float(row[0]), float(row[1])))
    if len(points) < 2:
        raise ValueError("MC Default protocol schedule has fewer than two points")
    points.sort()
    time = np.asarray([item[0] for item in points], dtype=float)
    pressure = np.asarray([item[1] for item in points], dtype=float)
    if np.any(np.diff(time) <= 0.0) or np.any(np.diff(pressure) < 0.0):
        raise ValueError("MC Default protocol schedule is not monotonic")
    return time - time[0], pressure


def read_mc_default_workbook(
    content: bytes,
    *,
    source_archive: str,
    source_member: str,
    flow_threshold_g_s: float = 1.0,
) -> H2ProtocolTrace:
    """Normalize one prospectively frozen Powertech MC Default bench test."""

    load_workbook = _require_openpyxl()
    workbook = load_workbook(BytesIO(content), read_only=True, data_only=True)
    worksheet = next(
        (
            sheet for sheet in workbook.worksheets
            if any(
                _text(value) in {"FM", "FM(R)"}
                for value in next(sheet.iter_rows(values_only=True))
            )
        ),
        None,
    )
    if worksheet is None:
        raise ValueError(f"{source_member} has no measurement worksheet")
    header_values = next(worksheet.iter_rows(values_only=True))
    headers = {_text(value): index for index, value in enumerate(header_values) if _text(value)}
    time_name = next(
        (name for name in ("Time (s)", "Elapsed") if name in headers), None
    )
    if time_name is None:
        raise ValueError(f"{source_member} is missing an elapsed-time column")
    rows = [
        row for row in worksheet.iter_rows(min_row=2, values_only=True)
        if isinstance(row[headers[time_name]], (int, float))
    ]
    _, time_all = _select_numeric_column(rows, headers, (time_name,))
    flow_name, flow_all = _select_numeric_column(
        rows, headers, ("FM(R)", "FM"), min_finite_fraction=0.01
    )
    active_indices = np.flatnonzero(
        np.isfinite(time_all) & np.isfinite(flow_all) & (flow_all > flow_threshold_g_s)
    )
    if active_indices.size < 3:
        raise ValueError(f"{source_member} has no sustained active fill")
    active_rows = rows[int(active_indices[0]): int(active_indices[-1]) + 1]
    _, time = _select_numeric_column(active_rows, headers, (time_name,))
    pressure_source, pressure = _select_numeric_column(
        active_rows, headers, ("Ptank", "Pinlet", "Precep")
    )
    soc_source, soc = _select_numeric_column(
        active_rows, headers, ("SOC (%)", "SOC %", "SOC")
    )
    _, flow = _select_numeric_column(active_rows, headers, (flow_name,))
    _, inlet_temperature = _select_numeric_column(
        active_rows, headers, ("Tinlet_G", "Trecep_G", "Tpc_out")
    )
    _, chamber = _select_numeric_column(
        active_rows, headers, ("Tchamber", "Tchamb")
    )
    tank_columns = []
    for name in headers:
        if not re.match(r"^Ttank\d*_G(?:\s|\(|$)", name):
            continue
        values = np.asarray([
            float(row[headers[name]])
            if isinstance(row[headers[name]], (int, float)) else np.nan
            for row in active_rows
        ])
        if np.isfinite(values).mean() >= 0.95:
            tank_columns.append(values)
    if not tank_columns:
        raise ValueError(f"{source_member} has no complete tank-temperature channel")
    tank_stack = np.vstack(tank_columns)

    optional_pressures: dict[str, np.ndarray | None] = {}
    for name in ("875PT1", "875PT3"):
        try:
            _, values = _select_numeric_column(active_rows, headers, (name,))
        except ValueError:
            values = None
        optional_pressures[name] = values
    valid = (
        np.isfinite(time) & np.isfinite(pressure) & np.isfinite(soc)
        & np.isfinite(flow) & np.isfinite(inlet_temperature) & np.isfinite(chamber)
        & np.all(np.isfinite(tank_stack), axis=0)
    )
    time = time[valid]
    if len(time) < 3 or np.any(np.diff(time) <= 0.0):
        raise ValueError(f"{source_member} has invalid active-fill timestamps")
    case_code, tank_capacity_kg = _mc_default_case_code(source_member)
    protocol_time, protocol_pressure = _protocol_pressure_profile(workbook)
    metadata = H2ProtocolMetadata(
        fill_number="",
        lab_test_number=None,
        test_code=case_code,
        nominal_pressure_mpa=70.0,
        tank_capacity_kg=tank_capacity_kg,
        chamber_temperature_c=float(np.median(chamber[valid])),
        scheduled_aprr_mpa_min=None,
        notes="Powertech Labs MC Default bench test",
    )
    return H2ProtocolTrace(
        case_id=f"H2P-MC-{case_code}",
        metadata=metadata,
        source_archive=source_archive,
        source_member=source_member,
        pressure_source=pressure_source,
        soc_source=soc_source,
        time_s=time - time[0],
        pressure_mpa=pressure[valid],
        soc_percent=soc[valid],
        mass_flow_g_s=flow[valid],
        tank_temperature_mean_c=np.mean(tank_stack[:, valid], axis=0),
        tank_temperature_max_c=np.max(tank_stack[:, valid], axis=0),
        inlet_gas_temperature_c=inlet_temperature[valid],
        chamber_temperature_c=chamber[valid],
        source_pressure_1_mpa=(
            optional_pressures["875PT1"][valid]
            if optional_pressures["875PT1"] is not None else None
        ),
        source_pressure_3_mpa=(
            optional_pressures["875PT3"][valid]
            if optional_pressures["875PT3"] is not None else None
        ),
        protocol_pressure_time_s=protocol_time,
        protocol_pressure_mpa=protocol_pressure,
    )


def iter_mc_default_traces(
    archive: Path,
    selected_members: Sequence[str],
) -> Iterable[H2ProtocolTrace]:
    """Yield exactly the prospectively selected MC Default workbooks."""

    with zipfile.ZipFile(archive) as bundle:
        available = set(bundle.namelist())
        missing = set(selected_members) - available
        if missing:
            raise ValueError(f"MC Default archive is missing: {sorted(missing)}")
        for member in selected_members:
            yield read_mc_default_workbook(
                bundle.read(member),
                source_archive=archive.name,
                source_member=member,
            )


def compare_traces(
    *,
    experimental_time_s: Sequence[float],
    experimental_pressure_mpa: Sequence[float],
    experimental_temperature_c: Sequence[float],
    experimental_soc_percent: Sequence[float],
    predicted_time_s: Sequence[float],
    predicted_pressure_mpa: Sequence[float],
    predicted_temperature_c: Sequence[float],
    predicted_soc_percent: Sequence[float],
) -> TraceAgreement:
    """Compare predictions on the experimental clock without time warping."""

    exp_t = np.asarray(experimental_time_s, dtype=float)
    pred_t = np.asarray(predicted_time_s, dtype=float)
    if len(exp_t) < 2 or len(pred_t) < 2:
        raise ValueError("Both traces require at least two samples")
    if np.any(np.diff(exp_t) <= 0.0) or np.any(np.diff(pred_t) <= 0.0):
        raise ValueError("Trace time values must be strictly increasing")
    mask = exp_t <= pred_t[-1]
    if np.count_nonzero(mask) < 2:
        raise ValueError("Predicted trace does not overlap the experiment")
    exp_t = exp_t[mask]
    exp_p = np.asarray(experimental_pressure_mpa, dtype=float)[mask]
    exp_temp = np.asarray(experimental_temperature_c, dtype=float)[mask]
    exp_soc = np.asarray(experimental_soc_percent, dtype=float)[mask]
    pred_p = np.interp(exp_t, pred_t, np.asarray(predicted_pressure_mpa, dtype=float))
    pred_temp = np.interp(exp_t, pred_t, np.asarray(predicted_temperature_c, dtype=float))
    pred_soc = np.interp(exp_t, pred_t, np.asarray(predicted_soc_percent, dtype=float))

    def rmse(actual: np.ndarray, predicted: np.ndarray) -> float:
        return float(np.sqrt(np.mean(np.square(predicted - actual))))

    def mae(actual: np.ndarray, predicted: np.ndarray) -> float:
        return float(np.mean(np.abs(predicted - actual)))

    return TraceAgreement(
        sample_count=len(exp_t),
        pressure_rmse_mpa=rmse(exp_p, pred_p),
        pressure_mae_mpa=mae(exp_p, pred_p),
        pressure_final_error_mpa=float(pred_p[-1] - exp_p[-1]),
        temperature_rmse_c=rmse(exp_temp, pred_temp),
        temperature_mae_c=mae(exp_temp, pred_temp),
        temperature_peak_error_c=float(np.max(pred_temp) - np.max(exp_temp)),
        soc_rmse_percentage_points=rmse(exp_soc, pred_soc),
        soc_final_error_percentage_points=float(pred_soc[-1] - exp_soc[-1]),
    )


def _sheet_records(workbook, sheet_name: str) -> dict[str, dict[str, object]]:
    worksheet = workbook[sheet_name]
    headers = [_text(value) for value in next(worksheet.iter_rows(values_only=True))]
    records: dict[str, dict[str, object]] = {}
    for row in worksheet.iter_rows(min_row=2, values_only=True):
        event_id = _text(row[0])
        if event_id:
            records[event_id] = dict(zip(headers, row))
    return records


def read_hiad_hrs_cases(path: Path, *, include_related_refuelling: bool = False) -> list[HIADCase]:
    """Join HIAD sheets and return the hydrogen-refuelling-station subset."""

    load_workbook = _require_openpyxl()
    workbook = load_workbook(path, read_only=True, data_only=True)
    events = _sheet_records(workbook, "EVENTS")
    facilities = _sheet_records(workbook, "FACILITY")
    lessons = _sheet_records(workbook, "LESSONS LEARNT")
    nature = _sheet_records(workbook, "EVENT NATURE")
    references = _sheet_records(workbook, "REFERENCES")
    result = []
    for event_id, facility in facilities.items():
        application = _text(facility.get("Application"))
        combined = " ".join((
            application, _text(facility.get("Sub-application")),
            _text(facility.get("Event Title")), _text(facility.get("Event full description")),
        )).lower()
        is_core = application == "Hydrogen refuelling station"
        is_related = "refuel" in combined or "fueling station" in combined
        if not is_core and not (include_related_refuelling and is_related):
            continue
        event = events.get(event_id, {})
        lesson = lessons.get(event_id, {})
        event_nature = nature.get(event_id, {})
        reference = references.get(event_id, {})
        links = tuple(
            _text(reference.get(name))
            for name in (
                "1st Reference & weblink", "2nd Reference &weblink",
                "3rd Reference & weblink", "4th Reference & weblink",
                "5th Documents & links", "6th Documents & links", "7th Documents & links",
            )
            if _text(reference.get(name))
        )
        result.append(HIADCase(
            event_id=event_id,
            quality=_text(event.get("Q")),
            title=_text(event.get("Event Title")),
            description=_text(event.get("Event full description")),
            initiating_system=_text(event.get("Event Initiating system")),
            physical_effect=_text(event.get("Classification of the physical effects")),
            consequence_nature=_text(event.get("Nature of the consequences")),
            root_cause_summary=_text(event.get("Summary root causes")),
            application=application,
            sub_application=_text(facility.get("Sub-application")),
            supply_chain_stage=_text(facility.get("Hydrogen supply chain stage")),
            operational_condition=_text(facility.get("Operational condition")),
            emergency_action=_text(event_nature.get("Emergency action")),
            lesson_learnt=_text(lesson.get("Lesson Learnt")),
            corrective_measures=_text(lesson.get("Corrective Measures")),
            references=links,
        ))
    return sorted(
        result,
        key=lambda item: (
            0, int(item.event_id)
        ) if item.event_id.isdigit() else (1, item.event_id),
    )


def _metadata_number(rows: Sequence[Mapping[str, str]], label: str) -> float | None:
    pattern = re.compile(rf"{re.escape(label)}\s*-\s*([-+]?\d+(?:\.\d+)?)", re.I)
    for row in rows[:12]:
        for value in row.values():
            match = pattern.search(value or "")
            if match:
                return float(match.group(1))
    return None


def _dispersion_article_ids(readme: Path) -> dict[str, str]:
    result = {}
    for line in readme.read_text(encoding="utf-8-sig").splitlines():
        match = re.search(r"(23_FFI_P101_T\d+)\s+-\s+(test-\S+)", line)
        if match:
            result[match.group(1)] = match.group(2)
    return result


def iter_dispersion_experiments(raw_directory: Path) -> Iterable[DispersionExperiment]:
    """Read the CC BY 4.0 USN/FFI channel-dispersion measurements."""

    article_ids = _dispersion_article_ids(raw_directory / "ReadMe.txt")
    archives = sorted(raw_directory.glob("*.zip"))
    if not archives:
        raise FileNotFoundError(f"No dispersion ZIP archives found under {raw_directory}")
    for archive in archives:
        with zipfile.ZipFile(archive) as bundle:
            member = next(
                (name for name in bundle.namelist() if name.lower().endswith(".csv")),
                None,
            )
            if member is None:
                raise ValueError(f"No CSV found in {archive}")
            rows = list(csv.DictReader(StringIO(bundle.read(member).decode("utf-8-sig"))))
        if not rows:
            raise ValueError(f"No rows found in {archive}")
        sensor_columns = tuple(
            name for name in rows[0]
            if name and name.startswith("sensor ") and "h2 concentration" in name
        )
        sensor_records = []
        for row in rows:
            try:
                time_s = float(row["h2 sensor time"])
                values = [float(row[name]) for name in sensor_columns]
            except (KeyError, TypeError, ValueError):
                continue
            sensor_records.append((time_s, values))
        if len(sensor_records) < 3:
            raise ValueError(f"Insufficient hydrogen-sensor data in {archive}")
        mean_flow = _metadata_number(rows, "mean mass flow")
        baseline = _metadata_number(rows, "baseline duration [s]")
        duration = _metadata_number(rows, "filling duration [s]")
        if mean_flow is None or baseline is None or duration is None:
            raise ValueError(f"Missing experiment metadata in {archive}")
        case_id = archive.stem
        yield DispersionExperiment(
            case_id=case_id,
            article_test_id=article_ids.get(case_id, ""),
            source_archive=archive.name,
            mean_mass_flow_g_s=mean_flow,
            mean_filling_pressure_bar=_metadata_number(rows, "mean pressure [bar]"),
            channel_temperature_c=_metadata_number(rows, "mean channel temperature [C]"),
            baseline_duration_s=baseline,
            filling_duration_s=duration,
            sensor_ids=tuple(re.search(r"sensor (\d+)", name).group(1) for name in sensor_columns),
            sensor_time_s=np.asarray([item[0] for item in sensor_records]),
            concentrations_percent=np.asarray([item[1] for item in sensor_records]),
        )
