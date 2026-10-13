"""Post-access diagnostic for PRESLHY E3.1 concentration-channel delay.

The E3.1 workbooks contain five hydrogen-concentration channels in addition to
vessel/nozzle pressure.  The released documentation reports a non-negligible
transport/measurement delay for the concentration system.  This module keeps
that observation explicit by comparing uncorrected and delay-corrected event
times.  It is deliberately a diagnostic: it does not fit the station model,
change runtime detector routing, or close an independent validation gate.
"""

from __future__ import annotations

from dataclasses import asdict, dataclass
from io import BytesIO
from pathlib import Path
import re
from typing import Iterable
from zipfile import ZipFile

import numpy as np
from openpyxl import load_workbook


DOCUMENTED_TRANSPORT_DELAY_S = 3.75
CONCENTRATION_CHANNELS = ("C3", "C2", "C1", "C1-", "C4")
_DIAMETER_BY_TOKEN = {"D05": 0.5, "D1": 1.0, "D2": 2.0, "D4": 4.0}


@dataclass(frozen=True)
class SensorDelaySummary:
    channel: str
    baseline_vol_percent: float
    peak_vol_percent: float
    onset_time_s: float | None
    peak_time_s: float | None
    raw_onset_lag_s: float | None
    corrected_onset_lag_s: float | None
    raw_peak_lag_s: float | None
    corrected_peak_lag_s: float | None


@dataclass(frozen=True)
class ConcentrationTraceSummary:
    case_id: str
    source_package: str
    source_member: str
    thermal_domain: str
    nozzle_diameter_mm: float
    samples: int
    release_onset_time_s: float | None
    sensors: tuple[SensorDelaySummary, ...]


def _numeric(value: object) -> float | None:
    try:
        value = float(value)
    except (TypeError, ValueError):
        return None
    return value if np.isfinite(value) else None


def _normalise(label: object) -> str:
    return re.sub(r"[^a-z0-9]+", "", str(label or "").lower())


def _channel_name(label: object) -> str | None:
    # The workbook uses both ``C1 -`` and ``C1-`` spellings.  Preserve the
    # hyphen before applying the general header normalisation.
    raw = str(label or "").strip().upper().replace(" ", "")
    if raw in {"C3", "C2", "C1", "C4"}:
        return raw
    if raw in {"C1-", "C1–"}:
        return "C1-"
    return None


def _header_columns(sheet) -> tuple[int, int, dict[str, int]]:
    rows = list(
        sheet.iter_rows(min_row=1, max_row=6, max_col=min(sheet.max_column, 32), values_only=True)
    )
    if len(rows) < 6:
        raise ValueError("E3.1 concentration worksheet has fewer than six header rows")
    time_column = None
    pressure_column = None
    channels: dict[str, int] = {}
    for column in range(max(len(row) for row in rows)):
        labels = [row[column] for row in rows if column < len(row)]
        normal = {_normalise(value) for value in labels}
        if "xvalue" in normal or "times" in normal or "synchronizedtimes" in normal:
            time_column = column
        if "pvessel" in normal or "druck1" in normal:
            pressure_column = column
        for label in labels:
            channel = _channel_name(label)
            if channel is not None:
                channels.setdefault(channel, column)
    if time_column is None or pressure_column is None:
        raise ValueError("E3.1 concentration worksheet time/pressure column not found")
    missing = [channel for channel in CONCENTRATION_CHANNELS if channel not in channels]
    if missing:
        raise ValueError(f"E3.1 concentration worksheet missing channels: {missing}")
    return time_column, pressure_column, channels


def _read_concentration_sheet(sheet) -> tuple[np.ndarray, np.ndarray, np.ndarray]:
    time_column, pressure_column, channels = _header_columns(sheet)
    ordered = [channels[channel] for channel in CONCENTRATION_CHANNELS]
    values: list[tuple[float, ...]] = []
    for row in sheet.iter_rows(min_row=7, max_col=max(time_column, pressure_column, *ordered) + 1, values_only=True):
        time = _numeric(row[time_column] if time_column < len(row) else None)
        pressure = _numeric(row[pressure_column] if pressure_column < len(row) else None)
        measurements = [_numeric(row[column] if column < len(row) else None) for column in ordered]
        if time is None or pressure is None or any(value is None for value in measurements):
            continue
        values.append((time, float(pressure), *(float(value) for value in measurements)))
    if len(values) < 20:
        raise ValueError("E3.1 concentration worksheet has fewer than 20 numeric samples")
    array = np.asarray(values, dtype=float)
    order = np.argsort(array[:, 0], kind="stable")
    array = array[order]
    # Remove duplicate timestamps while preserving the first complete row.
    _, unique = np.unique(array[:, 0], return_index=True)
    array = array[np.sort(unique)]
    return array[:, 0], array[:, 1], array[:, 2:]


def _release_onset(time_s: np.ndarray, vessel_pressure_bar: np.ndarray) -> float | None:
    """Find a stable 0.5% vessel-pressure drop from the initial value."""
    count = min(max(5, int(round(1.0 / max(np.median(np.diff(time_s)), 1e-6))),), len(time_s))
    baseline = float(np.median(vessel_pressure_bar[:count]))
    threshold = baseline * 0.995
    below = vessel_pressure_bar <= threshold
    run = max(3, int(round(0.5 / max(np.median(np.diff(time_s)), 1e-6))))
    for index in range(0, len(time_s) - run + 1):
        if bool(np.all(below[index : index + run])):
            return float(time_s[index])
    return None


def _sensor_summary(time_s: np.ndarray, values: np.ndarray, release_time_s: float | None, channel: str) -> SensorDelaySummary:
    baseline_count = min(max(5, int(round(5.0 / max(np.median(np.diff(time_s)), 1e-6))),), len(time_s))
    baseline = float(np.median(values[:baseline_count]))
    peak_index = int(np.argmax(values))
    peak = float(values[peak_index])
    excursion = peak - baseline
    onset_time: float | None = None
    if excursion > 0.0:
        threshold = baseline + 0.10 * excursion
        run = max(3, int(round(0.3 / max(np.median(np.diff(time_s)), 1e-6))))
        above = values >= threshold
        for index in range(0, len(time_s) - run + 1):
            if bool(np.all(above[index : index + run])):
                onset_time = float(time_s[index])
                break
    peak_time = float(time_s[peak_index]) if np.isfinite(peak) else None

    def lag(event_time: float | None, corrected: bool = False) -> float | None:
        if event_time is None or release_time_s is None:
            return None
        value = event_time - release_time_s
        return float(value - DOCUMENTED_TRANSPORT_DELAY_S) if corrected else float(value)

    return SensorDelaySummary(
        channel=channel,
        baseline_vol_percent=baseline,
        peak_vol_percent=peak,
        onset_time_s=onset_time,
        peak_time_s=peak_time,
        raw_onset_lag_s=lag(onset_time),
        corrected_onset_lag_s=lag(onset_time, corrected=True),
        raw_peak_lag_s=lag(peak_time),
        corrected_peak_lag_s=lag(peak_time, corrected=True),
    )


def summarise_workbook(payload: bytes, *, source_package: str, source_member: str, nozzle_diameter_mm: float) -> ConcentrationTraceSummary:
    workbook = load_workbook(BytesIO(payload), read_only=True, data_only=True)
    try:
        sheets = [sheet for sheet in workbook.worksheets if "ch2" in sheet.title.lower()]
        if not sheets:
            raise ValueError("E3.1 concentration worksheet not found")
        time_s, pressure_bar, concentration = _read_concentration_sheet(sheets[0])
        release_time = _release_onset(time_s, pressure_bar)
        sensors = tuple(
            _sensor_summary(time_s, concentration[:, index], release_time, channel)
            for index, channel in enumerate(CONCENTRATION_CHANNELS)
        )
    finally:
        workbook.close()
    thermal_domain = "80 K" if "80K" in source_package.upper() else "300 K"
    case_id = f"{source_package}:{source_member}"
    return ConcentrationTraceSummary(
        case_id=case_id,
        source_package=source_package,
        source_member=source_member,
        thermal_domain=thermal_domain,
        nozzle_diameter_mm=float(nozzle_diameter_mm),
        samples=int(time_s.size),
        release_onset_time_s=release_time,
        sensors=sensors,
    )


def _diameter_for_package(package_name: str) -> float:
    upper = package_name.upper()
    for token, diameter in _DIAMETER_BY_TOKEN.items():
        if token in upper:
            return diameter
    raise ValueError(f"no E3.1 nozzle diameter mapping for {package_name}")


def run_archive_diagnostic(raw_root: str | Path) -> dict:
    root = Path(raw_root)
    traces: list[ConcentrationTraceSummary] = []
    errors: list[dict[str, str]] = []
    for archive in sorted(root.glob("PRE*.zip")):
        diameter = _diameter_for_package(archive.name)
        with ZipFile(archive) as package:
            for member in package.namelist():
                if not member.lower().endswith(".xlsx"):
                    continue
                try:
                    traces.append(
                        summarise_workbook(
                            package.read(member),
                            source_package=archive.name,
                            source_member=member,
                            nozzle_diameter_mm=diameter,
                        )
                    )
                except (OSError, ValueError, KeyError, TypeError) as exc:
                    errors.append({"source_package": archive.name, "source_member": member, "error": str(exc)})

    raw_onsets = [
        sensor.raw_onset_lag_s
        for trace in traces
        for sensor in trace.sensors
        if sensor.raw_onset_lag_s is not None
    ]
    corrected_onsets = [
        sensor.corrected_onset_lag_s
        for trace in traces
        for sensor in trace.sensors
        if sensor.corrected_onset_lag_s is not None
    ]
    raw_peaks = [
        sensor.raw_peak_lag_s
        for trace in traces
        for sensor in trace.sensors
        if sensor.raw_peak_lag_s is not None
    ]
    corrected_peaks = [
        sensor.corrected_peak_lag_s
        for trace in traces
        for sensor in trace.sensors
        if sensor.corrected_peak_lag_s is not None
    ]

    def stats(values: Iterable[float]) -> dict[str, float | int | None]:
        data = np.asarray(list(values), dtype=float)
        if data.size == 0:
            return {"count": 0, "median_s": None, "p10_s": None, "p90_s": None}
        return {
            "count": int(data.size),
            "median_s": float(np.median(data)),
            "p10_s": float(np.percentile(data, 10)),
            "p90_s": float(np.percentile(data, 90)),
        }

    return {
        "schema_version": 1,
        "diagnostic_type": "POST_ACCESS_CONCENTRATION_TRANSPORT_DELAY_DIAGNOSTIC",
        "runtime_promotion": False,
        "independent_validation_gate_impact": "none",
        "source_description": {
            "dataset": "PRESLHY E3.1 DisCha public workbooks",
            "doi": "10.35097/1187",
            "concentration_channels": list(CONCENTRATION_CHANNELS),
            "documented_transport_delay_s": DOCUMENTED_TRANSPORT_DELAY_S,
        },
        "archive_count": len(list(root.glob("PRE*.zip"))),
        "trace_count": len(traces),
        "error_count": len(errors),
        "raw_onset_lag": stats(raw_onsets),
        "corrected_onset_lag": stats(corrected_onsets),
        "raw_peak_lag": stats(raw_peaks),
        "corrected_peak_lag": stats(corrected_peaks),
        "traces": [
            {
                **asdict(trace),
                "sensors": [asdict(sensor) for sensor in trace.sensors],
            }
            for trace in traces
        ],
        "errors": errors,
        "interpretation": (
            "Subtracting the documented transport delay is a timestamp diagnostic only. "
            "It must not be used to advance an alarm, change detector routing, or claim "
            "independent station validation without a pre-access frozen protocol and a "
            "separate holdout."
        ),
    }


__all__ = [
    "CONCENTRATION_CHANNELS",
    "DOCUMENTED_TRANSPORT_DELAY_S",
    "ConcentrationTraceSummary",
    "SensorDelaySummary",
    "run_archive_diagnostic",
    "summarise_workbook",
]
