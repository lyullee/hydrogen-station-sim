"""80 K PRESLHY holdout reader corrections.

The original validation reader is hash-locked by the 300 K prospective
protocol.  This module applies the publisher's declared-unit and gas-side
thermocouple interpretation only for the separately frozen 80 K holdout;
the archived 300 K implementation and outcomes remain untouched.
"""

from __future__ import annotations

from io import BytesIO
import math
import re

import numpy as np
from openpyxl import load_workbook

from .preslhy_validation import (
    PreslhyTrace,
    _data_time_column,
    _header_layout,
    _normalise,
    _profile_label,
    read_preslhy_workbook,
)


def _declared_gas_temperature(workbook, case_prefix: str) -> tuple[float, bool]:
    """Read T1/T2/T3 gas sensors, excluding T1o/T2o/T3o wall sensors."""

    candidates = [
        sheet
        for sheet in workbook.worksheets
        if "temp" in sheet.title.lower() and case_prefix in sheet.title
    ] or [sheet for sheet in workbook.worksheets if "temp" in sheet.title.lower()]
    for sheet in candidates:
        try:
            rows, profiles = _header_layout(sheet)
            header_row, time_column = _data_time_column(rows)
        except ValueError:
            continue
        columns = [
            index
            for index, profile in enumerate(profiles)
            if any(
                re.fullmatch(r"t[123](?:degc|c|kelvin)?", value)
                for value in profile
            )
        ]
        samples: list[float] = []
        for row in sheet.iter_rows(
            min_row=header_row + 1,
            min_col=1,
            max_col=max([time_column, *columns], default=time_column) + 1,
            values_only=True,
        ):
            try:
                time_s = float(row[time_column])
            except (TypeError, ValueError, IndexError):
                continue
            if not -0.5 <= time_s < 0.0:
                continue
            for column in columns:
                try:
                    value = float(row[column])
                except (TypeError, ValueError, IndexError):
                    continue
                if not math.isfinite(value):
                    continue
                label = _normalise(_profile_label(rows, column))
                if "kelvin" in label:
                    value_k = value
                elif "degc" in label or "celsius" in label or "gradc" in label:
                    value_k = value + 273.15
                else:
                    value_k = value + 273.15 if value < 200.0 else value
                samples.append(value_k)
        if samples:
            return float(np.median(samples)), False
    return 293.15, True


def read_preslhy_80k_workbook(
    payload: bytes,
    *,
    source_package: str,
    source_member: str,
    nozzle_diameter_mm: float,
) -> PreslhyTrace:
    """Read an 80 K trace without changing the hash-locked 300 K reader."""

    trace = read_preslhy_workbook(
        payload,
        source_package=source_package,
        source_member=source_member,
        nozzle_diameter_mm=nozzle_diameter_mm,
    )
    workbook = load_workbook(BytesIO(payload), read_only=True, data_only=True)
    try:
        case_prefix = source_member.rsplit("/", 1)[-1].split("-")[0].split(".")[0]
        temperature_k, substituted = _declared_gas_temperature(workbook, case_prefix)
    finally:
        workbook.close()
    return trace.__class__(
        case_id=trace.case_id,
        source_package=trace.source_package,
        source_member=trace.source_member,
        nozzle_diameter_mm=trace.nozzle_diameter_mm,
        time_s=trace.time_s,
        pressure_bar_abs=trace.pressure_bar_abs,
        initial_temperature_k=temperature_k,
        ambient_pressure_pa=trace.ambient_pressure_pa,
        pressure_unit_interpretation=trace.pressure_unit_interpretation,
        temperature_substituted=substituted,
        ambient_pressure_substituted=trace.ambient_pressure_substituted,
    )
