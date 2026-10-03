"""Structural adapter for the prospectively reserved PRESLHY E5.1 holdout."""

from __future__ import annotations

from dataclasses import dataclass
from io import BytesIO
import math
from pathlib import PurePosixPath
import re
from typing import Iterable
import zipfile

import numpy as np
from openpyxl import load_workbook

from .preslhy_validation import PreslhyTrace


@dataclass(frozen=True, slots=True)
class E5CaseDefinition:
    case_id: str
    se_member_suffix: str
    pfast_case_id: str
    nozzle_diameter_mm: float
    nominal_pressure_bar: float
    nominal_temperature_class: str
    ignition_distance_cm: float


# Frozen from DisChaIgn-Explanations.pdf before numerical workbook access.
E5_CASES = (
    E5CaseDefinition("20200122-133946", "20200122-133946-SE-39cm5.xlsx", "20200122-133946", 4.0, 50.0, "ambient", 39.5),
    E5CaseDefinition("20200130-124939", "20200130-124937-SE-39cm5.xlsx", "20200130-124939", 4.0, 200.0, "ambient", 39.5),
    E5CaseDefinition("20200131-141755", "20200131-141755-SE-39cm5.xlsx", "20200131-141755", 4.0, 50.0, "cryogenic", 39.5),
    E5CaseDefinition("20200131-120937", "20200131-120937-SE-39cm5.xlsx", "20200131-120937", 4.0, 200.0, "cryogenic", 39.5),
    E5CaseDefinition("20200130-144119", "20200130-144119-SE.xlsx", "20200130-144119", 2.0, 50.0, "ambient", 39.5),
    E5CaseDefinition("20200122-144854", "20200122-144854-SE.xlsx", "20200122-144854", 4.0, 50.0, "ambient", 62.5),
    E5CaseDefinition("20200130-133910", "20200130-133910-SE.xlsx", "20200130-133910", 4.0, 50.0, "ambient", 106.5),
)


def _normalise(value: object) -> str:
    return re.sub(r"[^a-z0-9]+", "", str(value or "").lower())


def _find_sheet(workbook, suffix: str):
    matches = [sheet for sheet in workbook.worksheets if sheet.title.endswith(suffix)]
    if len(matches) != 1:
        raise ValueError(f"expected one {suffix} worksheet, found {len(matches)}")
    return matches[0]


def _header_columns(sheet, header_row: int = 6) -> dict[str, int]:
    values = next(
        sheet.iter_rows(
            min_row=header_row, max_row=header_row, values_only=True
        )
    )
    return {
        _normalise(value): index
        for index, value in enumerate(values)
        if value not in (None, "")
    }


def _column(columns: dict[str, int], *tokens: str) -> int:
    for token in tokens:
        normal = _normalise(token)
        if normal in columns:
            return columns[normal]
    raise ValueError(f"required column not found: {tokens}")


def _numeric_rows(sheet, columns: tuple[int, ...], start_row: int = 7):
    blank_run = 0
    for row in sheet.iter_rows(
        min_row=start_row,
        min_col=1,
        max_col=max(columns) + 1,
        values_only=True,
    ):
        values = tuple(row[index] if index < len(row) else None for index in columns)
        if all(value is None for value in values):
            blank_run += 1
            if blank_run >= 50:
                return
            continue
        blank_run = 0
        try:
            converted = tuple(float(value) for value in values)
        except (TypeError, ValueError):
            continue
        if all(math.isfinite(value) for value in converted):
            yield converted


def _absolute_pressure(pressure_bar: np.ndarray) -> tuple[np.ndarray, str]:
    tail_count = max(20, len(pressure_bar) // 20)
    terminal = float(np.median(pressure_bar[-tail_count:]))
    if -0.5 <= terminal <= 0.5:
        return pressure_bar + 1.01325, "terminal_near_zero_gauge_plus_standard_ambient"
    if 0.5 < terminal <= 1.5:
        return pressure_bar, "terminal_near_ambient_absolute"
    raise ValueError(
        f"ambiguous E5.1 vessel-pressure basis; terminal median={terminal:.3f} bar"
    )


def _initial_temperature(workbook) -> float:
    sheet = _find_sheet(workbook, "-TE")
    columns = _header_columns(sheet)
    time_column = _column(columns, "Time [s]")
    temperature_columns = (
        _column(columns, "Behälter_Innen", "Beh鋖ter_Innen"),
        _column(columns, "Behälter_Mitte", "Beh鋖ter_Mitte"),
        _column(columns, "Behälter_Aussen", "Beh鋖ter_Aussen"),
    )
    samples: list[float] = []
    for values in _numeric_rows(
        sheet, (time_column, *temperature_columns)
    ):
        time_s, *temperatures = values
        if -0.5 <= time_s < 0.0:
            samples.extend(temperature for temperature in temperatures if temperature > 0.0)
    if not samples:
        raise ValueError("no finite in-vessel temperature samples in the pre-release window")
    return float(np.median(samples))


def read_preslhy_e5_workbook(
    payload: bytes,
    *,
    source_package: str,
    source_member: str,
    case: E5CaseDefinition,
) -> PreslhyTrace:
    workbook = load_workbook(BytesIO(payload), read_only=True, data_only=True)
    try:
        pressure_sheet = _find_sheet(workbook, "-presV")
        columns = _header_columns(pressure_sheet)
        time_column = _column(columns, "Time [s]")
        pressure_column = _column(
            columns, "Druck_Behälter", "Druck_Beh鋖ter"
        )
        samples = list(
            _numeric_rows(pressure_sheet, (time_column, pressure_column))
        )
        if len(samples) < 20:
            raise ValueError("fewer than 20 finite synchronized source-pressure samples")
        values = np.asarray(samples, dtype=float)
        order = np.argsort(values[:, 0], kind="stable")
        time_s = values[order, 0]
        pressure_bar, pressure_basis = _absolute_pressure(values[order, 1])
        unique = np.concatenate(([True], np.diff(time_s) > 0.0))
        time_s, pressure_bar = time_s[unique], pressure_bar[unique]
        if not np.any(time_s < 0.0) or not np.any(time_s >= 0.1):
            raise ValueError("synchronized trace does not span the release onset")
        temperature_k = _initial_temperature(workbook)
        return PreslhyTrace(
            case_id=case.case_id,
            source_package=source_package,
            source_member=source_member,
            nozzle_diameter_mm=case.nozzle_diameter_mm,
            time_s=time_s,
            pressure_bar_abs=pressure_bar,
            initial_temperature_k=temperature_k,
            ambient_pressure_pa=101_325.0,
            pressure_unit_interpretation=pressure_basis,
            temperature_substituted=False,
            ambient_pressure_substituted=True,
        )
    finally:
        workbook.close()


def iter_preslhy_e5_traces(archive: str | PurePosixPath) -> Iterable[tuple[E5CaseDefinition, PreslhyTrace]]:
    archive_name = str(archive)
    with zipfile.ZipFile(archive_name) as bundle:
        members = bundle.namelist()
        for case in E5_CASES:
            matches = [
                member for member in members
                if PurePosixPath(member).name == case.se_member_suffix
            ]
            if len(matches) != 1:
                raise ValueError(
                    f"expected one member named {case.se_member_suffix}, found {len(matches)}"
                )
            member = matches[0]
            yield case, read_preslhy_e5_workbook(
                bundle.read(member),
                source_package=PurePosixPath(archive_name).name,
                source_member=member,
                case=case,
            )


__all__ = [
    "E5_CASES",
    "E5CaseDefinition",
    "iter_preslhy_e5_traces",
    "read_preslhy_e5_workbook",
]
