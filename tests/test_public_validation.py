from __future__ import annotations

from io import BytesIO
import hashlib
import json
from pathlib import Path
import zipfile

import numpy as np
from openpyxl import Workbook
import pytest

from h2station.public_validation import (
    _active_fill_bounds,
    compare_traces,
    DispersionExperiment,
    evaluate_dispersion_detector_logic,
    first_persistent_threshold_time,
    iter_dispersion_experiments,
    iter_mc_default_traces,
    read_mc_default_workbook,
    read_h2protocol_overview,
    read_h2protocol_workbook,
    read_hiad_hrs_cases,
)
from h2station.api import SimulationInput
from h2station.risk.runtime_backend import UnavailableHyRAMBackend
from h2station.scenario import (
    ReferenceScenario,
    build_reference_scenario,
    capacity_eos_volume_m3,
)
from h2station.tabulated import PropsSI


def _workbook_bytes(workbook: Workbook) -> bytes:
    output = BytesIO()
    workbook.save(output)
    return output.getvalue()


def _overview_archive(path: Path) -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.append([
        "Fill Number", "Lab Test #", "Test #", "P ", "Tank Size",
        "Chamber Temp", "Fill Rate", "Notes",
    ])
    sheet.append([9, 6, "1-1 B", "70 MPa", "4.7 kg", "-20 C", "12 MPa/min", "cold"])
    sheet.append([31, 29, "9-1 A", "70 MPa", "2x 2.3kg", "+20 C", 21.8, "two vessels"])
    with zipfile.ZipFile(path, "w") as bundle:
        bundle.writestr("Test Overview.xlsx", _workbook_bytes(workbook))


def test_h2protocol_reader_handles_source_header_variants_and_active_fill(tmp_path):
    archive = tmp_path / "SAE J2601 Tables Validation Files 1.zip"
    _overview_archive(archive)
    overview = read_h2protocol_overview(archive)
    assert overview["91a"].tank_capacity_kg == pytest.approx(4.6)

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "BMW Test 1-1 B, 2012Oct01"
    sheet.append([
        "Elapsed", "SOC %", "FM", "Pinlet", "Tinlet_G", "Tchamber",
        "Ttank1_G (T10)", "Ttank2_G (T11)", "Ttank3_G (T15)",
    ])
    sheet.append([0.0, 5.0, None, 2.0, -38.0, -20.0, 20.0, 21.0, 22.0])
    sheet.append([0.5, 6.0, 5.0, 3.0, -38.0, -20.0, 21.0, 22.0, 23.0])
    sheet.append([1.0, 7.0, 6.0, 4.0, -37.0, -20.0, 22.0, 23.0, 24.0])
    sheet.append([1.5, 8.0, 7.0, 5.0, -36.0, -20.0, 23.0, 24.0, 25.0])
    sheet.append([2.0, 9.0, None, 6.0, -36.0, -20.0, 24.0, 25.0, 26.0])
    sheet.append([None, None, None, "formatted trailing row", None, None, None, None, None])

    trace = read_h2protocol_workbook(
        _workbook_bytes(workbook),
        source_archive=archive.name,
        source_member="test.xlsx",
        overview=overview,
    )
    assert trace.case_id == "H2P-L06"
    assert trace.metadata.lab_test_number == 6
    assert trace.metadata.tank_capacity_kg == pytest.approx(4.7)
    assert trace.pressure_source == "Pinlet"
    assert trace.time_s.tolist() == pytest.approx([0.0, 0.5, 1.0])
    assert trace.tank_temperature_mean_c.tolist() == pytest.approx([22.0, 23.0, 24.0])


def test_compare_traces_uses_experimental_clock_without_time_warping():
    agreement = compare_traces(
        experimental_time_s=[0, 1, 2],
        experimental_pressure_mpa=[2, 4, 6],
        experimental_temperature_c=[20, 22, 24],
        experimental_soc_percent=[5, 10, 15],
        predicted_time_s=[0, 2],
        predicted_pressure_mpa=[2, 8],
        predicted_temperature_c=[20, 26],
        predicted_soc_percent=[5, 20],
    )
    assert agreement.pressure_mae_mpa == pytest.approx(1.0)
    assert agreement.pressure_rmse_mpa == pytest.approx(np.sqrt(5 / 3))
    assert agreement.temperature_mae_c == pytest.approx(1.0)
    assert agreement.soc_final_error_percentage_points == pytest.approx(5.0)


def test_active_fill_bounds_rejects_negligible_leading_meter_pulse():
    time_s = np.arange(0.0, 43.0)
    flow_g_s = np.zeros_like(time_s)
    flow_g_s[0] = 2.0
    flow_g_s[30:43] = 10.0

    first, last = _active_fill_bounds(
        time_s, flow_g_s, flow_threshold_g_s=1.0
    )

    assert first == 30
    assert last == 42


def test_mc_default_reader_preserves_protocol_schedule_and_source_pressures():
    workbook = Workbook()
    measurement = workbook.active
    measurement.title = "MC Validation Test 1-A"
    measurement.append([
        "Time (s)", "SOC (%)", "875PT1", "875PT3", "Pinlet", "FM",
        "Tinlet_G", "Ttank1_G", "Ttank2_G", "Ttank3_G", "Tchamber",
    ])
    for index, flow in enumerate((0.0, 2.0, 3.0, 4.0, 0.0)):
        measurement.append([
            float(index), 5.0 + index, 88.0 - index, 89.0 - index,
            2.0 + index, flow, -35.0, 20.0 + index, 21.0 + index,
            22.0 + index, 40.0,
        ])
    schedule = workbook.create_sheet("Sheet1")
    schedule.append(["Time", "Pressure"])
    schedule.append([0.0, 2.0])
    schedule.append([300.0, 82.0])

    trace = read_mc_default_workbook(
        _workbook_bytes(workbook),
        source_archive="mc.zip",
        source_member="MC Default Validation Test 1-A, 2013Aug01.xlsx",
    )
    assert trace.case_id == "H2P-MC-1-A"
    assert trace.metadata.tank_capacity_kg == pytest.approx(4.7)
    assert trace.time_s.tolist() == pytest.approx([0.0, 1.0, 2.0])
    assert trace.source_pressure_1_mpa.tolist() == pytest.approx([87.0, 86.0, 85.0])
    assert trace.protocol_pressure_mpa.tolist() == pytest.approx([2.0, 82.0])
    assert trace.summary()["protocol_effective_aprr_mpa_min"] == pytest.approx(16.0)


def test_mc_default_iterator_requires_every_frozen_member(tmp_path):
    archive = tmp_path / "mc.zip"
    with zipfile.ZipFile(archive, "w") as bundle:
        bundle.writestr("unrelated.txt", "data")
    with pytest.raises(ValueError, match="archive is missing"):
        list(iter_mc_default_traces(archive, ["frozen.xlsx"]))


def test_hiad_reader_joins_sheets_and_filters_core_hrs_cases(tmp_path):
    workbook = Workbook()
    workbook.remove(workbook.active)

    def add_sheet(name, headers, rows):
        sheet = workbook.create_sheet(name)
        sheet.append(headers)
        for row in rows:
            sheet.append(row)

    add_sheet("EVENTS", [
        "Event ID", "Q", "Event Title", "Event full description",
        "Event Initiating system", "Classification of the physical effects",
        "Nature of the consequences", "Summary root causes",
    ], [
        [2, "Q2", "HRS leak", "Hose released hydrogen", "Dispenser", "Leak", "Near miss", "Seal"],
        [10, "Q3", "Other", "Unrelated process", "Pipe", "Leak", "None", "Joint"],
    ])
    add_sheet("FACILITY", [
        "Event ID", "Application", "Sub-application", "Hydrogen supply chain stage",
        "Operational condition", "Event Title", "Event full description",
    ], [
        [2, "Hydrogen refuelling station", "Road vehicle", "Distribution", "Fuelling", "", ""],
        [10, "Chemical plant", "Process", "Production", "Operating", "", ""],
    ])
    add_sheet("LESSONS LEARNT", ["Event ID", "Lesson Learnt", "Corrective Measures"], [
        [2, "Verify coupling", "Replace seal"], [10, "Other", "Other"],
    ])
    add_sheet("EVENT NATURE", ["Event ID", "Emergency action"], [
        [2, "Stop fuelling and isolate"], [10, "Stop"],
    ])
    add_sheet("REFERENCES", ["Event ID", "1st Reference & weblink"], [
        [2, "https://example.test/2"], [10, "https://example.test/10"],
    ])
    add_sheet("CONSEQUENCES", ["Event ID"], [[2], [10]])
    path = tmp_path / "HIAD 2.2.xlsx"
    workbook.save(path)

    cases = read_hiad_hrs_cases(path)
    assert [case.event_id for case in cases] == ["2"]
    assert cases[0].emergency_action == "Stop fuelling and isolate"
    assert cases[0].lesson_learnt == "Verify coupling"
    assert cases[0].corrective_measures == "Replace seal"
    assert cases[0].references == ("https://example.test/2",)


def test_reference_scenario_supports_independent_experimental_tank_volumes():
    built = build_reference_scenario(
        ReferenceScenario(
            vehicle_internal_volume_m3=0.060,
            vehicle_2_internal_volume_m3=0.180,
            vehicle_effective_volume_multiplier=1.05,
            vehicle_gas_liner_ua_multiplier=2.5,
        ),
        UnavailableHyRAMBackend(),
    )
    first = built.station.partial_station.vehicle_tank
    second = built.station.secondary_partial_station.vehicle_tank
    assert first.parameters.internal_volume_m3 == pytest.approx(0.060)
    assert second.parameters.internal_volume_m3 == pytest.approx(0.180)
    assert first.effective_volume_m3 == pytest.approx(0.060 * 1.05)
    assert first.fit.gas_liner_ua_multiplier == pytest.approx(2.5)
    assert second.fit == first.fit
    assert (
        built.initial_state.secondary_partial_station.vehicle.hydrogen_mass_kg
        > built.initial_state.partial_station.vehicle.hydrogen_mass_kg
    )


def test_capacity_eos_geometry_uses_declared_public_capacity_for_both_vehicles():
    built = build_reference_scenario(
        ReferenceScenario(
            vehicle_geometry_basis="capacity_eos",
            vehicle_capacity_kg=9.8,
            vehicle_2_capacity_kg=5.0,
        ),
        UnavailableHyRAMBackend(),
    )
    first = built.station.partial_station.vehicle_tank
    second = built.station.secondary_partial_station.vehicle_tank
    assert first.parameters.internal_volume_m3 == pytest.approx(
        capacity_eos_volume_m3(9.8, 70.0e6), rel=1e-10
    )
    assert second.parameters.internal_volume_m3 == pytest.approx(
        capacity_eos_volume_m3(5.0, 70.0e6), rel=1e-10
    )
    assert first.parameters.internal_volume_m3 != pytest.approx(0.122)


def test_capacity_eos_geometry_requires_both_declared_capacities():
    with pytest.raises(ValueError, match="both vehicle capacities"):
        build_reference_scenario(
            ReferenceScenario(
                vehicle_geometry_basis="capacity_eos",
                vehicle_capacity_kg=9.8,
            ),
            UnavailableHyRAMBackend(),
        )


def test_api_rejects_capacity_eos_without_both_vehicle_capacities():
    with pytest.raises(ValueError, match="vehicle_capacity_kg"):
        SimulationInput(
            vehicle_geometry_basis="capacity_eos",
            vehicle_capacity_kg=9.8,
        )


def test_reference_scenario_rejects_nonpositive_experimental_tank_volume():
    with pytest.raises(ValueError, match="tank volumes"):
        build_reference_scenario(
            ReferenceScenario(vehicle_internal_volume_m3=0.0),
            UnavailableHyRAMBackend(),
        )


def test_reference_scenario_rejects_nonpositive_tank_fit_multiplier():
    with pytest.raises(ValueError, match="fit multipliers"):
        build_reference_scenario(
            ReferenceScenario(vehicle_gas_liner_ua_multiplier=0.0),
            UnavailableHyRAMBackend(),
        )


def test_reference_scenario_applies_one_global_precooler_duty_multiplier():
    built = build_reference_scenario(
        ReferenceScenario(precooler_duty_multiplier=3.0),
        UnavailableHyRAMBackend(),
    )
    first = built.station.partial_station.fit
    second = built.station.secondary_partial_station.fit
    assert first.precooler_ua_multiplier == pytest.approx(3.0)
    assert first.chiller_ua_multiplier == pytest.approx(3.0)
    assert first.precooler_capacity_multiplier == pytest.approx(1.0)
    assert second == first


def test_reference_scenario_rejects_nonpositive_precooler_duty_multiplier():
    with pytest.raises(ValueError, match="fit multipliers"):
        build_reference_scenario(
            ReferenceScenario(precooler_duty_multiplier=0.0),
            UnavailableHyRAMBackend(),
        )


def test_reference_scenario_uses_vehicle_specific_nominal_pressure_for_soc():
    built = build_reference_scenario(
        ReferenceScenario(vehicle_nominal_working_pressure_pa=35.0e6),
        UnavailableHyRAMBackend(),
    )
    schedule = built.station.partial_station.controller.schedule
    reference_density = built.station.partial_station.controller.soc_model.reference_density_kg_m3
    expected_density = float(PropsSI(
        "Dmass", "P", 35.0e6, "T", schedule.reference_temperature_k,
        "Hydrogen",
    ))
    assert schedule.nominal_working_pressure_pa == pytest.approx(35.0e6)
    assert reference_density == pytest.approx(expected_density)


def test_dispersion_reader_uses_declared_steady_interval_and_clips_sensor_noise(tmp_path):
    (tmp_path / "ReadMe.txt").write_text(
        "23_FFI_P101_T00001 - test-01-0\n", encoding="utf-8"
    )
    csv_text = (
        ",mass flow meter 1,filling pressure sensor,various data,"
        "h2 sensor time,sensor 1 h2 concentration [%],sensor 2 h2 concentration [%]\n"
        "0,species - h2,mean pressure [bar] - 2.5,total,0,-0.1,-0.2\n"
        "1,mean mass flow - 0.5,std,mean channel temperature [C] - 25,10,1,2\n"
        "2,mass fraction - 1,,baseline duration [s] - 30,50,3,5\n"
        "3,logging frequency,,filling duration [s] - 30,55,4,6\n"
        "4,,,,60,5,7\n"
    )
    archive = tmp_path / "23_FFI_P101_T00001.zip"
    with zipfile.ZipFile(archive, "w") as bundle:
        bundle.writestr("case/case.csv", csv_text)
    experiment = list(iter_dispersion_experiments(tmp_path))[0]
    summary = experiment.summary()
    assert experiment.article_test_id == "test-01-0"
    assert experiment.mean_mass_flow_g_s == pytest.approx(0.5)
    assert experiment.mean_filling_pressure_bar == pytest.approx(2.5)
    assert summary["steady_interval_start_s"] == pytest.approx(50.0)
    assert summary["steady_mean_concentration_percent"] == pytest.approx(5.0)
    assert summary["steady_peak_concentration_percent"] == pytest.approx(7.0)
    assert summary["steady_lfl_exceedance_fraction"] == pytest.approx(5 / 6)


def test_detector_threshold_requires_declared_persistence_and_resets_on_gap():
    times = np.asarray([0.0, 0.5, 1.0, 1.6, 3.0])
    values = np.asarray([0.0, 1.2, 1.3, 0.0, 1.5])
    assert first_persistent_threshold_time(times, values, 1.0, 0.5) == pytest.approx(1.0)
    assert first_persistent_threshold_time(times, values, 1.0, 1.0) is None


def test_dispersion_detector_logic_reports_sensor_coverage_and_latency():
    experiment = DispersionExperiment(
        case_id="case-1",
        article_test_id="test-1",
        source_archive="case-1.zip",
        mean_mass_flow_g_s=0.5,
        mean_filling_pressure_bar=1.0,
        channel_temperature_c=25.0,
        baseline_duration_s=1.0,
        filling_duration_s=3.0,
        sensor_ids=("1", "2"),
        sensor_time_s=np.asarray([0.0, 1.0, 1.5, 2.0, 3.0]),
        concentrations_percent=np.asarray([
            [0.0, 0.0],
            [0.0, 1.2],
            [2.2, 2.2],
            [2.2, 2.2],
            [0.0, 0.0],
        ]),
    )
    result = evaluate_dispersion_detector_logic(
        experiment,
        alarm_threshold_percent=1.0,
        trip_threshold_percent=2.0,
        persistence_s=0.5,
    )
    assert result["alarm"]["detected_sensor_count"] == 2
    assert result["trip"]["detected_sensor_count"] == 2
    assert result["alarm"]["first_detection_after_fill_start_s"] == pytest.approx(0.5)
    assert result["trip"]["first_detection_after_fill_start_s"] == pytest.approx(1.0)


def test_frozen_dispersion_detector_artifact_matches_the_public_archives():
    """Keep the published detector-logic evidence tied to the checked-in bytes.

    This is deliberately a bounded integrity check: it verifies that the
    aggregate reported in the frozen artifact is reproducible from the public
    archives, but it does not promote the open-channel experiment to a full
    hydrogen-refuelling-station validation set.
    """
    root = Path(__file__).resolve().parents[1]
    artifact = json.loads(
        (root / "research/dispersion_detector_logic_validation.json").read_text(
            encoding="utf-8"
        )
    )
    raw = root / "data/public_validation/raw/hydrogen_dispersion_channel"
    archives = sorted(raw.glob("*.zip"))
    manifest = artifact["source"]["archive_manifest"]

    assert artifact["status"] == "completed_bounded_instrumented_detector_logic_evidence"
    assert artifact["evidence_role"] == "instrumented_detector_logic_evidence"
    assert artifact["aggregate"] == {
        "case_count": 22,
        "sensor_count_per_case": [29],
        "cases_with_alarm_detection": 22,
        "cases_with_trip_detection": 22,
        "mean_alarm_sensor_coverage_fraction": pytest.approx(0.8871473354231975),
        "mean_trip_sensor_coverage_fraction": pytest.approx(0.8087774294670846),
        "median_first_alarm_after_fill_start_s": pytest.approx(10.674999999995634),
        "median_first_trip_after_fill_start_s": pytest.approx(10.674999999995634),
    }
    assert len(archives) == len(manifest) == 22

    actual = []
    for archive in archives:
        digest = hashlib.sha256(archive.read_bytes()).hexdigest()
        actual.append({
            "name": archive.name,
            "bytes": archive.stat().st_size,
            "sha256": digest,
        })
    assert actual == manifest
