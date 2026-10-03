from __future__ import annotations

from pathlib import Path
import sys

import pytest
from openpyxl import Workbook

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "scripts"))

from run_nrel_h2fills_hdvs_validation import TANK_IDS, read_nrel_workbook  # noqa: E402


def _make_workbook(path: Path) -> None:
    workbook = Workbook()
    description = workbook.active
    description.title = "Description"
    description.append(["Test conditions", "Name", "Value"])
    description.append([None, "CHSS size", "68.6 kg (7 tanks at 9.8 kg each)"])
    data = workbook.create_sheet("Data")
    headers = ["Time [s]"]
    for tank_id in TANK_IDS:
        headers.extend(
            [
                f"HDVS_ tank#{tank_id}_inlet_press [MPa]",
                f"HDVS_ tank#{tank_id}_inlet_temp [degC]",
                f"HDVS_ tank#{tank_id}_mass [kg]",
                f"HDVS_ tank#{tank_id}_internal_press [MPa]",
                f"HDVS_ tank#{tank_id}_internal_temp [degC]",
            ]
        )
    data.append(headers)
    for second in range(3):
        row = [second]
        for tank_id in TANK_IDS:
            row.extend(
                [
                    1.0 + second,
                    15.0,
                    0.4 + 0.1 * second,
                    1.8 + 1.0 * second,
                    15.0 + 2.0 * second,
                ]
            )
        data.append(row)
    workbook.create_sheet("Charts")
    workbook.save(path)


def test_nrel_reader_preserves_seven_tank_common_clock(tmp_path):
    path = tmp_path / "sample.xlsx"
    _make_workbook(path)

    dataset = read_nrel_workbook(path)

    assert dataset.sample_count == 3
    assert dataset.summary()["tank_count"] == 7
    assert [trace.tank_id for trace in dataset.traces] == list(TANK_IDS)
    assert dataset.traces[0].time_s.tolist() == [0.0, 1.0, 2.0]
    assert dataset.traces[0].mass_kg.tolist() == pytest.approx([0.4, 0.5, 0.6])


def test_nrel_reader_rejects_missing_required_channel(tmp_path):
    path = tmp_path / "invalid.xlsx"
    _make_workbook(path)
    workbook = __import__("openpyxl").load_workbook(path)
    sheet = workbook["Data"]
    sheet.cell(row=1, column=2).value = "wrong"
    workbook.save(path)

    try:
        read_nrel_workbook(path)
    except ValueError as exc:
        assert "missing columns" in str(exc)
    else:
        raise AssertionError("reader accepted a workbook with a missing channel")
