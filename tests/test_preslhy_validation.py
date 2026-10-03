from __future__ import annotations

from io import BytesIO

import numpy as np
from openpyxl import Workbook
import pytest

from h2station.preslhy_validation import (
    eligible_pressure_window,
    evaluate_preslhy_trace,
    read_preslhy_workbook,
    simulate_preslhy_blowdown,
)


def _workbook_bytes() -> bytes:
    workbook = Workbook()
    pressure = workbook.active
    pressure.title = "20190101_010101-Press"
    for _ in range(6):
        pressure.append([])
    pressure.append(["Synchronized Time [s]", "PVes [bar_g]", "PNoz [bar_g]"])
    for time_s in np.linspace(-0.5, 5.0, 111):
        pressure.append([float(time_s), max(0.0, 100.0 - 20.0 * max(time_s, 0.0)), 1.0])
    temperature = workbook.create_sheet("20190101_010101-Temp")
    for _ in range(6):
        temperature.append([])
    temperature.append(["Synchronized Time [s]", "T1 [degC]", "T2 [degC]", "T3 [degC]"])
    for time_s in np.linspace(-0.5, 1.0, 31):
        temperature.append([float(time_s), 20.0, 21.0, 22.0])
    ambient = workbook.create_sheet("20190101_010101-cH2-Amb")
    ambient.append(["Synchronized Time [s]", "Ambient pressure [hPa]"])
    for time_s in np.linspace(-5.0, 1.0, 13):
        ambient.append([float(time_s), 1005.0])
    output = BytesIO()
    workbook.save(output)
    return output.getvalue()


def test_preslhy_reader_uses_synchronised_pves_and_internal_temperature():
    trace = read_preslhy_workbook(
        _workbook_bytes(),
        source_package="PRE3P1A_KIT_D1_300K_DATA.zip",
        source_member="20190101_010101.xlsx",
        nozzle_diameter_mm=1.0,
    )
    assert trace.case_id == "20190101_010101"
    assert trace.initial_temperature_k == pytest.approx(294.15)
    assert trace.temperature_substituted is False
    assert trace.ambient_pressure_pa == pytest.approx(100_500.0)
    assert trace.ambient_pressure_substituted is False
    assert trace.pressure_unit_interpretation == "header_gauge_plus_standard_ambient"
    assert trace.initial_pressure_pa == pytest.approx(101.01325e5)
    time_s, pressure = eligible_pressure_window(trace)
    assert time_s[0] >= 0.1
    assert pressure[-1] == pytest.approx(1.01325)


def test_frozen_blowdown_model_closes_mass_and_reduces_pressure():
    trace = read_preslhy_workbook(
        _workbook_bytes(),
        source_package="PRE3P1A_KIT_D1_300K_DATA.zip",
        source_member="20190101_010101.xlsx",
        nozzle_diameter_mm=1.0,
    )
    times = np.linspace(0.1, 2.0, 20)
    pressure, mass, peak = simulate_preslhy_blowdown(trace, times)
    assert np.all(np.diff(pressure) <= 1.0e-3)
    assert np.all(np.diff(mass) <= 1.0e-12)
    assert pressure[-1] < pressure[0]
    assert mass[-1] < mass[0]
    assert peak[0] > 0.0


def test_case_evaluator_reports_frozen_primary_endpoints():
    trace = read_preslhy_workbook(
        _workbook_bytes(),
        source_package="PRE3P1A_KIT_D1_300K_DATA.zip",
        source_member="20190101_010101.xlsx",
        nozzle_diameter_mm=1.0,
    )
    result = evaluate_preslhy_trace(trace)
    assert result.samples >= 20
    assert result.pressure_nrmse_percent_initial_absolute_pressure >= 0.0
    assert result.pressure_screen_pass == (
        result.pressure_nrmse_percent_initial_absolute_pressure <= 10.0
    )
    assert result.half_time_screen_pass == (
        result.time_to_50_percent_gauge_relative_error_percent <= 20.0
    )
