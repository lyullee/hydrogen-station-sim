from __future__ import annotations

import json
from pathlib import Path
from zipfile import ZipFile

from openpyxl import Workbook

from scripts.audit_controlled_data_schema import inventory_schema


def test_header_only_inventory_identifies_candidate_shapes_without_disclosure(tmp_path: Path):
    full_loop = tmp_path / "owner_site_private_trace.csv"
    full_loop.write_text(
        "private_time,private_vehicle_pressure,private_inlet_temp,private_mass_flow,"
        "private_vehicle_id,private_cascade_pressure,private_controller_state\n"
        "0,1,2,3,4,5,6\n"
        "1,1,2,3,4,5,6\n"
        "2,1,2,3,4,5,6\n",
        encoding="utf-8",
    )
    vehicle = tmp_path / "sensitive_vehicle_run.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "confidential worksheet"
    sheet.append([
        "private_time", "private_vehicle_pressure", "private_temperature",
        "private_mass_flow", "private_vehicle", "private_dispenser",
    ])
    sheet.append([0, 1, 2, 3, 4, 5])
    sheet.append([1, 1, 2, 3, 4, 5])
    workbook.save(vehicle)

    report = inventory_schema([tmp_path])
    rendered = json.dumps(report, ensure_ascii=False)

    assert report["schema_search_max_rows_per_table"] == 40
    assert report["schema_search_rows_examined"] == 7
    assert report["raw_rows_persisted"] is False
    assert report["candidate_schema_counts"]["full_loop_candidate"] == 1
    assert report["candidate_schema_counts"]["vehicle_fill_candidate"] == 1
    assert report["candidate_schema_counts"]["station_recharge_candidate"] == 0
    assert "owner_site_private_trace" not in rendered
    assert "private_vehicle_pressure" not in rendered
    assert "confidential worksheet" not in rendered
    assert "sensitive-value" not in rendered


def test_inventory_finds_header_after_title_rows_without_retaining_them(tmp_path: Path):
    source = tmp_path / "controlled_export.csv"
    source.write_text(
        "confidential title\n"
        "operator-only note\n"
        "clock,vehicle_pressure,inlet_temperature,mass_flow,vehicle,cascade,controller_state\n"
        "0,1,2,3,4,5,6\n"
        "1,1,2,3,4,5,6\n",
        encoding="utf-8",
    )

    report = inventory_schema([tmp_path])
    rendered = json.dumps(report, ensure_ascii=False)

    assert report["candidate_schema_counts"]["full_loop_candidate"] == 1
    assert report["measurement_rows_persisted"] is False
    assert "confidential title" not in rendered
    assert "operator-only note" not in rendered
    assert "private-value" not in rendered


def test_schema_inventory_counts_unreadable_tabular_file_without_disclosure(tmp_path: Path):
    unreadable = tmp_path / "restricted.xlsx"
    unreadable.write_bytes(b"not a workbook")

    report = inventory_schema([tmp_path])
    rendered = json.dumps(report, ensure_ascii=False)

    assert report["source_files_scanned"] == 1
    assert report["source_tables_scanned"] == 0
    assert report["unreadable_tables"] == 1
    assert "restricted.xlsx" not in rendered


def test_schema_inventory_screens_zip_members_without_disclosing_member_names(tmp_path: Path):
    archive = tmp_path / "controlled_bundle.zip"
    with ZipFile(archive, "w") as bundle:
        bundle.writestr(
            "private-folder/secret-full-loop.csv",
            "time,vehicle_pressure,temperature,mass_flow,vehicle,cascade,controller_state\n"
            "0,1,2,3,4,5,6\n"
            "1,1,2,3,4,5,6\n",
        )

    report = inventory_schema([tmp_path])
    rendered = json.dumps(report, ensure_ascii=False)

    assert report["archive_members_scanned"] == 1
    assert report["candidate_schema_counts"]["full_loop_candidate"] == 1
    assert report["tabular_content_format_counts"] == {"csv": 1}
    assert "secret-full-loop" not in rendered
    assert "private-folder" not in rendered
    assert "private-value" not in rendered


def test_schema_inventory_detects_utf16_logger_headers_without_disclosure(tmp_path: Path):
    source = tmp_path / "controlled_utf16_logger.csv"
    source.write_text(
        "time,vehicle_pressure,temperature,mass_flow,vehicle,cascade,controller_state\n"
        "0,1,2,3,4,5,6\n"
        "1,1,2,3,4,5,6\n",
        encoding="utf-16",
    )

    report = inventory_schema([tmp_path])
    rendered = json.dumps(report, ensure_ascii=False)

    assert report["candidate_schema_counts"]["full_loop_candidate"] == 1
    assert report["measurement_like_tables"] == 1
    assert "controlled_utf16_logger" not in rendered
    assert "private-value" not in rendered


def test_schema_inventory_recognizes_us_locale_historian_timestamps(tmp_path: Path):
    source = tmp_path / "private-historian-export.csv"
    source.write_text(
        "Date/Time,vehicle_pressure,temperature,mass_flow,vehicle,cascade,controller_state\n"
        "1/7/2026 1:00:00 PM,1,2,3,4,5,6\n"
        "1/7/2026 1:00:01 PM,1,2,3,4,5,6\n",
        encoding="utf-8",
    )

    report = inventory_schema([tmp_path])

    assert report["measurement_like_tables"] == 1
    assert report["candidate_schema_counts"]["full_loop_candidate"] == 1


def test_schema_inventory_recognizes_reverse_order_historian_timestamps(tmp_path: Path):
    source = tmp_path / "private-reverse-historian-export.csv"
    source.write_text(
        "Date/Time,vehicle_pressure,temperature,mass_flow,vehicle,cascade,controller_state\n"
        "1/7/2026 1:00:01 PM,1,2,3,4,5,6\n"
        "1/7/2026 1:00:00 PM,1,2,3,4,5,6\n",
        encoding="utf-8",
    )

    report = inventory_schema([tmp_path])

    assert report["measurement_like_tables"] == 1
    assert report["candidate_schema_counts"]["full_loop_candidate"] == 1


def test_schema_inventory_recognizes_space_separated_historian_timestamps(tmp_path: Path):
    source = tmp_path / "private-space-separated-historian.csv"
    source.write_text(
        "time,vehicle_pressure,temperature,mass_flow,vehicle,cascade,controller_state\n"
        "2026 01 07 13:00:00,1,2,3,4,5,6\n"
        "2026 01 07 13:00:01,1,2,3,4,5,6\n",
        encoding="cp949",
    )

    report = inventory_schema([tmp_path])

    assert report["measurement_like_tables"] == 1
    assert report["candidate_schema_counts"]["full_loop_candidate"] == 1


def test_schema_inventory_streams_csv_schema_without_whole_file_read(
    tmp_path: Path, monkeypatch,
):
    source = tmp_path / "private-large-logger.csv"
    source.write_text(
        "time,vehicle_pressure,temperature,mass_flow,vehicle,cascade,controller_state\n"
        "0,1,2,3,4,5,6\n"
        "1,1,2,3,4,5,6\n",
        encoding="utf-8",
    )

    def no_full_file_read(_: Path) -> bytes:
        raise AssertionError("whole controlled CSV must not be read")

    monkeypatch.setattr(Path, "read_bytes", no_full_file_read)
    report = inventory_schema([tmp_path])

    assert report["measurement_like_tables"] == 1
    assert report["raw_rows_persisted"] is False


def test_inventory_groups_aligned_flat_files_without_publishing_clock_or_names(
    tmp_path: Path,
):
    timestamps = (
        "2026-01-07T13:00:00", "2026-01-07T13:00:01",
        "2026-01-07T13:00:02", "2026-01-07T13:00:03",
    )
    (tmp_path / "private-pressure.csv").write_text(
        "time,storage_pressure,compressor_state\n"
        + "".join(f"{stamp},50,run\n" for stamp in timestamps),
        encoding="utf-8",
    )
    (tmp_path / "private-flow.csv").write_text(
        "time,mass_flow,vehicle,dispenser_temperature\n"
        + "".join(f"{stamp},1,private-car,20\n" for stamp in timestamps),
        encoding="utf-8",
    )

    report = inventory_schema([tmp_path])
    rendered = json.dumps(report, ensure_ascii=False)

    assert report["schema_version"] == 3
    assert report["flat_time_axis_candidate_summary"] == {
        "candidate_groups": 1,
        "tables_in_candidate_groups": 2,
        "largest_candidate_group_tables": 2,
        "fingerprints_published": False,
        "absolute_time_samples_published": False,
    }
    assert report["candidate_schema_counts"]["synchronized_flat_full_loop_candidate"] == 1
    assert "2026-01-07" not in rendered
    assert "private-pressure" not in rendered
    assert "private-car" not in rendered


def test_inventory_does_not_group_flat_files_with_different_tail_clocks(tmp_path: Path):
    (tmp_path / "first.csv").write_text(
        "time,storage_pressure,compressor_state\n"
        "0,50,run\n1,51,run\n2,52,run\n3,53,run\n",
        encoding="utf-8",
    )
    (tmp_path / "second.csv").write_text(
        "time,mass_flow,vehicle,temperature\n"
        "0,1,car,20\n1,1,car,20\n2,1,car,20\n4,1,car,20\n",
        encoding="utf-8",
    )

    report = inventory_schema([tmp_path])

    assert report["flat_time_axis_candidate_summary"]["candidate_groups"] == 0
    assert report["candidate_schema_counts"]["synchronized_flat_full_loop_candidate"] == 0


def test_schema_inventory_recognizes_compact_instrument_tags_without_tag_disclosure(
    tmp_path: Path,
):
    source = tmp_path / "controlled_tag_logger.csv"
    source.write_text(
        "LocalTimeCol,COMP.AI.PT_201,COMP.AI.TT_201,FQI_0001,COMP.STATUS.RUN\n"
        "0,1,2,3,1\n"
        "1,1,2,3,1\n",
        encoding="utf-8",
    )

    report = inventory_schema([tmp_path])
    rendered = json.dumps(report, ensure_ascii=False)

    assert report["measurement_like_tables"] == 1
    assert report["semantic_channel_table_counts"] == {
        "controller_state": 1,
        "mass_flow": 1,
        "pressure": 1,
        "temperature": 1,
        "time": 1,
    }
    assert "COMP.AI.PT_201" not in rendered
    assert "private-value" not in rendered


def test_inventory_flags_complementary_workbook_tables_only_as_co_located_candidate(
    tmp_path: Path,
):
    source = tmp_path / "controlled_subsystem_export.xlsx"
    workbook = Workbook()
    storage = workbook.active
    storage.title = "private storage signals"
    storage.append(["시간", "저장탱크 압력", "저장용기 온도", "운전 상태"])
    storage.append([0, 1, 2, "private-state"])
    storage.append([1, 1, 2, "private-state"])
    dispenser = workbook.create_sheet("private dispenser signals")
    dispenser.append(["시각", "차량", "충전기 유량"])
    dispenser.append([0, "private-vehicle", 3])
    dispenser.append([1, "private-vehicle", 3])
    workbook.save(source)

    report = inventory_schema([tmp_path])
    rendered = json.dumps(report, ensure_ascii=False)

    assert report["candidate_schema_counts"]["full_loop_candidate"] == 0
    assert report["candidate_schema_counts"]["co_located_full_loop_candidate"] == 1
    assert "private storage signals" not in rendered
    assert "private dispenser signals" not in rendered
    assert "private-time" not in rendered


def test_inventory_rejects_documentation_sheets_that_only_mention_channels(
    tmp_path: Path,
):
    source = tmp_path / "controlled_documentation.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "private notes"
    sheet.append(["시간", "차량 압력", "온도", "유량"])
    sheet.append(["설계 검토 문서이며 계측 로그가 아니다."])
    other = workbook.create_sheet("private response notes")
    other.append(["시간", "저장 압력", "압축기", "제어 상태"])
    other.append(["실제 시계열 또는 동기화된 운전 기록이 아니다."])
    workbook.save(source)

    report = inventory_schema([tmp_path])

    assert report["measurement_like_tables"] == 0
    assert report["candidate_schema_counts"]["co_located_full_loop_candidate"] == 0
    assert report["candidate_schema_counts"]["rejected_nonmeasurement_candidate_container"] == 1
    assert report["sample_data_rows_structurally_inspected_in_memory"] is True


def test_inventory_rejects_hazop_style_rows_without_a_monotonic_logger_clock(
    tmp_path: Path,
):
    """Do not treat case tables as traces just because their headers are familiar."""

    source = tmp_path / "controlled_hazop_reference.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "private case reference"
    sheet.append(["시간", "저장 압력", "온도", "유량", "운전 상태"])
    sheet.append(["N01", 40, 20, 0, "대기"])
    sheet.append(["N02", 70, 25, 0, "대기"])
    sheet.append(["N03", 95, 30, 0, "대기"])
    workbook.save(source)

    report = inventory_schema([tmp_path])

    assert report["measurement_like_tables"] == 0
    assert report["candidate_schema_counts"]["full_loop_candidate"] == 0


def test_inventory_rejects_repeated_timestamp_values_even_when_rows_are_numeric(
    tmp_path: Path,
):
    source = tmp_path / "controlled_static_table.csv"
    source.write_text(
        "time,vehicle_pressure,temperature,mass_flow,vehicle,cascade,controller_state\n"
        "0,1,2,3,4,5,6\n"
        "0,1,2,3,4,5,6\n"
        "0,1,2,3,4,5,6\n",
        encoding="utf-8",
    )

    report = inventory_schema([tmp_path])

    assert report["measurement_like_tables"] == 0
    assert report["candidate_schema_counts"]["full_loop_candidate"] == 0
