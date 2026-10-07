"""Audit and replay the public MetHyTrucks Hy-SaM measurement workbooks.

The Zenodo files were numerically inspected before this diagnostic was
specified.  Results produced here are therefore post-access external
diagnostics, not a prospective holdout.  Raw workbook rows remain in the
gitignored data tree; only provenance, aggregate channel checks and model
errors are written to the research artifact.
"""

from __future__ import annotations

import argparse
from dataclasses import dataclass
from datetime import datetime, timezone
import hashlib
import json
from pathlib import Path
import subprocess
from typing import Any

import numpy as np
from openpyxl import load_workbook
from scipy.integrate import solve_ivp

from h2station.public_tank_calibration import load_public_type_iv_tank_calibration
from h2station.scenario import build_vehicle_tank
from h2station.tabulated import PropsSI
from h2station.vehicle import CompositeTankFitParameters, CompositeTankState, TankBoundaryFlow


FILES = {
    "20241023_Test_1_HySaM.xlsx": {
        "sha256": "fbdd81a7ec9c1191d69322f06087e8b30a33acc8b340b861a636173c08d9913a",
        "expected_rows": 6600,
    },
    "20241023_Test_4_HySaM_CESAME.xlsx": {
        "sha256": "bc9e5cbd23350b7320a96d627854cc5154f1b92ff987e087c957b1da691ea4aa",
        "expected_rows": 720,
    },
    "20241024_Test_9_HySaM.xlsx": {
        "sha256": "f33d3e451b42b1d3aff905cc5deb1e9a2ed11eeaf7d4f0d559446a609bdc7563",
        "expected_rows": 2160,
    },
}

COMMON_COLUMNS = (
    "Zeit",
    "TT_D04 ValueY",
    "PTX05 ValueY",
    "PT03 ValueY",
    "TEX01 ValueY",
    "TT24 ValueY",
    "QT_D02 ValueY",
    "PTD10 ValueY",
    "PTD11 ValueY",
    "FWg35_Masse ValueY",
)
TANK_CANDIDATE_COLUMNS = (
    "PT01 ValueY",
    "TT08 ValueY",
    "TT09 ValueY",
    "TT10 ValueY",
    "TT11 ValueY",
)
MASS_CLOSURE_RATIO_MIN = 0.8
MASS_CLOSURE_RATIO_MAX = 1.2
MINIMUM_SESSION_MASS_KG = 0.05
MINIMUM_SESSION_DURATION_S = 10.0


@dataclass(frozen=True)
class WorkbookTrace:
    filename: str
    elapsed_s: np.ndarray
    channels: dict[str, np.ndarray]


def _sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def _git_commit() -> str:
    try:
        return subprocess.run(
            ["git", "rev-parse", "HEAD"], check=True, capture_output=True, text=True,
        ).stdout.strip()
    except (OSError, subprocess.CalledProcessError):
        return "unavailable"


def _clock_seconds(value: Any) -> float:
    return float(
        value.hour * 3600
        + value.minute * 60
        + value.second
        + value.microsecond / 1.0e6
    )


def load_trace(path: Path) -> WorkbookTrace:
    expected = FILES[path.name]
    observed_hash = _sha256(path)
    if observed_hash != expected["sha256"]:
        raise ValueError(f"SHA-256 mismatch for {path.name}: {observed_hash}")
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        sheet = workbook[workbook.sheetnames[0]]
        rows = sheet.iter_rows(values_only=True)
        headers = list(next(rows))
        missing = [column for column in COMMON_COLUMNS if column not in headers]
        if missing:
            raise ValueError(f"{path.name} is missing columns: {missing}")
        records = list(rows)
    finally:
        workbook.close()
    if len(records) != expected["expected_rows"]:
        raise ValueError(
            f"{path.name} row count changed: {len(records)} != {expected['expected_rows']}"
        )
    clock = np.asarray([_clock_seconds(row[headers.index("Zeit")]) for row in records])
    elapsed = clock - clock[0]
    channels = {
        str(header): np.asarray([row[index] for row in records], dtype=float)
        for index, header in enumerate(headers)
        if header is not None and all(isinstance(row[index], (int, float)) for row in records)
    }
    return WorkbookTrace(path.name, elapsed, channels)


def active_sessions(
    elapsed_s: np.ndarray,
    flow_g_s: np.ndarray,
    *,
    threshold_g_s: float = 1.0,
    join_gap_s: float = 180.0,
) -> list[tuple[int, int]]:
    """Group active flow pulses into sessions without smoothing or time warping."""

    indices = np.flatnonzero(flow_g_s > threshold_g_s)
    if len(indices) == 0:
        return []
    runs = np.split(indices, np.where(np.diff(indices) > 1)[0] + 1)
    sessions: list[list[int]] = [[int(runs[0][0]), int(runs[0][-1])]]
    for run in runs[1:]:
        start, end = int(run[0]), int(run[-1])
        if elapsed_s[start] - elapsed_s[sessions[-1][1]] <= join_gap_s:
            sessions[-1][1] = end
        else:
            sessions.append([start, end])
    return [(start, end) for start, end in sessions]


def _session_summary(trace: WorkbookTrace, start: int, end: int) -> dict[str, Any]:
    sl = slice(start, end + 1)
    time_s = trace.elapsed_s[sl] - trace.elapsed_s[start]
    raw_flow = trace.channels["QT_D02 ValueY"][sl]
    baseline = float(np.median(np.partition(raw_flow, min(20, len(raw_flow) - 1))[:20]))
    flow = np.maximum(raw_flow - baseline, 0.0)
    integrated_mass = float(np.trapezoid(flow / 1000.0, time_s))
    cumulative = trace.channels["FWg35_Masse ValueY"][sl]
    scale_delta = float(cumulative[-1] - cumulative[0])
    result: dict[str, Any] = {
        "start_elapsed_s": float(trace.elapsed_s[start]),
        "end_elapsed_s": float(trace.elapsed_s[end]),
        "duration_s": float(time_s[-1]),
        "sample_count": int(len(time_s)),
        "flow_baseline_g_s": baseline,
        "peak_flow_g_s": float(np.max(flow)),
        "active_mean_flow_g_s": float(np.mean(flow[flow > 1.0])),
        "integrated_flow_mass_kg": integrated_mass,
        "cumulative_scale_delta_kg": scale_delta,
        "flow_to_scale_mass_ratio": (
            integrated_mass / scale_delta if scale_delta > 0.0 else None
        ),
    }
    for tag in ("PT01 ValueY", "PTD10 ValueY", "PTD11 ValueY"):
        if tag in trace.channels:
            values = trace.channels[tag][sl]
            result.setdefault("pressure_candidates_bar", {})[tag] = {
                "initial": float(values[0]),
                "final": float(values[-1]),
                "minimum": float(np.min(values)),
                "maximum": float(np.max(values)),
            }
    if all(tag in trace.channels for tag in TANK_CANDIDATE_COLUMNS[1:]):
        tank_temperature = np.mean(
            np.stack([trace.channels[tag][sl] for tag in TANK_CANDIDATE_COLUMNS[1:]]),
            axis=0,
        )
        result["four_sensor_temperature_mean_c"] = {
            "initial": float(tank_temperature[0]),
            "final": float(tank_temperature[-1]),
            "maximum": float(np.max(tank_temperature)),
        }
    return result


def _candidate_session_eligible(summary: dict[str, Any]) -> bool:
    """Apply measurement-only eligibility before inspecting model errors."""

    ratio = summary.get("flow_to_scale_mass_ratio")
    return bool(
        isinstance(ratio, (int, float))
        and MASS_CLOSURE_RATIO_MIN <= ratio <= MASS_CLOSURE_RATIO_MAX
        and summary.get("integrated_flow_mass_kg", 0.0) >= MINIMUM_SESSION_MASS_KG
        and summary.get("duration_s", 0.0) >= MINIMUM_SESSION_DURATION_S
    )


def _replay_candidate_session(
    trace: WorkbookTrace,
    start: int,
    end: int,
    *,
    session_selection: str,
) -> dict[str, Any]:
    """Replay one candidate tank mapping without fitting parameters."""

    required = set(TANK_CANDIDATE_COLUMNS) | {
        "QT_D02 ValueY", "TEX01 ValueY", "PTD10 ValueY", "TT_D04 ValueY"
    }
    missing = sorted(required - set(trace.channels))
    if missing:
        raise ValueError(f"candidate replay is missing {missing}")
    sl = slice(start, end + 1)
    time_s = trace.elapsed_s[sl] - trace.elapsed_s[start]
    flow_raw = trace.channels["QT_D02 ValueY"][sl]
    baseline = float(np.median(np.partition(flow_raw, min(20, len(flow_raw) - 1))[:20]))
    flow_kg_s = np.maximum(flow_raw - baseline, 0.0) / 1000.0
    pressure_mpa = trace.channels["PT01 ValueY"][sl] / 10.0
    source_pressure_mpa = trace.channels["PTD10 ValueY"][sl] / 10.0
    inlet_temperature_c = trace.channels["TEX01 ValueY"][sl]
    tank_temperature_c = np.mean(
        np.stack([trace.channels[tag][sl] for tag in TANK_CANDIDATE_COLUMNS[1:]]),
        axis=0,
    )
    ambient_k = float(np.median(trace.channels["TT_D04 ValueY"][sl]) + 273.15)
    calibration = load_public_type_iv_tank_calibration()
    if calibration is None:
        raise RuntimeError("frozen public Type-IV tank calibration is unavailable")
    fit = CompositeTankFitParameters(
        effective_volume_multiplier=calibration.effective_volume_multiplier,
        gas_liner_ua_multiplier=calibration.gas_liner_ua_multiplier,
    )
    # The paper identifies a 244 L Type-IV sink for Hy-SaM in set-up 1.  The
    # workbook does not contain a test-to-device dictionary, so this geometry
    # remains an explicit candidate mapping rather than a confirmed identity.
    tank = build_vehicle_tank(0.244, fit)
    initial = tank.initial_state(
        float(pressure_mpa[0] * 1.0e6), float(tank_temperature_c[0] + 273.15)
    )

    def rhs(local_time_s: float, values: np.ndarray) -> np.ndarray:
        state = CompositeTankState.from_vector(values)
        mass_flow = float(np.interp(local_time_s, time_s, flow_kg_s))
        inlet_temperature_k = float(
            np.interp(local_time_s, time_s, inlet_temperature_c) + 273.15
        )
        inlet_pressure_pa = float(
            max(0.2, np.interp(local_time_s, time_s, source_pressure_mpa)) * 1.0e6
        )
        inlet_enthalpy = float(PropsSI(
            "Hmass", "P", inlet_pressure_pa, "T", inlet_temperature_k, "Hydrogen"
        ))
        return tank.derivative(state, TankBoundaryFlow(
            inlet_mass_flow_kg_s=mass_flow,
            inlet_specific_enthalpy_j_kg=inlet_enthalpy,
            ambient_temperature_k=ambient_k,
        )).as_vector()

    solution = solve_ivp(
        rhs,
        (float(time_s[0]), float(time_s[-1])),
        initial.as_vector(),
        method="BDF",
        t_eval=time_s,
        rtol=5.0e-6,
        atol=5.0e-8,
    )
    if not solution.success:
        raise RuntimeError(solution.message)
    gas = [tank.gas_state(CompositeTankState.from_vector(row)) for row in solution.y.T]
    predicted_pressure = np.asarray([item.pressure_pa / 1.0e6 for item in gas])
    predicted_temperature = np.asarray([item.temperature_k - 273.15 for item in gas])
    pressure_error = predicted_pressure - pressure_mpa
    temperature_error = predicted_temperature - tank_temperature_c
    cumulative = trace.channels["FWg35_Masse ValueY"][sl]
    integrated_mass = float(np.trapezoid(flow_kg_s, time_s))
    scale_delta = float(cumulative[-1] - cumulative[0])
    return {
        "workbook": trace.filename,
        "session_selection": session_selection,
        "event_elapsed_start_s": float(trace.elapsed_s[start]),
        "event_elapsed_end_s": float(trace.elapsed_s[end]),
        "event_duration_s": float(time_s[-1]),
        "sample_count": int(len(time_s)),
        "candidate_mapping": {
            "tank_pressure": "PT01 ValueY interpreted as bar absolute",
            "tank_temperature": "arithmetic mean of TT08..TT11 interpreted as degC",
            "mass_flow": "QT_D02 ValueY interpreted as g/s after measured idle baseline",
            "inlet_temperature": "TEX01 ValueY interpreted as degC",
            "upstream_pressure": "PTD10 ValueY interpreted as bar absolute",
            "tank_internal_volume_m3": 0.244,
            "mapping_status": "inferred_from_article_and_signal_dynamics_not_confirmed_by_channel_dictionary",
        },
        "fit": {
            "source": calibration.evidence_artifact,
            "effective_volume_multiplier": calibration.effective_volume_multiplier,
            "gas_liner_ua_multiplier": calibration.gas_liner_ua_multiplier,
            "case_specific_fitting": False,
        },
        "mass_boundary": {
            "integrated_flow_kg": integrated_mass,
            "cumulative_scale_delta_kg": scale_delta,
            "flow_to_scale_mass_ratio": integrated_mass / scale_delta,
        },
        "pressure": {
            "observed_initial_mpa": float(pressure_mpa[0]),
            "observed_final_mpa": float(pressure_mpa[-1]),
            "predicted_final_mpa": float(predicted_pressure[-1]),
            "rmse_mpa": float(np.sqrt(np.mean(pressure_error**2))),
            "mae_mpa": float(np.mean(np.abs(pressure_error))),
            "final_error_mpa": float(pressure_error[-1]),
        },
        "temperature": {
            "observed_initial_c": float(tank_temperature_c[0]),
            "observed_peak_c": float(np.max(tank_temperature_c)),
            "predicted_peak_c": float(np.max(predicted_temperature)),
            "rmse_c": float(np.sqrt(np.mean(temperature_error**2))),
            "mae_c": float(np.mean(np.abs(temperature_error))),
            "peak_error_c": float(
                np.max(predicted_temperature) - np.max(tank_temperature_c)
            ),
        },
    }


def analyze(source_dir: Path) -> dict[str, Any]:
    traces = [load_trace(source_dir / name) for name in FILES]
    workbooks: list[dict[str, Any]] = []
    eligible_sessions: list[tuple[WorkbookTrace, int, int, dict[str, Any]]] = []
    for trace in traces:
        flow = trace.channels["QT_D02 ValueY"]
        sessions = active_sessions(trace.elapsed_s, flow)
        summaries = [_session_summary(trace, start, end) for start, end in sessions]
        workbooks.append({
            "filename": trace.filename,
            "sha256": FILES[trace.filename]["sha256"],
            "sample_count": int(len(trace.elapsed_s)),
            "sampling_interval_s": float(np.median(np.diff(trace.elapsed_s))),
            "duration_s": float(trace.elapsed_s[-1]),
            "column_count": int(len(trace.channels)),
            "tank_candidate_columns_present": all(
                tag in trace.channels for tag in TANK_CANDIDATE_COLUMNS
            ),
            "active_session_count": len(sessions),
            "active_sessions": summaries,
        })
        if all(tag in trace.channels for tag in TANK_CANDIDATE_COLUMNS):
            eligible_sessions.extend(
                (trace, start, end, summary)
                for (start, end), summary in zip(sessions, summaries, strict=True)
                if _candidate_session_eligible(summary)
            )
    if not eligible_sessions:
        raise RuntimeError("No mass-consistent candidate tank session was found")
    replays = [
        _replay_candidate_session(
            trace,
            start,
            end,
            session_selection=(
                "all sessions passing pre-model mass-closure, minimum-mass and "
                "minimum-duration criteria"
            ),
        )
        for trace, start, end, _summary in eligible_sessions
    ]
    replay = max(
        replays,
        key=lambda item: item["mass_boundary"]["integrated_flow_kg"],
    )
    pressure_rmse = np.asarray([item["pressure"]["rmse_mpa"] for item in replays])
    temperature_rmse = np.asarray([item["temperature"]["rmse_c"] for item in replays])
    mass_ratios = np.asarray([
        item["mass_boundary"]["flow_to_scale_mass_ratio"] for item in replays
    ])
    joint_screen = (pressure_rmse <= 5.0) & (temperature_rmse <= 10.0)
    aggregate = {
        "case_count": len(replays),
        "selection_is_independent_of_model_prediction": True,
        "selection_criteria": {
            "flow_to_scale_mass_ratio": [
                MASS_CLOSURE_RATIO_MIN, MASS_CLOSURE_RATIO_MAX,
            ],
            "minimum_integrated_flow_mass_kg": MINIMUM_SESSION_MASS_KG,
            "minimum_duration_s": MINIMUM_SESSION_DURATION_S,
            "required_candidate_tank_columns": list(TANK_CANDIDATE_COLUMNS),
        },
        "pressure_rmse_mpa": {
            "case_mean": float(np.mean(pressure_rmse)),
            "case_median": float(np.median(pressure_rmse)),
            "minimum": float(np.min(pressure_rmse)),
            "maximum": float(np.max(pressure_rmse)),
        },
        "temperature_rmse_c": {
            "case_mean": float(np.mean(temperature_rmse)),
            "case_median": float(np.median(temperature_rmse)),
            "minimum": float(np.min(temperature_rmse)),
            "maximum": float(np.max(temperature_rmse)),
        },
        "flow_to_scale_mass_ratio": {
            "minimum": float(np.min(mass_ratios)),
            "maximum": float(np.max(mass_ratios)),
        },
        "project_screen": {
            "pressure_rmse_mpa_max": 5.0,
            "temperature_rmse_c_max": 10.0,
            "joint_pass_count": int(np.sum(joint_screen)),
            "joint_pass_fraction": float(np.mean(joint_screen)),
            "interpretation": "descriptive post-access screen, not confirmatory validation",
        },
    }
    return {
        "schema_version": 1,
        "generated_at": datetime.now(timezone.utc).isoformat(),
        "status": "completed_post_access_external_diagnostic",
        "evidence_role": "post_access_external_diagnostic_only",
        "model_commit": _git_commit(),
        "source": {
            "dataset_title": "Group B -- HySam system measurement data",
            "dataset_doi": "10.5281/zenodo.20590842",
            "dataset_license": "CC BY 4.0",
            "article_title": "Representative Hydrogen Sampling at Hydrogen Refuelling Stations: Interplay of Sampling Strategy and Station Parameters",
            "article_doi": "10.3390/cleantechnol8030091",
            "article_license": "CC BY 4.0",
            "article_reported_context": {
                "station": "physical test HRS with seven storage banks",
                "dispensing": "35 MPa and 70 MPa",
                "maximum_mass_flow_g_s": 120.0,
                "protocols": ["SAE J2601", "MC Formula", "PHRYDE", "free configurable"],
                "hysam_setup_1_sink": "244 L, 70 MPa Type-IV tank",
                "hysam_setup_1_reported_flow_g_s": "approximately 14-16",
                "logged_scope": "all dispenser and tank data",
            },
        },
        "access_integrity": {
            "all_expected_sha256_match": True,
            "raw_rows_committed": False,
            "raw_files_location": "gitignored data/public_validation/raw/methytrucks_hysam",
            "outcomes_inspected_before_protocol": True,
            "parameter_fitting_performed": False,
        },
        "channel_semantics": {
            "confirmed_by_internal_consistency": [
                "Zeit is a common 0.5 s time base",
                "QT_D02 is flow-like and its integral tracks FWg35_Masse changes",
                "FWg35_Masse is cumulative transferred-mass-like",
            ],
            "plausible_but_unconfirmed": [
                "PT01 as sink/tank pressure",
                "TT08..TT11 as sink/tank temperature measurements",
                "PTD10 as upstream/dispenser pressure",
                "TEX01 as delivered-gas temperature",
            ],
            "missing": [
                "publisher channel dictionary and engineering units",
                "unambiguous workbook/test-to-sampling-device crosswalk",
                "sensor calibration and uncertainty metadata",
                "controller state and storage-bank selection tags",
            ],
        },
        "workbooks": workbooks,
        "candidate_session_replays": replays,
        "candidate_session_aggregate": aggregate,
        "candidate_tank_replay_alias": (
            "largest integrated flow mass among eligible candidate_session_replays"
        ),
        "candidate_tank_replay": replay,
        "eligibility": {
            "component_diagnostic_eligible": True,
            "prospective_holdout_eligible": False,
            "quantitative_full_loop_validation_eligible": False,
            "reason": (
                "The rows are public and synchronized, but numerical outcomes were inspected "
                "before this protocol and the released package lacks a channel dictionary, "
                "controller/bank-state mapping and calibration metadata."
            ),
        },
        "highest_value_next_action": (
            "Obtain the publisher channel dictionary and test-to-device crosswalk. Then freeze "
            "an event mapping and rerun the existing model without tuning; request a disjoint "
            "uninspected logger event for prospective full-loop validation."
        ),
        "claim_boundary": (
            "This artifact proves file integrity, synchronized 0.5 s measurement availability, "
            "flow-to-cumulative-mass consistency and a no-fit candidate tank replay. It is a "
            "post-access diagnostic, not an untouched holdout, a full station-to-vehicle "
            "validation result, SAE certification or field-safety validation."
        ),
    }


def main() -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument(
        "--source-dir",
        type=Path,
        default=Path("data/public_validation/raw/methytrucks_hysam"),
    )
    parser.add_argument(
        "--output",
        type=Path,
        default=Path("research/methytrucks_hysam_postaccess_diagnostic_2026_10_08.json"),
    )
    args = parser.parse_args()
    report = analyze(args.source_dir)
    args.output.parent.mkdir(parents=True, exist_ok=True)
    args.output.write_text(
        json.dumps(report, ensure_ascii=False, indent=2) + "\n",
        encoding="utf-8",
        newline="\n",
    )
    replay = report["candidate_tank_replay"]
    aggregate = report["candidate_session_aggregate"]
    print(json.dumps({
        "status": report["status"],
        "workbook_count": len(report["workbooks"]),
        "candidate_session_count": aggregate["case_count"],
        "candidate_joint_screen_pass_count": aggregate["project_screen"]["joint_pass_count"],
        "candidate_pressure_rmse_mpa": replay["pressure"]["rmse_mpa"],
        "candidate_temperature_rmse_c": replay["temperature"]["rmse_c"],
        "full_loop_eligible": report["eligibility"]["quantitative_full_loop_validation_eligible"],
        "output": str(args.output),
    }, ensure_ascii=False))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
