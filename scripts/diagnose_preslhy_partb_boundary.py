"""Build a post-outcome boundary diagnostic for the PRESLHY Part-B holdout.

The primary Part-B result is frozen and remains negative.  This script only
quantifies whether the measured nozzle-pressure channel exposes a boundary
condition that is absent from the locked free-aperture model.  The implied
discharge-coefficient value is a sensitivity indicator, not a fitted runtime
parameter and must never be routed into the simulator.
"""

from __future__ import annotations

import argparse
from datetime import datetime, timezone
import hashlib
import json
import math
from pathlib import Path
import re
import subprocess

import numpy as np
from openpyxl import load_workbook


LOCKED_CD = 0.8


def _sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for block in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(block)
    return digest.hexdigest()


def _number(value: object) -> float | None:
    try:
        value = float(value)
    except (TypeError, ValueError):
        return None
    return value if math.isfinite(value) else None


def _read_pressure_channels(path: Path) -> dict[str, float | int | None]:
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        sheet = next(
            sheet for sheet in workbook.worksheets
            if "-vd" in sheet.title.lower() or "-v10" in sheet.title.lower()
        )
        rows = list(sheet.iter_rows(min_row=6, values_only=True))
        headers = [str(value or "").strip() for value in rows[0]]
        time_index = headers.index("Time new [s]")
        vessel_index = headers.index("P-Innen")
        nozzle_index = headers.index("P-Duese")
        values: list[tuple[float, float, float]] = []
        for row in rows[1:]:
            time_s = _number(row[time_index] if time_index < len(row) else None)
            vessel = _number(row[vessel_index] if vessel_index < len(row) else None)
            nozzle = _number(row[nozzle_index] if nozzle_index < len(row) else None)
            if time_s is not None and vessel is not None and nozzle is not None:
                values.append((time_s, vessel, nozzle))
        if len(values) < 20:
            raise ValueError("fewer than 20 finite pressure rows")
        array = np.asarray(values, dtype=float)
        release = np.flatnonzero(array[:, 2] >= 0.1)
        if not release.size:
            raise ValueError("no release start in P-Duese")
        start = int(release[0])
        pre = array[max(0, start - 500):start]
        post = array[start:min(len(array), start + 1000)]
        delta = post[:, 1] - post[:, 2]
        ratio = np.divide(post[:, 2], post[:, 1], out=np.full_like(post[:, 2], np.nan), where=post[:, 1] > 0.0)
        return {
            "finite_rows": int(len(array)),
            "release_index": start,
            "initial_vessel_pressure_bar_g": float(np.median(pre[:, 1])) if len(pre) else None,
            "post_release_median_vessel_nozzle_delta_bar": float(np.median(delta)),
            "post_release_p05_vessel_nozzle_delta_bar": float(np.quantile(delta, 0.05)),
            "post_release_p95_vessel_nozzle_delta_bar": float(np.quantile(delta, 0.95)),
            "post_release_median_nozzle_to_vessel_ratio": float(np.nanmedian(ratio)),
            "post_release_max_nozzle_pressure_bar_g": float(np.max(post[:, 2])),
        }
    finally:
        workbook.close()


def main() -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument("--raw", type=Path, default=Path("data/public_validation/raw/preslhy_partb_xlsx"))
    parser.add_argument("--result", type=Path, default=Path("research/preslhy_partb_ambient_validation.json"))
    parser.add_argument("--protocol", type=Path, default=Path("research/preslhy_partb_holdout_protocol.json"))
    parser.add_argument("--output", type=Path, default=Path("research/preslhy_partb_boundary_diagnostic_2026_10_11.json"))
    args = parser.parse_args()
    root = Path(__file__).resolve().parents[1]
    result = json.loads(args.result.read_text(encoding="utf-8"))
    protocol = json.loads(args.protocol.read_text(encoding="utf-8"))
    records = []
    for case in result["cases"]:
        path = args.raw / case["source_member"]
        channel = _read_pressure_channels(path)
        observed_half = case["experimental_time_to_50_percent_gauge_s"]
        predicted_half = case["predicted_time_to_50_percent_gauge_s"]
        implied_cd = None
        if observed_half and predicted_half:
            implied_cd = LOCKED_CD * predicted_half / observed_half
        records.append({
            "case_id": case["case_id"],
            "nozzle_diameter_mm": case["nozzle_diameter_mm"],
            "joint_primary_screen_pass": case["joint_primary_screen_pass"],
            "observed_half_pressure_time_s": observed_half,
            "predicted_half_pressure_time_s": predicted_half,
            "locked_discharge_coefficient": LOCKED_CD,
            "implied_cd_if_only_cd_changed": implied_cd,
            "pressure_channel_diagnostic": channel,
            "source_file_sha256": _sha256(path),
        })
    payload = {
        "schema_version": 1,
        "artifact_type": "preslhy_partb_post_outcome_boundary_diagnostic",
        "generated_at": datetime.now(timezone.utc).isoformat(),
        "source_commit": subprocess.check_output(["git", "rev-parse", "HEAD"], cwd=root, text=True).strip(),
        "primary_result": str(args.result),
        "protocol": str(args.protocol),
        "protocol_sha256": _sha256(args.protocol),
        "status": "POST_OUTCOME_DIAGNOSTIC_ONLY",
        "model_boundary": {
            "locked_model": "well-mixed adiabatic vessel with free-aperture real-gas discharge",
            "locked_discharge_coefficient": LOCKED_CD,
            "runtime_parameter_application": False,
            "primary_gate_effect": "none",
        },
        "cases": records,
        "interpretation": {
            "observation": "The 2 mm cases fail the frozen half-time screen while showing a smaller measured vessel-to-nozzle pressure difference than the 4 mm cases.",
            "diagnostic_hypotheses": [
                "The locked free-aperture boundary may not represent diameter-dependent contraction, valve/orifice coefficient, or instrument location.",
                "The P-Duese channel is a measured diagnostic boundary and is not substituted into the primary model after outcome access.",
            ],
            "not_supported": [
                "No discharge coefficient is fitted or promoted to runtime.",
                "No primary validation gate is changed.",
                "No station safety-distance, ignition, dispersion or vehicle-loop claim is made.",
            ],
        },
        "claim_boundary": "Post-outcome development diagnostic only. The frozen Part-B result remains a negative 3/5 joint screen and this artifact cannot replace it or authorize a runtime parameter change.",
    }
    args.output.write_text(json.dumps(payload, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    print(args.output)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
