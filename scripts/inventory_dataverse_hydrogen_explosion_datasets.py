"""Inventory two CC0 DataverseNO hydrogen-explosion datasets.

The inventories intentionally stop at provenance, file coverage, and the
published experiment-summary workbooks.  They do not claim that an explosion
dataset validates the complete HRS fueling loop or a consequence model until a
pre-registered model comparison is run against the raw traces.
"""

from __future__ import annotations

import argparse
import hashlib
import json
import re
import urllib.parse
import urllib.request
from datetime import datetime, timezone
from pathlib import Path
from typing import Any

from openpyxl import load_workbook


SOURCES: tuple[dict[str, Any], ...] = (
    {
        "id": "wskbij_large_scale_obstructed_releases",
        "doi": "10.18710/WSKBIJ",
        "title": "Replication dataset for: Large-Scale Hydrogen Explosion Experiments: Obstructed Releases in open atmosphere",
        "record": "https://dataverse.no/dataset.xhtml?persistentId=doi:10.18710/WSKBIJ",
        "api": "https://dataverse.no/api/datasets/:persistentId/?persistentId=doi%3A10.18710%2FWSKBIJ",
        "license": "CC0 1.0",
        "summary_file": "wsk_results.xlsx",
        "expected_summary_md5": "73ee4f29ca08612f46b7a37980cbe4df",
    },
    {
        "id": "x044qk_ignited_jets_overpressure",
        "doi": "10.18710/X044QK",
        "title": "Replication data for: Laboratory-Scale Experiments on Ignited Hydrogen Jets: Flame Acceleration and Overpressure Analysis",
        "record": "https://dataverse.no/dataset.xhtml?persistentId=doi:10.18710/X044QK",
        "api": "https://dataverse.no/api/datasets/:persistentId/?persistentId=doi%3A10.18710%2FX044QK",
        "license": "CC0 1.0",
        "summary_file": "x_results.xlsx",
        "expected_summary_md5": "6c929ee530ac393ee8e76f6db2eff05c",
    },
)


def _hashes(path: Path) -> dict[str, Any]:
    md5 = hashlib.md5()
    sha256 = hashlib.sha256()
    with path.open("rb") as stream:
        for block in iter(lambda: stream.read(1024 * 1024), b""):
            md5.update(block)
            sha256.update(block)
    return {"bytes": path.stat().st_size, "md5": md5.hexdigest(), "sha256": sha256.hexdigest()}


def _fetch_json(url: str) -> dict[str, Any]:
    with urllib.request.urlopen(url, timeout=60) as response:
        return json.load(response)


def _api_summary(payload: dict[str, Any]) -> dict[str, Any]:
    version = payload["data"]["latestVersion"]
    files = [item.get("dataFile", item) for item in version.get("files", [])]
    text_files = [item for item in files if str(item.get("filename", "")).lower().endswith(".txt")]
    csv_files = [item for item in files if str(item.get("filename", "")).lower().endswith(".csv")]
    video_files = [item for item in files if str(item.get("filename", "")).lower().endswith((".cine", ".zip"))]
    workbook_files = [item for item in files if str(item.get("filename", "")).lower().endswith(".xlsx")]
    return {
        "identifier": payload["data"].get("identifier"),
        "publication_date": payload["data"].get("publicationDate"),
        "version": version.get("versionNumber"),
        "license": version.get("license", {}).get("name"),
        "file_count": len(files),
        "file_groups": {
            "text_trace_files": len(text_files),
            "csv_trace_files": len(csv_files),
            "video_or_archive_files": len(video_files),
            "workbook_files": len(workbook_files),
        },
        "total_bytes": sum(int(item.get("filesize", 0)) for item in files),
        "largest_file_bytes": max((int(item.get("filesize", 0)) for item in files), default=0),
        "readme_files": [item.get("filename") for item in files if "readme" in str(item.get("filename", "")).lower()],
        "summary_workbooks": [
            {
                "filename": item.get("filename"),
                "id": item.get("id"),
                "bytes": item.get("filesize"),
                "md5": item.get("checksum", {}).get("value"),
            }
            for item in workbook_files
        ],
    }


def _number(value: Any) -> float | None:
    try:
        if value in (None, "", "-"):
            return None
        return float(value)
    except (TypeError, ValueError):
        return None


def _wsk_workbook(path: Path) -> dict[str, Any]:
    workbook = load_workbook(path, data_only=True, read_only=True)
    sheet = workbook["Flow parameters"]
    rows = list(sheet.iter_rows(values_only=True))
    entries = [row for row in rows[4:] if re.fullmatch(r"\d+\*{0,2}", str(row[0]).strip())]
    nozzle = [_number(row[1]) for row in entries]
    reservoir = [_number(row[2]) for row in entries]
    sensor_values = [[_number(row[index]) for row in entries] for index in range(4, 8)]
    measured = [value for column in sensor_values for value in column if value is not None]
    return {
        "summary_kind": "published_experiment_summary",
        "experiment_count": len(entries),
        "nozzle_diameter_mm_range": [min(v for v in nozzle if v is not None), max(v for v in nozzle if v is not None)],
        "reservoir_pressure_bar_range": [min(v for v in reservoir if v is not None), max(v for v in reservoir if v is not None)],
        "explosion_pressure_sensor_count": 4,
        "explosion_pressure_kpa_range": [min(measured), max(measured)],
        "experiments_with_numeric_pressure_sensor": sum(
            any(_number(value) is not None for value in row[4:8]) for row in entries
        ),
        "rows_with_missing_explosion_pressure": sum(
            not all(_number(value) is not None for value in row[4:8]) for row in entries
        ),
        "source_notes": [
            "Summary workbook contains pressure peaks by experiment; raw text traces and selected high-speed camera files are separate API files.",
            "The README states that mass-flow instrumentation was available for experiments 1-27 only and ambient weather was not documented.",
        ],
    }


def _x044_workbook(path: Path) -> dict[str, Any]:
    workbook = load_workbook(path, data_only=True, read_only=True)
    rows = list(workbook["Sheet1"].iter_rows(values_only=True))
    entries = [row for row in rows[2:] if _number(row[0]) is not None]
    max_pressure = [_number(row[1]) for row in entries]
    max_mfr = [_number(row[2]) for row in entries]
    average_pressure = [_number(row[3]) for row in entries]
    average_mfr = [_number(row[4]) for row in entries]
    obstacles = [_number(row[6]) for row in entries]
    comments = [str(row[9]) for row in entries if len(row) > 9 and row[9]]
    return {
        "summary_kind": "published_experiment_summary",
        "experiment_count": len(entries),
        "max_pressure_bar_range": [min(v for v in max_pressure if v is not None), max(v for v in max_pressure if v is not None)],
        "max_mass_flow_kg_s_range": [min(v for v in max_mfr if v is not None), max(v for v in max_mfr if v is not None)],
        "average_pressure_bar_range": [min(v for v in average_pressure if v is not None), max(v for v in average_pressure if v is not None)],
        "average_mass_flow_kg_s_range": [min(v for v in average_mfr if v is not None), max(v for v in average_mfr if v is not None)],
        "obstacle_distance_cm_range": [min(v for v in obstacles if v is not None), max(v for v in obstacles if v is not None)],
        "commented_experiments": comments,
        "source_notes": [
            "Workbook reports pressure and mass-flow summaries for 40 numbered experiments; CSV files contain the large raw traces.",
            "The README describes four high-frequency piezoelectric pressure sensors and an upstream pressure transmitter.",
        ],
    }


def build(raw_directory: Path, fetch_api: bool = True) -> dict[str, Any]:
    sources: list[dict[str, Any]] = []
    for source in SOURCES:
        item = dict(source)
        item["api_observation"] = None
        item["api_error"] = None
        if fetch_api:
            try:
                item["api_observation"] = _api_summary(_fetch_json(source["api"]))
            except Exception as exc:  # preserve availability failure in the record
                item["api_error"] = str(exc)
        summary_path = raw_directory / source["summary_file"]
        item["summary_file_present"] = summary_path.is_file()
        item["summary_file_identity_match"] = False
        item["summary_file_hashes"] = None
        item["published_summary"] = None
        if summary_path.is_file():
            hashes = _hashes(summary_path)
            item["summary_file_hashes"] = hashes
            item["summary_file_identity_match"] = hashes["md5"] == source["expected_summary_md5"]
            item["published_summary"] = (
                _wsk_workbook(summary_path) if source["id"].startswith("wskbij") else _x044_workbook(summary_path)
            )
        sources.append(item)
    return {
        "schema_version": 1,
        "generated_at": datetime.now(timezone.utc).isoformat(),
        "status": "completed_dataverse_provenance_and_summary_inventory",
        "evidence_role": "independent_hydrogen_release_ignition_overpressure_component_evidence",
        "claim_boundary": (
            "These CC0 datasets provide independent release, ignition, pressure and mass-flow evidence for consequence-component checks. "
            "They do not contain a synchronized gaseous H70 station-to-vehicle fueling loop, and this inventory does not perform a model comparison. "
            "They therefore cannot close the full-loop external-validation gate or establish SAGA effectiveness."
        ),
        "sources": sources,
        "validation_status": {
            "model_comparison_performed": False,
            "numeric_validation_gate_closed": False,
            "full_loop_external_validation_supported": False,
            "next_step": (
                "Freeze a consequence-model protocol and inclusion criteria, then run the locked model against raw traces without changing parameters after reading outcomes."
            ),
        },
    }


def markdown(result: dict[str, Any]) -> str:
    lines = [
        "# DataverseNO hydrogen explosion dataset inventory",
        "",
        f"Generated: `{result['generated_at']}`",
        "",
        "This report fixes the public provenance and summary-workbook boundary. It is not a model-validation result.",
        "",
        "## Public sources",
        "",
        "| Dataset | DOI | License | Files | Trace groups | Summary identity |",
        "|---|---|---|---:|---|---|",
    ]
    for source in result["sources"]:
        api = source.get("api_observation") or {}
        groups = api.get("file_groups", {})
        trace_groups = ", ".join(f"{key}={value}" for key, value in groups.items()) or "unavailable"
        lines.append(
            f"| [{source['title']}]({source['record']}) | `{source['doi']}` | `{source['license']}` | "
            f"{api.get('file_count', 'n/a')} | {trace_groups} | **{source['summary_file_identity_match']}** |"
        )
    lines += ["", "## Published summary coverage", ""]
    for source in result["sources"]:
        lines += [f"### {source['title']}", ""]
        summary = source.get("published_summary")
        if summary is None:
            lines.append("The local summary workbook is unavailable; no numerical summary was reproduced.")
            continue
        for key, value in summary.items():
            if key != "source_notes":
                lines.append(f"- **{key}:** `{value}`")
        for note in summary["source_notes"]:
            lines.append(f"- {note}")
        lines.append("")
    lines += [
        "## Claim boundary",
        "",
        result["claim_boundary"],
        "",
        f"- Model comparison performed: **{result['validation_status']['model_comparison_performed']}**",
        f"- Numeric validation gate closed: **{result['validation_status']['numeric_validation_gate_closed']}**",
        f"- Full-loop external validation supported: **{result['validation_status']['full_loop_external_validation_supported']}**",
        "",
        result["validation_status"]["next_step"],
        "",
    ]
    return "\n".join(lines)


def main() -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument("--raw-directory", type=Path, default=Path("tmp/dataverse_hydrogen_explosion_summaries"))
    parser.add_argument("--json-output", type=Path, default=Path("research/dataverse_hydrogen_explosion_dataset_inventory_2026_10_05.json"))
    parser.add_argument("--report-output", type=Path, default=Path("research/DATAVERSE_HYDROGEN_EXPLOSION_DATASET_INVENTORY_2026_10_05.md"))
    parser.add_argument("--no-api", action="store_true")
    args = parser.parse_args()
    result = build(args.raw_directory, fetch_api=not args.no_api)
    args.json_output.parent.mkdir(parents=True, exist_ok=True)
    args.report_output.parent.mkdir(parents=True, exist_ok=True)
    args.json_output.write_text(json.dumps(result, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    args.report_output.write_text(markdown(result), encoding="utf-8")
    print(json.dumps({"sources": len(result["sources"]), "api_ok": sum(item["api_observation"] is not None for item in result["sources"]), "summary_ok": sum(item["summary_file_identity_match"] for item in result["sources"])}))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
