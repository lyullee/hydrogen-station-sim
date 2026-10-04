"""Verify the locally acquired JRC HIAD 2.2 workbook without opening response text.

The workbook is ignored from Git because it is a third-party source file. This
script records its immutable digest, sheet dimensions and the deterministic HRS
selection used by the public-evidence inventory. It does not freeze or approve
a blinded expert holdout.
"""

from __future__ import annotations

import argparse
from hashlib import sha256
import json
from pathlib import Path

from openpyxl import load_workbook

from h2station.public_validation import read_hiad_hrs_cases


DOWNLOAD_URL = (
    "https://minerva.jrc.ec.europa.eu/en/shorturl/capri/"
    "hiad_22_export_for_users_2026_01_01xlsx"
)
LANDING_PAGE = "https://minerva.jrc.ec.europa.eu/en/shorturl/capri/hiadpt"


def build_verification(source: Path) -> dict:
    try:
        display_file = source.resolve().relative_to(Path.cwd().resolve()).as_posix()
    except ValueError:
        display_file = source.as_posix()
    digest = sha256(source.read_bytes()).hexdigest()
    workbook = load_workbook(source, read_only=True, data_only=True)
    dimensions = {
        sheet.title: {"rows": sheet.max_row, "columns": sheet.max_column}
        for sheet in workbook.worksheets
    }
    cases = read_hiad_hrs_cases(source)
    event_ids = [case.event_id for case in cases]
    return {
        "schema_version": 1,
        "verified_on": "2026-10-04",
        "source": {
            "title": "European Hydrogen Incidents and Accidents Database HIAD 2.2",
            "publisher": "European Commission Joint Research Centre",
            "landing_page": LANDING_PAGE,
            "download_url": DOWNLOAD_URL,
            "local_file": display_file,
            "sha256": digest,
            "size_bytes": source.stat().st_size,
            "version": "HIAD 2.2",
            "event_cutoff": "2025-12-31",
        },
        "workbook": {
            "sheet_count": len(dimensions),
            "sheet_dimensions": dimensions,
        },
        "selection": {
            "rule": "FACILITY.Application exactly equals 'Hydrogen refuelling station'",
            "case_count": len(event_ids),
            "unique_event_ids": len(event_ids) == len(set(event_ids)),
            "event_ids_sha256": sha256("\n".join(event_ids).encode("utf-8")).hexdigest(),
            "response_text_read": False,
            "coordinator_approval": False,
        },
        "claim_boundary": (
            "This verifies provenance and deterministic extraction of public HIAD HRS metadata. "
            "It is not a frozen casebook, does not approve response collection, and does not "
            "establish accident probabilities, consequence physics or SAGA effectiveness."
        ),
    }


def main() -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument(
        "--source",
        type=Path,
        default=Path("data/public_validation/raw/hiad_2_2/HIAD 2.2.xlsx"),
    )
    parser.add_argument(
        "--output",
        type=Path,
        default=Path("research/hiad_2_2_source_verification_2026_10_04.json"),
    )
    args = parser.parse_args()
    result = build_verification(args.source)
    args.output.parent.mkdir(parents=True, exist_ok=True)
    args.output.write_text(json.dumps(result, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    print(json.dumps({"sha256": result["source"]["sha256"], "case_count": result["selection"]["case_count"]}))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
