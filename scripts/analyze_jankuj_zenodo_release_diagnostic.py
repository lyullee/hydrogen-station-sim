"""Create a claim-bounded diagnostic from the Jankuj actual-H2 release archive.

The Zenodo archive contains pressure and thermocouple workbooks from a 50 L,
200 bar cylinder penetration experiment.  The workbooks are kept outside the
repository.  This script stores only aggregate descriptors and provenance
boundaries so the evidence can guide consequence and emergency-response text
without pretending that an undocumented breach geometry is a calibrated
release model.
"""

from __future__ import annotations

import argparse
import hashlib
import json
from datetime import datetime, timezone
from pathlib import Path

import numpy as np
import openpyxl


ROOT = Path(__file__).resolve().parents[1]
DOI = "10.5281/zenodo.17913628"


def _md5(path: Path) -> str:
    digest = hashlib.md5()
    with path.open("rb") as stream:
        for block in iter(lambda: stream.read(1024 * 1024), b""):
            digest.update(block)
    return digest.hexdigest()


def _numeric_columns(path: Path, *, header_row: int = 1) -> list[np.ndarray]:
    workbook = openpyxl.load_workbook(path, read_only=True, data_only=True)
    sheet = workbook.active
    columns: list[list[float]] = [[] for _ in range(sheet.max_column)]
    for row in sheet.iter_rows(min_row=header_row + 1, values_only=True):
        for index, value in enumerate(row):
            if isinstance(value, (int, float)) and np.isfinite(value):
                columns[index].append(float(value))
    return [np.asarray(values, dtype=float) for values in columns]


def _pressure_summary(path: Path) -> dict[str, object]:
    workbook = openpyxl.load_workbook(path, read_only=True, data_only=True)
    sheet = workbook.active
    rows = [row for row in sheet.iter_rows(min_row=2, values_only=True) if row[0] is not None]
    values = np.asarray(
        [[float(row[0]), float(row[1]), float(row[2])] for row in rows],
        dtype=float,
    )
    time = values[:, 0]
    channels: list[dict[str, object]] = []
    for channel_index, name in ((1, "pressure_n1"), (2, "pressure_n2")):
        pressure = values[:, channel_index]
        initial = float(np.median(pressure[:10]))
        half_target = initial * 0.5
        tenth_target = initial * 0.1
        half_indices = np.flatnonzero(pressure <= half_target)
        tenth_indices = np.flatnonzero(pressure <= tenth_target)
        jumps = np.abs(np.diff(pressure))
        channels.append({
            "channel": name,
            "sample_count": int(pressure.size),
            "time_start_s": float(time[0]),
            "time_end_s": float(time[-1]),
            "initial_median_bar": initial,
            "final_median_bar": float(np.median(pressure[-10:])),
            "minimum_bar": float(np.min(pressure)),
            "maximum_bar": float(np.max(pressure)),
            "half_pressure_time_s": (
                float(time[int(half_indices[0])]) if half_indices.size else None
            ),
            "ten_percent_pressure_time_s": (
                float(time[int(tenth_indices[0])]) if tenth_indices.size else None
            ),
            "large_jump_count_gt_20_bar": int(np.count_nonzero(jumps > 20.0)),
            "interpretation": (
                "continuous pressure-release channel retained for descriptive replay"
                if name == "pressure_n2"
                else "channel contains a >20 bar reset after the first release and is retained as a segmented diagnostic"
            ),
        })
    return {
        "file": path.name,
        "file_md5": _md5(path),
        "time_unit": "s",
        "pressure_unit": "bar as published",
        "channels": channels,
        "primary_descriptive_channel": "pressure_n2",
    }


def _temperature_summary(path: Path, scenario: str) -> dict[str, object]:
    columns = _numeric_columns(path)
    # Column 0 is the elapsed-second index.  The remaining columns are the
    # thermocouple channels; empty cells before/after the event are omitted by
    # _numeric_columns without inventing values.
    data = [column for column in columns[1:] if column.size]
    return {
        "file": path.name,
        "file_md5": _md5(path),
        "scenario": scenario,
        "channel_count": len(data),
        "sample_count_by_channel": [int(column.size) for column in data],
        "minimum_c": float(min(np.min(column) for column in data)),
        "maximum_c": float(max(np.max(column) for column in data)),
        "first_window_median_c": [float(np.median(column[:5])) for column in data],
        "last_window_median_c": [float(np.median(column[-5:])) for column in data],
    }


def build_report(raw_dir: Path, protocol_path: Path) -> dict[str, object]:
    pressure_path = raw_dir / "Figure 7. Pressure drop during hydrogen release.xlsx"
    ignition_path = raw_dir / "Figure 8. Hydrogen release and subsequent spark ignition.xlsx"
    no_ignition_path = raw_dir / "Figure 9. Hydrogen release without ignition.xlsx"
    required = (pressure_path, ignition_path, no_ignition_path)
    if not all(path.is_file() for path in required):
        missing = [str(path) for path in required if not path.is_file()]
        raise FileNotFoundError("missing Jankuj archive files: " + ", ".join(missing))

    pressure = _pressure_summary(pressure_path)
    temperatures = [
        _temperature_summary(ignition_path, "spark_ignition"),
        _temperature_summary(no_ignition_path, "no_ignition"),
    ]
    return {
        "schema_version": 1,
        "artifact_type": "post_access_zenodo_17913628_actual_hydrogen_release_diagnostic",
        "generated_at": datetime.now(timezone.utc).isoformat(),
        "status": "POST_ACCESS_DESCRIPTIVE_RELEASE_DIAGNOSTIC_NOT_INDEPENDENT_MODEL_VALIDATION",
        "source": {
            "dataset_doi": DOI,
            "landing_page": "https://zenodo.org/records/17913628",
            "title": "Dataset for: Accidental self-ignition of hydrogen released from pressurized cylinder (Figures 7-9)",
            "publisher": "Zenodo / Jankuj, Lepik, Salzano and Mynarz",
            "license": "CC BY 4.0",
            "actual_hydrogen": True,
            "cylinder_volume_l": 50.0,
            "initial_pressure_bar": 200.0,
            "raw_files_external_only": True,
        },
        "provenance_boundary": {
            "raw_accessed_before_this_artifact": True,
            "independent_holdout": False,
            "prospective_validation": False,
            "parameter_fitting": False,
            "runtime_parameter_change": False,
            "readiness_gate_effect": False,
            "reason": (
                "The archive was accessed after the current validation program was frozen; "
                "the penetration/breach geometry is not documented as a calibrated nozzle."
            ),
        },
        # Keep the committed artifact portable and avoid leaking a developer's
        # local checkout path into the LLM evidence envelope.
        "protocol_reference": "research/jankuj_zenodo_17913628_pressure_holdout_protocol_2026_10_14.json",
        "analysis_contract": {
            "pressure": "aggregate replay of both published pressure channels; channel n2 is retained as the continuous descriptive release trace",
            "temperature": "aggregate extrema and baseline/end medians for the spark-ignition and no-ignition workbooks",
            "raw_rows_exposed": False,
            "no_fit_or_interpolation": True,
        },
        "pressure": pressure,
        "temperature": temperatures,
        "eligibility": {
            "descriptive_actual_hydrogen_release_evidence": True,
            "quantitative_release_model_validation": False,
            "reason": "No documented breach diameter, upstream restriction state or discharge coefficient is supplied for a locked release-model comparison.",
            "component_consequence_context": True,
            "ignition_scenario_context": True,
            "runtime_parameter_application": False,
        },
        "claim_supported": False,
        "claim_limit": (
            "This artifact supports descriptive actual-hydrogen pressure-depletion, cooling and "
            "ignition-scenario context. It does not validate a release coefficient, pressure model, "
            "heat flux, radiation, dispersion, detector placement, ESD effectiveness, safety distance "
            "or station-to-vehicle behavior."
        ),
    }


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--raw-dir", type=Path, required=True)
    parser.add_argument(
        "--protocol",
        type=Path,
        default=ROOT / "research/jankuj_zenodo_17913628_pressure_holdout_protocol_2026_10_14.json",
    )
    parser.add_argument(
        "--output",
        type=Path,
        default=ROOT / "research/jankuj_zenodo_17913628_postaccess_diagnostic_2026_10_14.json",
    )
    parser.add_argument(
        "--markdown-output",
        type=Path,
        default=ROOT / "research/JANKUJ_ZENODO_17913628_POSTACCESS_DIAGNOSTIC_2026_10_14.md",
    )
    args = parser.parse_args()
    report = build_report(args.raw_dir, args.protocol)
    args.output.write_text(json.dumps(report, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    args.markdown_output.write_text(
        "# Jankuj Zenodo 17913628 actual-hydrogen release diagnostic\n\n"
        "This is a post-access descriptive diagnostic. It is not an independent model holdout.\n\n"
        f"- DOI: {DOI}\n"
        f"- Pressure samples: {report['pressure']['channels'][1]['sample_count']} on the continuous n2 channel\n"
        f"- Pressure half-time (descriptive): {report['pressure']['channels'][1]['half_pressure_time_s']} s\n"
        f"- Ignition temperature range: {report['temperature'][0]['minimum_c']} to {report['temperature'][0]['maximum_c']} °C\n"
        f"- No-ignition temperature range: {report['temperature'][1]['minimum_c']} to {report['temperature'][1]['maximum_c']} °C\n\n"
        f"**Boundary:** {report['claim_limit']}\n",
        encoding="utf-8",
    )
    print(json.dumps({"pressure": report["pressure"], "temperature": report["temperature"]}, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
