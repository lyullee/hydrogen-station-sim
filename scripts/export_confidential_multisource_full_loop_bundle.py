"""Create a controlled full-loop evaluator bundle from one multi-sheet workbook.

The tool exists for authorised station records whose signals are separated by
subsystem worksheet.  Source headers, sheet names, identities, absolute times,
and raw records remain in the controlled workspace.  Only a de-identified,
canonical trace is written to an output directory outside the repository.

A successful export proves only that the declared channels passed a controlled
schema, time-alignment, and trace-quality screen.  It is not model validation,
safety certification, or evidence of journal readiness.
"""

from __future__ import annotations

import argparse
from bisect import bisect_left
import csv
from datetime import datetime, timezone
import hashlib
from io import StringIO
import json
import math
from pathlib import Path
from typing import Any

try:  # Supports both ``python scripts/...`` and test-module imports.
    from export_confidential_full_loop_bundle import (
        NUMERIC_COLUMNS,
        OUTPUT_COLUMNS,
        REQUIRED_METADATA,
        ROOT,
        STATE_COLUMNS,
        _finite,
        _json,
        _outside_repository,
        _relative_time,
        _sha256,
    )
    from audit_controlled_data_schema import _csv_encodings
except ModuleNotFoundError:  # pragma: no cover - import style depends on launcher
    from scripts.export_confidential_full_loop_bundle import (
        NUMERIC_COLUMNS,
        OUTPUT_COLUMNS,
        REQUIRED_METADATA,
        ROOT,
        STATE_COLUMNS,
        _finite,
        _json,
        _outside_repository,
        _relative_time,
        _sha256,
    )
    from scripts.audit_controlled_data_schema import _csv_encodings

from intake_external_hrs_bundle import build_manifest
from validate_external_hrs_full_loop import validate_full_loop_trace
from validate_external_hrs_manifest import validate as validate_manifest


def _require_attestation(attestation: dict[str, Any]) -> None:
    if attestation.get("schema_version") != 1:
        raise ValueError("attestation schema_version must be 1")
    if attestation.get("authorised_controlled_evaluation") is not True:
        raise ValueError("attestation must explicitly authorize controlled evaluation")
    if attestation.get("outcomes_accessed_before_protocol_freeze") is not False:
        raise ValueError("attestation must confirm no outcome access before protocol freeze")

    units = attestation.get("units")
    expected_units = {
        "vehicle_pressure_mpa_abs": "MPa_abs",
        "temperature_degC": "degC",
        "mass_flow_g_s": "g/s",
        "station_pressure_mpa_abs": "MPa_abs",
        "delivered_gas_temperature_degC": "degC",
        "cascade_source_pressure_mpa_abs": "MPa_abs",
    }
    if not isinstance(units, dict) or any(
        units.get(column) != expected for column, expected in expected_units.items()
    ):
        raise ValueError("attestation must declare the canonical numeric units")
    states = attestation.get("state_semantics")
    if not isinstance(states, dict) or not all(
        isinstance(states.get(column), str) and states[column].strip()
        for column in STATE_COLUMNS
    ):
        raise ValueError("attestation must provide nonempty semantics for every state column")
    metadata = attestation.get("metadata")
    if not isinstance(metadata, dict) or not all(
        isinstance(metadata.get(key), str) and metadata[key].strip()
        for key in REQUIRED_METADATA
    ):
        raise ValueError("attestation must provide all required generic metadata declarations")
    synchronization = attestation.get("source_synchronization")
    if not isinstance(synchronization, dict):
        raise ValueError("attestation must declare source synchronization")
    if synchronization.get("common_time_basis_confirmed") is not True:
        raise ValueError("attestation must confirm a common time basis")
    if not isinstance(synchronization.get("alignment_method"), str) or not synchronization["alignment_method"].strip():
        raise ValueError("attestation must declare the alignment method")


def _specifications(mapping: dict[str, Any]) -> tuple[list[dict[str, Any]], int, float]:
    if mapping.get("schema_version") != 1:
        raise ValueError("mapping schema_version must be 1")
    sources = mapping.get("sources")
    if not isinstance(sources, list) or not sources:
        raise ValueError("mapping requires at least one source specification")
    normalized: list[dict[str, Any]] = []
    canonical_owners: dict[str, int] = {}
    for index, raw in enumerate(sources):
        if not isinstance(raw, dict):
            raise ValueError("each source specification must be an object")
        worksheet = raw.get("worksheet")
        time_column = raw.get("time_column")
        source_file = raw.get("file")
        header_row = raw.get("header_row", 1)
        column_map = raw.get("column_map")
        if worksheet is None:
            worksheet = ""
        if not isinstance(worksheet, str):
            raise ValueError("worksheet mapping must be a string when supplied")
        if not isinstance(time_column, str) or not time_column.strip():
            raise ValueError("every source requires a nonempty time_column mapping")
        if not isinstance(header_row, int) or header_row < 1:
            raise ValueError("every source header_row must be a positive integer")
        if not isinstance(column_map, dict) or not column_map:
            raise ValueError("every source requires a nonempty column_map")
        if source_file is not None:
            if not isinstance(source_file, str) or not source_file.strip():
                raise ValueError("source file must be a nonempty relative path")
            source_path = Path(source_file)
            if (
                source_path.is_absolute()
                or ".." in source_path.parts
            ):
                raise ValueError("source file must be a nonempty relative path")
        for canonical, source_name in column_map.items():
            if canonical not in OUTPUT_COLUMNS[1:]:
                raise ValueError("source column_map contains an unsupported canonical channel")
            if not isinstance(source_name, str) or not source_name.strip():
                raise ValueError("source column_map values must be nonempty strings")
            if canonical in canonical_owners:
                raise ValueError("each canonical channel must be mapped by exactly one source")
            canonical_owners[canonical] = index
        normalized.append({
            "worksheet": worksheet,
            "file": source_file,
            "time_column": time_column,
            "time_format": raw.get("time_format"),
            "header_row": header_row,
            "column_map": column_map,
        })
    missing = [channel for channel in OUTPUT_COLUMNS[1:] if channel not in canonical_owners]
    if missing:
        raise ValueError("mapping does not cover every canonical full-loop channel")

    alignment = mapping.get("alignment")
    if not isinstance(alignment, dict):
        raise ValueError("mapping requires an alignment object")
    if alignment.get("method") != "nearest_observation":
        raise ValueError("mapping alignment.method must be nearest_observation")
    anchor = alignment.get("anchor_source", 0)
    maximum_offset = alignment.get("maximum_offset_s")
    if not isinstance(anchor, int) or not 0 <= anchor < len(normalized):
        raise ValueError("alignment.anchor_source must identify a mapped source")
    if not isinstance(maximum_offset, (int, float)) or not math.isfinite(maximum_offset) or maximum_offset < 0:
        raise ValueError("alignment.maximum_offset_s must be a finite nonnegative number")
    return normalized, anchor, float(maximum_offset)


def _source_path(input_data: Path, specification: dict[str, Any]) -> Path:
    """Resolve a declared source without accepting an arbitrary external path."""

    if input_data.is_file():
        if specification.get("file") is not None:
            raise ValueError("a file mapping is only valid when input is a directory")
        return input_data
    source_file = specification.get("file")
    if not isinstance(source_file, str):
        raise ValueError("directory intake requires every source to declare file")
    resolved_root = input_data.resolve()
    resolved_source = (resolved_root / source_file).resolve()
    try:
        resolved_source.relative_to(resolved_root)
    except ValueError as exc:
        raise ValueError("source file must remain under the controlled input directory") from exc
    if not resolved_source.is_file():
        raise ValueError("a declared source file is not present in the controlled input directory")
    return resolved_source


def _read_csv_source(
    source_path: Path,
    specification: dict[str, Any],
    *,
    include_rows: bool,
) -> tuple[tuple[str, ...], list[dict[str, object]] | None]:
    raw = source_path.read_bytes()
    for encoding in _csv_encodings(raw):
        try:
            parsed = list(csv.reader(StringIO(raw.decode(encoding))))
        except UnicodeError:
            continue
        header_index = specification["header_row"] - 1
        if header_index >= len(parsed):
            return (), [] if include_rows else None
        header = tuple(value.strip() for value in parsed[header_index])
        if not include_rows:
            return header, None
        rows = [
            {
                header[position]: value
                for position, value in enumerate(values)
                if position < len(header)
            }
            for values in parsed[header_index + 1:]
            if any(value.strip() for value in values)
        ]
        return header, rows
    raise ValueError("a declared CSV source cannot be decoded with supported logger encodings")


def _read_source(
    input_data: Path,
    specification: dict[str, Any],
    *,
    include_rows: bool,
) -> tuple[tuple[str, ...], list[dict[str, object]] | None]:
    source_path = _source_path(input_data, specification)
    if source_path.suffix.casefold() == ".csv":
        return _read_csv_source(source_path, specification, include_rows=include_rows)
    if source_path.suffix.casefold() not in {".xlsx", ".xlsm"}:
        raise ValueError("controlled multi-source intake supports CSV, XLSX, or XLSM sources")
    try:
        from openpyxl import load_workbook
    except ImportError as exc:  # pragma: no cover - environment-specific
        raise RuntimeError("openpyxl is required for controlled Excel intake") from exc
    workbook = load_workbook(source_path, read_only=True, data_only=True)
    try:
        worksheet_name = specification["worksheet"]
        if not worksheet_name:
            raise ValueError("XLSX/XLSM sources require a nonempty worksheet mapping")
        if worksheet_name not in workbook.sheetnames:
            raise ValueError("a specified worksheet is not present in the controlled input")
        worksheet = workbook[worksheet_name]
        header_row = specification["header_row"]
        header_values = next(
            worksheet.iter_rows(min_row=header_row, max_row=header_row, values_only=True),
            None,
        )
        if header_values is None:
            return (), [] if include_rows else None
        header = tuple(str(value).strip() if value is not None else "" for value in header_values)
        if not include_rows:
            return header, None
        rows: list[dict[str, object]] = []
        for values in worksheet.iter_rows(min_row=header_row + 1, values_only=True):
            if not any(value is not None and str(value).strip() for value in values):
                continue
            rows.append({
                header[position]: value
                for position, value in enumerate(values)
                if position < len(header)
            })
        return header, rows
    finally:
        workbook.close()


def _missing_mapped_columns(header: tuple[str, ...], specification: dict[str, Any]) -> list[str]:
    available = set(header)
    missing = [] if specification["time_column"] in available else ["time_s"]
    missing.extend(
        canonical
        for canonical, source in specification["column_map"].items()
        if source not in available
    )
    return missing


def _source_input_digest(input_data: Path, specifications: list[dict[str, Any]]) -> str:
    """Hash the actual source bytes without putting their private names in a receipt."""

    if input_data.is_file():
        return _sha256(input_data)
    digest = hashlib.sha256()
    for specification in specifications:
        digest.update(_sha256(_source_path(input_data, specification)).encode("ascii"))
    return digest.hexdigest()


def preflight_bundle(
    input_data: Path,
    mapping_path: Path,
    attestation_path: Path,
    protocol_path: Path,
) -> dict[str, Any]:
    """Confirm a multi-sheet mapping without reading measurement rows."""

    input_data = _outside_repository(input_data, field="input data")
    mapping_path = _outside_repository(mapping_path, field="mapping")
    attestation_path = _outside_repository(attestation_path, field="attestation")
    if not input_data.is_dir() and input_data.suffix.casefold() not in {".csv", ".xlsx", ".xlsm"}:
        raise ValueError("multi-source controlled intake requires a CSV/XLSX/XLSM file or directory")
    mapping = _json(mapping_path)
    attestation = _json(attestation_path)
    _require_attestation(attestation)
    specifications, _, _ = _specifications(mapping)
    protocol = _json(protocol_path)
    if protocol.get("status") != "prospective_intake_contract":
        raise ValueError("intake protocol must be prospective")

    missing: set[str] = set()
    for specification in specifications:
        header, _ = _read_source(input_data, specification, include_rows=False)
        missing.update(_missing_mapped_columns(header, specification))
    return {
        "schema_version": 1,
        "artifact_type": "controlled_deidentified_hrs_multisource_preflight",
        "ready_for_controlled_export": not missing,
        "source_format": "directory" if input_data.is_dir() else input_data.suffix.casefold().lstrip("."),
        "source_table_count": len(specifications),
        "alignment_method": "nearest_observation",
        "required_canonical_channels": list(OUTPUT_COLUMNS),
        "missing_canonical_channels": sorted(missing),
        "source_headers_exposed": False,
        "source_rows_read": False,
        "mapping_sha256": _sha256(mapping_path),
        "attestation_sha256": _sha256(attestation_path),
        "protocol_sha256": _sha256(protocol_path),
        "claim_boundary": (
            "This is a controlled schema-and-attestation preflight only. It does not "
            "read outcome rows, confirm source synchronization, validate the model, "
            "or establish safety or IJHE readiness."
        ),
    }


def _timed_rows(
    rows: list[dict[str, object]], specification: dict[str, Any], *, source_index: int,
) -> list[tuple[float, dict[str, object]]]:
    timed: list[tuple[float, dict[str, object]]] = []
    previous: float | None = None
    time_mapping = {"time_format": specification.get("time_format")}
    for row_number, row in enumerate(rows, start=specification["header_row"] + 1):
        value = _relative_time(
            row.get(specification["time_column"]), time_mapping, row_number=row_number,
        )
        if previous is not None and value <= previous:
            raise ValueError(f"source {source_index + 1}: time values must be strictly increasing")
        previous = value
        timed.append((value, row))
    if len(timed) < 20:
        raise ValueError(f"source {source_index + 1}: fewer than 20 usable rows")
    return timed


def _nearest_row(
    timed: list[tuple[float, dict[str, object]]], target: float, maximum_offset: float,
) -> tuple[dict[str, object], float]:
    times = [entry[0] for entry in timed]
    position = bisect_left(times, target)
    candidates = [entry for entry in (position - 1, position) if 0 <= entry < len(timed)]
    best = min(candidates, key=lambda entry: abs(times[entry] - target))
    offset = abs(times[best] - target)
    if offset > maximum_offset:
        raise ValueError("a mapped source cannot be aligned within the declared time tolerance")
    return timed[best][1], offset


def _quantile(values: list[float], fraction: float) -> float:
    """Return a deterministic linear quantile without retaining source rows."""

    if not values:
        return 0.0
    ordered = sorted(values)
    position = (len(ordered) - 1) * fraction
    lower = int(math.floor(position))
    upper = int(math.ceil(position))
    if lower == upper:
        return ordered[lower]
    return ordered[lower] + (ordered[upper] - ordered[lower]) * (position - lower)


def _alignment_diagnostics(
    offsets_by_source: dict[int, list[float]],
    *,
    anchor_index: int,
    maximum_offset: float,
) -> list[dict[str, object]]:
    """Summarize alignment error by anonymous source index only.

    The receipt contains neither source names nor original timestamps.  These
    aggregate offsets make the common-time-basis attestation reviewable before
    an exported trace is evaluated by the simulator.
    """

    threshold = maximum_offset * 0.8
    return [
        {
            "source_index": index + 1,
            "anchor_source": index == anchor_index,
            "matched_row_count": len(offsets),
            "minimum_offset_s": _quantile(offsets, 0.0),
            "median_offset_s": _quantile(offsets, 0.5),
            "p95_offset_s": _quantile(offsets, 0.95),
            "maximum_offset_s": _quantile(offsets, 1.0),
            "matches_at_or_above_80pct_of_tolerance": sum(
                offset >= threshold for offset in offsets
            ),
        }
        for index, offsets in sorted(offsets_by_source.items())
    ]


def _declaration() -> dict[str, Any]:
    return {
        "schema_version": 1,
        "declaration_type": "external_hrs_bundle_metadata",
        "outcomes_accessed_before_freeze": False,
        "controlled_access_only": True,
        "source_identifiers_published": False,
        "absolute_timestamps_published": False,
        "original_column_names_published": False,
        "channels": {
            "common_time_base": {"present": True},
            "vehicle_or_receptacle_pressure": {"present": True, "unit": "MPa_abs"},
            "gas_or_tank_temperature": {"present": True, "unit": "degC"},
            "mass_flow_or_transferred_mass": {"present": True, "unit": "g/s"},
        },
        "metadata": {key: "declared" for key in REQUIRED_METADATA},
        "claim_boundary": (
            "A controlled, de-identified multi-source export that has passed channel and "
            "time-base screens is an evaluator input only. It is not an independently "
            "public holdout, safety certification, or an IJHE completion result."
        ),
    }


def export_bundle(
    input_data: Path,
    mapping_path: Path,
    attestation_path: Path,
    output_directory: Path,
    protocol_path: Path,
) -> dict[str, Any]:
    """Create one de-identified, nearest-time aligned multi-source bundle."""

    input_data = _outside_repository(input_data, field="input data")
    mapping_path = _outside_repository(mapping_path, field="mapping")
    attestation_path = _outside_repository(attestation_path, field="attestation")
    output_directory = _outside_repository(output_directory, field="output directory")
    if not input_data.is_dir() and input_data.suffix.casefold() not in {".csv", ".xlsx", ".xlsm"}:
        raise ValueError("multi-source controlled intake requires a CSV/XLSX/XLSM file or directory")
    if output_directory.exists() and any(output_directory.iterdir()):
        raise ValueError("output directory must be empty to avoid mixing controlled data bundles")

    mapping = _json(mapping_path)
    attestation = _json(attestation_path)
    _require_attestation(attestation)
    specifications, anchor_index, maximum_offset = _specifications(mapping)
    protocol = _json(protocol_path)
    if protocol.get("status") != "prospective_intake_contract":
        raise ValueError("intake protocol must be prospective")

    source_rows: list[list[tuple[float, dict[str, object]]]] = []
    channel_source: dict[str, int] = {}
    for index, specification in enumerate(specifications):
        header, rows = _read_source(input_data, specification, include_rows=True)
        missing = _missing_mapped_columns(header, specification)
        if missing:
            raise ValueError("a source is missing mapped canonical channels")
        assert rows is not None
        source_rows.append(_timed_rows(rows, specification, source_index=index))
        for canonical in specification["column_map"]:
            channel_source[canonical] = index

    anchor_rows = source_rows[anchor_index]
    origin = anchor_rows[0][0]
    normalized: list[dict[str, str]] = []
    offsets_by_source: dict[int, list[float]] = {
        index: [] for index in range(len(source_rows))
    }
    for row_number, (anchor_time, _) in enumerate(anchor_rows, start=1):
        joined = {"time_s": f"{anchor_time - origin:.9g}"}
        aligned_rows: dict[int, dict[str, object]] = {}
        for source_index, timed in enumerate(source_rows):
            source_row, offset = _nearest_row(timed, anchor_time, maximum_offset)
            aligned_rows[source_index] = source_row
            offsets_by_source[source_index].append(offset)
        for canonical in OUTPUT_COLUMNS[1:]:
            source_index = channel_source[canonical]
            source_row = aligned_rows[source_index]
            raw = source_row.get(specifications[source_index]["column_map"][canonical])
            if canonical in NUMERIC_COLUMNS:
                joined[canonical] = f"{_finite(raw, column=canonical, row_number=row_number):.9g}"
            else:
                value = raw.strip() if isinstance(raw, str) else str(raw or "").strip()
                if not value:
                    raise ValueError(f"row {row_number}: {canonical} is empty")
                joined[canonical] = value
        normalized.append(joined)

    output_directory.mkdir(parents=True, exist_ok=True)
    trace_path = output_directory / "full_loop_event.csv"
    with trace_path.open("w", encoding="utf-8", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=OUTPUT_COLUMNS)
        writer.writeheader()
        writer.writerows(normalized)
    declaration = _declaration()
    declaration_path = output_directory / "declaration.json"
    declaration_path.write_text(json.dumps(declaration, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    manifest = build_manifest(output_directory, protocol_path)
    manifest_path = output_directory / "manifest.json"
    manifest_path.write_text(json.dumps(manifest, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    intake = validate_manifest(manifest_path, declaration_path, protocol_path, bundle_root=output_directory)
    screen = validate_full_loop_trace(trace_path, manifest_path, declaration_path, protocol_path, bundle_root=output_directory)
    if not screen.get("full_loop_trace_ready"):
        raise RuntimeError("export did not pass full-loop quality screen: " + "; ".join(screen.get("reasons") or []))
    receipt = {
        "schema_version": 1,
        "artifact_type": "controlled_deidentified_hrs_multisource_export_receipt",
        "created_at_utc": datetime.now(timezone.utc).isoformat(),
        "controlled_access_only": True,
        "source_identifiers_published": False,
        "absolute_timestamps_published": False,
        "original_column_names_published": False,
        "raw_rows_persisted_in_repository": False,
        "source_table_count": len(specifications),
        "alignment_method": "nearest_observation",
        "maximum_alignment_offset_s": maximum_offset,
        "alignment_diagnostics": _alignment_diagnostics(
            offsets_by_source,
            anchor_index=anchor_index,
            maximum_offset=maximum_offset,
        ),
        "output_file_count": 4,
        "rows": len(normalized),
        "time_duration_s": float(normalized[-1]["time_s"]),
        "generic_output_columns": list(OUTPUT_COLUMNS),
        "source_input_sha256": _source_input_digest(input_data, specifications),
        "mapping_sha256": _sha256(mapping_path),
        "attestation_sha256": _sha256(attestation_path),
        "protocol_sha256": _sha256(protocol_path),
        "trace_sha256": _sha256(trace_path),
        "manifest_sha256": _sha256(manifest_path),
        "declaration_sha256": _sha256(declaration_path),
        "intake_decision": intake.get("decision"),
        "quality_decision": screen.get("decision"),
        "full_loop_trace_ready": screen.get("full_loop_trace_ready"),
        "claim_boundary": declaration["claim_boundary"],
    }
    (output_directory / "receipt.json").write_text(
        json.dumps(receipt, ensure_ascii=False, indent=2) + "\n", encoding="utf-8",
    )
    return receipt


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument(
        "--input", type=Path, required=True,
        help="authorised CSV/XLSX/XLSM source, or a controlled directory named by the mapping",
    )
    parser.add_argument("--mapping", type=Path, required=True, help="Private multi-source mapping outside this repository")
    parser.add_argument("--attestation", type=Path, required=True, help="Private unit/state attestation outside this repository")
    parser.add_argument("--output-directory", type=Path, help="Controlled output outside this repository")
    parser.add_argument("--protocol", type=Path, default=ROOT / "research/external_hrs_intake_protocol.json")
    parser.add_argument("--preflight", action="store_true", help="Check schema coverage without reading measurement rows")
    args = parser.parse_args()
    if args.preflight:
        print(json.dumps(preflight_bundle(args.input, args.mapping, args.attestation, args.protocol), ensure_ascii=False, indent=2))
        return 0
    if args.output_directory is None:
        parser.error("--output-directory is required unless --preflight is used")
    print(json.dumps(export_bundle(args.input, args.mapping, args.attestation, args.output_directory, args.protocol), ensure_ascii=False, indent=2))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
