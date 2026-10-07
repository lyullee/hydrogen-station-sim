"""Run the claim-bounded Jankuj 2026 pressure-decay diagnostic.

The repository already contained a numerical intake record for this workbook
before the transfer protocol was committed.  The attempted prospective protocol
is therefore invalidated.  Curve 1 still identifies one effective breach
diameter and curve 2 is evaluated without refitting, but the result is retained
only as a post-access diagnostic.  All finite post-onset samples remain.
"""
from __future__ import annotations

import argparse
import hashlib
import json
import math
import sys
from datetime import datetime, timezone
from pathlib import Path

import numpy as np
from openpyxl import load_workbook
from scipy.optimize import minimize_scalar

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "src"))

from h2station.schefer_2007_validation import (  # noqa: E402
    Schefer2007Inputs,
    simulate_adiabatic_blowdown_pressure,
)


PROTOCOL = ROOT / "research/jankuj_2026_pressure_decay_protocol_2026_10_08.json"
PRIOR_ACCESS_RECORD = ROOT / (
    "research/accidental_self_ignition_public_evidence_2026_10_04.json"
)
WORKBOOK = ROOT / (
    "data/public_validation/raw/jankuj_2026/extracted/"
    "Figure 7. Pressure drop during hydrogen release.xlsx"
)
ARCHIVE = ROOT / (
    "data/public_validation/raw/jankuj_2026/"
    "fig_7-9_Accidental self-ignition of hydrogen_VJ.7z"
)
OUT_JSON = ROOT / "research/jankuj_2026_pressure_decay_transfer_result_2026_10_08.json"
OUT_MD = ROOT / "research/JANKUJ_2026_PRESSURE_DECAY_TRANSFER_RESULT_2026_10_08.md"
PROTOCOL_COMMIT = "b8b7ac542abab17fa2bdebeca83ce584bebb8831"
BAR_PER_PSI = 0.06894757293168


def _sha256(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def _md5(path: Path) -> str:
    return hashlib.md5(path.read_bytes()).hexdigest()  # noqa: S324 - publisher checksum


def _finite(value: float | None) -> float | None:
    return float(value) if value is not None and math.isfinite(float(value)) else None


def _crossing(time_s: np.ndarray, values: np.ndarray, target: float) -> float | None:
    reached = np.flatnonzero(values <= target)
    if not reached.size:
        return None
    index = int(reached[0])
    if index == 0:
        return float(time_s[0])
    t0, t1 = float(time_s[index - 1]), float(time_s[index])
    v0, v1 = float(values[index - 1]), float(values[index])
    if v1 == v0:
        return t1
    return t0 + (target - v0) * (t1 - t0) / (v1 - v0)


def _load_curves(path: Path) -> dict[str, dict[str, object]]:
    workbook = load_workbook(path, read_only=True, data_only=True)
    sheet = workbook.active
    headers = [cell.value for cell in next(sheet.iter_rows(min_row=1, max_row=1))]
    expected = ["Time [s]", "Pressure n.1 [bar]", "Pressure n.2 [bar]"]
    if headers != expected:
        raise ValueError(f"unexpected Figure 7 headers: {headers!r}")
    rows = list(sheet.iter_rows(min_row=2, values_only=True))
    curves: dict[str, dict[str, object]] = {}
    for column, label in ((1, expected[1]), (2, expected[2])):
        pairs = [
            (float(row[0]), float(row[column]))
            for row in rows
            if row[0] is not None
            and row[column] is not None
            and math.isfinite(float(row[0]))
            and math.isfinite(float(row[column]))
        ]
        time_s = np.asarray([item[0] for item in pairs], dtype=float)
        gauge_bar = np.asarray([item[1] for item in pairs], dtype=float)
        if len(time_s) < 20 or np.any(np.diff(time_s) <= 0.0):
            raise ValueError(f"{label} fails the frozen time-series eligibility rule")
        initial_gauge_bar = float(np.median(gauge_bar[:3]))
        onset = next(
            (
                index
                for index in range(len(gauge_bar) - 2)
                if gauge_bar[index] <= 0.99 * initial_gauge_bar
                and gauge_bar[index + 1] <= gauge_bar[index]
                and gauge_bar[index + 2] <= gauge_bar[index + 1]
            ),
            None,
        )
        if onset is None:
            raise ValueError(f"{label} has no release onset under the frozen rule")
        selected_time = time_s[onset:] - time_s[onset]
        selected_abs_bar = gauge_bar[onset:] + 0.970
        curves[label] = {
            "time_s": selected_time,
            "pressure_abs_bar": selected_abs_bar,
            "initial_pressure_abs_bar": initial_gauge_bar + 0.970,
            "onset_source_time_s": float(time_s[onset]),
            "source_sample_count": len(time_s),
            "evaluation_sample_count": len(selected_time),
            "median_sample_interval_s": float(np.median(np.diff(time_s))),
            "pressure_drop_fraction": float(
                (initial_gauge_bar - np.min(gauge_bar)) / initial_gauge_bar
            ),
            "large_positive_jump_count": int(np.sum(np.diff(selected_abs_bar) >= 50.0)),
            "nonphysical_absolute_pressure_count": int(np.sum(selected_abs_bar <= 0.0)),
            "minimum_reported_gauge_pressure_bar": float(np.min(gauge_bar)),
        }
    return curves


def _predict(time_s: np.ndarray, initial_abs_bar: float, diameter_mm: float) -> np.ndarray:
    pressure_psi = simulate_adiabatic_blowdown_pressure(
        time_s,
        inputs=Schefer2007Inputs(
            initial_pressure_pa_abs=initial_abs_bar * 1.0e5,
            initial_temperature_k=276.25,
            tank_volume_m3=0.05,
            restriction_diameter_m=diameter_mm / 1000.0,
            discharge_coefficient=1.0,
            ambient_pressure_pa=97_000.0,
        ),
    )
    return pressure_psi * BAR_PER_PSI


def _metrics(curve: dict[str, object], predicted_abs_bar: np.ndarray) -> dict[str, object]:
    time_s = np.asarray(curve["time_s"], dtype=float)
    measured = np.asarray(curve["pressure_abs_bar"], dtype=float)
    initial = float(curve["initial_pressure_abs_bar"])
    initial_excess = initial - 0.970
    error = predicted_abs_bar - measured
    nrmse = float(np.sqrt(np.mean(error**2)) / initial_excess * 100.0)
    positive = measured > 0.0
    median_ape = float(np.median(np.abs(error[positive]) / measured[positive]) * 100.0)
    half_target = 0.970 + 0.5 * initial_excess
    measured_half = _crossing(time_s, measured, half_target)
    predicted_half = _crossing(time_s, predicted_abs_bar, half_target)
    if measured_half is not None and measured_half > 0.0 and predicted_half is not None:
        half_error = abs(predicted_half - measured_half) / measured_half * 100.0
    else:
        half_error = None
    return {
        "sample_count": len(time_s),
        "pressure_nrmse_percent_initial_excess": nrmse,
        "median_absolute_percentage_error_percent": median_ape,
        "measured_half_pressure_time_s": measured_half,
        "predicted_half_pressure_time_s": predicted_half,
        "half_pressure_time_relative_error_percent": half_error,
        "pressure_nrmse_screen_pass": nrmse <= 10.0,
        "median_ape_screen_pass": median_ape <= 15.0,
        "half_pressure_time_screen_pass": half_error is not None and half_error <= 25.0,
        "minimum_measured_absolute_pressure_bar": float(np.min(measured)),
        "minimum_predicted_absolute_pressure_bar": float(np.min(predicted_abs_bar)),
        "maximum_absolute_error_bar": float(np.max(np.abs(error))),
    }


def run() -> dict[str, object]:
    protocol = json.loads(PROTOCOL.read_text(encoding="utf-8"))
    prior_access = json.loads(PRIOR_ACCESS_RECORD.read_text(encoding="utf-8"))
    if protocol.get("status") != "FROZEN_BEFORE_RAW_OUTCOME_ACCESS":
        raise ValueError("Jankuj protocol is not frozen")
    prior_pressure_file = next(
        (
            item
            for item in prior_access.get("files", [])
            if item.get("name") == "Figure 7. Pressure drop during hydrogen release.xlsx"
        ),
        None,
    )
    if not prior_pressure_file or prior_pressure_file.get("numeric_rows", 0) <= 0:
        raise ValueError("Jankuj prior numerical-access record is missing")
    if _md5(ARCHIVE) != protocol["source"]["archive_md5"]:
        raise ValueError("Jankuj archive checksum mismatch")
    curves = _load_curves(WORKBOOK)
    development = curves["Pressure n.1 [bar]"]
    holdout = curves["Pressure n.2 [bar]"]

    def objective(log_diameter_mm: float) -> float:
        diameter_mm = float(np.exp(log_diameter_mm))
        predicted = _predict(
            np.asarray(development["time_s"], dtype=float),
            float(development["initial_pressure_abs_bar"]),
            diameter_mm,
        )
        measured = np.asarray(development["pressure_abs_bar"], dtype=float)
        return float(np.sqrt(np.mean((predicted - measured) ** 2)))

    fit = minimize_scalar(
        objective,
        bounds=(math.log(0.5), math.log(20.0)),
        method="bounded",
        options={"xatol": 1.0e-5},
    )
    if not fit.success:
        raise RuntimeError(f"effective breach fit failed: {fit.message}")
    diameter_mm = float(np.exp(fit.x))
    development_prediction = _predict(
        np.asarray(development["time_s"], dtype=float),
        float(development["initial_pressure_abs_bar"]),
        diameter_mm,
    )
    holdout_prediction = _predict(
        np.asarray(holdout["time_s"], dtype=float),
        float(holdout["initial_pressure_abs_bar"]),
        diameter_mm,
    )
    development_metrics = _metrics(development, development_prediction)
    holdout_metrics = _metrics(holdout, holdout_prediction)
    joint = bool(
        holdout_metrics["pressure_nrmse_screen_pass"]
        and holdout_metrics["median_ape_screen_pass"]
        and holdout_metrics["half_pressure_time_screen_pass"]
    )
    return {
        "schema_version": 1,
        "artifact_type": "post_access_pressure_decay_transfer_diagnostic",
        "generated_at": datetime.now(timezone.utc).isoformat(),
        "status": "INVALIDATED_PRIOR_OUTCOME_ACCESS",
        "evidence_role": "post_access_diagnostic_only",
        "source": {
            "article_doi": protocol["source"]["article_doi"],
            "dataset_doi": protocol["source"]["dataset_doi"],
            "license": protocol["source"]["license"],
            "archive_md5": _md5(ARCHIVE),
            "workbook_sha256": _sha256(WORKBOOK),
            "raw_rows_committed": False,
        },
        "protocol": {
            "artifact": str(PROTOCOL.relative_to(ROOT)).replace("\\", "/"),
            "sha256": _sha256(PROTOCOL),
            "git_commit_before_current_archive_download": PROTOCOL_COMMIT,
            "protocol_frozen_before_first_numeric_outcome_access": False,
            "prospective_protocol_valid": False,
            "invalidation_reason": (
                "research/accidental_self_ignition_public_evidence_2026_10_04.json "
                "already recorded numeric Figure 7 rows and the workbook hash"
            ),
            "split": protocol["fixed_split"],
            "primary_screens": protocol["frozen_primary_screens"],
        },
        "prior_outcome_access": {
            "artifact": str(PRIOR_ACCESS_RECORD.relative_to(ROOT)).replace("\\", "/"),
            "recorded_at": prior_access.get("recorded_at"),
            "status": prior_access.get("status"),
            "pressure_workbook_numeric_rows": prior_pressure_file.get("numeric_rows"),
            "pressure_workbook_sha256": prior_pressure_file.get("sha256"),
            "matches_current_workbook": (
                prior_pressure_file.get("sha256") == _sha256(WORKBOOK)
            ),
        },
        "model": {
            "source": "src/h2station/schefer_2007_validation.py",
            "source_sha256": _sha256(ROOT / "src/h2station/schefer_2007_validation.py"),
            "name": protocol["fixed_model"]["name"],
            "fitted_effective_breach_diameter_mm": diameter_mm,
            "fit_function_evaluations": int(fit.nfev),
            "holdout_refitted": False,
        },
        "development_curve": {
            **{key: value for key, value in development.items() if key not in {"time_s", "pressure_abs_bar"}},
            "metrics": development_metrics,
        },
        "holdout_curve": {
            **{key: value for key, value in holdout.items() if key not in {"time_s", "pressure_abs_bar"}},
            "metrics": holdout_metrics,
        },
        "joint_primary_screen_pass": joint,
        "source_depletion_transfer_validation_supported": False,
        "interpretation": (
            "The attempted prospective transfer protocol is invalid because the repository "
            "already recorded numerical workbook outcomes before the protocol commit. The "
            "post-access diagnostic also failed: the development curve contains a large "
            "pressure reset after depressurization, the all-points rule retained it, and the "
            "second curve failed without refitting. No validation claim is made."
        ),
        "claim_boundary": protocol["claim_boundary"],
    }


def _write_markdown(result: dict[str, object], path: Path) -> None:
    dev = result["development_curve"]
    holdout = result["holdout_curve"]
    metrics = holdout["metrics"]
    half = metrics["half_pressure_time_relative_error_percent"]
    half_text = "not reached" if half is None else f"{half:.3f}%"
    text = f"""# Jankuj 2026 pressure-decay transfer result

**Decision: {result['status']}**

The repository had already recorded the Figure 7 numerical row count and
workbook hash on 2026-10-04. The later transfer protocol is therefore not a
valid prospective protocol, even though its commit predates this particular
archive download. This result is retained only as a post-access diagnostic.
One effective breach diameter
({result['model']['fitted_effective_breach_diameter_mm']:.4f} mm) was fitted to
`Pressure n.1`; `Pressure n.2` was evaluated without refitting.

| Holdout metric | Result | Frozen limit | Pass |
|---|---:|---:|---:|
| Pressure NRMSE / initial excess | {metrics['pressure_nrmse_percent_initial_excess']:.3f}% | <= 10% | {metrics['pressure_nrmse_screen_pass']} |
| Median absolute percentage error | {metrics['median_absolute_percentage_error_percent']:.3f}% | <= 15% | {metrics['median_ape_screen_pass']} |
| Half-pressure time relative error | {half_text} | <= 25% | {metrics['half_pressure_time_screen_pass']} |

The development curve contains {dev['large_positive_jump_count']} pressure jump
of at least 50 bar after the detected release onset. The protocol required all
finite post-onset samples to remain, so this discontinuity was not trimmed. The
holdout also contains {holdout['nonphysical_absolute_pressure_count']} samples
that become non-positive after the frozen gauge-to-absolute conversion. These
observations make the public columns unsuitable for the assumed two-repeat
transfer design and the failed result is retained.

This result does not validate breach geometry, ignition probability, flame
length, consequence distance, H70 release, station operation, emergency action
or SAGA effectiveness.
"""
    path.write_text(text, encoding="utf-8", newline="\n")


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--output", type=Path, default=OUT_JSON)
    parser.add_argument("--markdown", type=Path, default=OUT_MD)
    args = parser.parse_args()
    result = run()
    args.output.write_text(
        json.dumps(result, ensure_ascii=False, indent=2) + "\n",
        encoding="utf-8",
        newline="\n",
    )
    _write_markdown(result, args.markdown)
    print(json.dumps({
        "status": result["status"],
        "joint_primary_screen_pass": result["joint_primary_screen_pass"],
        "output": str(args.output),
    }, ensure_ascii=False))


if __name__ == "__main__":
    main()
