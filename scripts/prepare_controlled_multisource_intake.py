"""Prepare a private mapping workbench for one controlled HRS workbook.

The output is intentionally outside the repository.  It contains worksheet
names and original header labels so only an authorised custodian can connect
their actual logger channels to the canonical full-loop contract.  It reads at
most the first 43 rows of each worksheet and retains no measurement values.
Never commit, publish, or attach the generated files to a manuscript.
"""

from __future__ import annotations

import argparse
from datetime import datetime, timezone
import json
from pathlib import Path
from typing import Any

try:  # Supports both ``python scripts/...`` and test-module imports.
    from audit_controlled_data_schema import (
        DATA_LIKENESS_SAMPLE_ROWS,
        HEADER_SEARCH_MAX_ROWS,
        _classify_header,
        _outside_repository,
        _screen_rows,
    )
    from export_confidential_full_loop_bundle import OUTPUT_COLUMNS, ROOT, _sha256
except ModuleNotFoundError:  # pragma: no cover - import style depends on launcher
    from scripts.audit_controlled_data_schema import (
        DATA_LIKENESS_SAMPLE_ROWS,
        HEADER_SEARCH_MAX_ROWS,
        _classify_header,
        _outside_repository,
        _screen_rows,
    )
    from scripts.export_confidential_full_loop_bundle import OUTPUT_COLUMNS, ROOT, _sha256


_REQUIRED_METADATA = (
    "initial_conditions",
    "tank_capacity_or_geometry",
    "protocol_mode_or_pressure_ramp",
    "units_and_sampling_interval",
    "source_cascade_state",
    "compressor_state",
    "precooler_state",
    "stop_abort_fault_markers",
    "quality_and_calibration_metadata",
    "reuse_terms",
    "source_identity",
    "custodian_or_archive",
    "acquired_at_utc",
    "license_or_reuse_reference",
)


def _header_row(rows: tuple[tuple[object, ...], ...], labels: tuple[str, ...]) -> int:
    for number, row in enumerate(rows[:HEADER_SEARCH_MAX_ROWS], start=1):
        if tuple(str(value or "") for value in row) == labels:
            return number
    raise RuntimeError("selected header cannot be located")


def _time_columns(labels: tuple[str, ...]) -> list[str]:
    return [
        label for label in labels
        if "time" in _classify_header((label,))
    ]


def _workbench(input_data: Path) -> dict[str, Any]:
    try:
        from openpyxl import load_workbook
    except ImportError as exc:  # pragma: no cover - environment-specific
        raise RuntimeError("openpyxl is required for controlled Excel intake") from exc

    workbook = load_workbook(input_data, read_only=True, data_only=True)
    try:
        sources: list[dict[str, Any]] = []
        for worksheet in workbook.worksheets:
            rows = tuple(
                worksheet.iter_rows(
                    min_row=1,
                    max_row=HEADER_SEARCH_MAX_ROWS + DATA_LIKENESS_SAMPLE_ROWS,
                    values_only=True,
                )
            )
            screen = _screen_rows(rows)
            if screen is None:
                continue
            labels, _, record_shaped = screen
            semantic_channels = _classify_header(labels)
            if not (record_shaped and "time" in semantic_channels):
                continue
            sources.append({
                "worksheet": worksheet.title,
                "header_row": _header_row(rows, labels),
                "time_column_candidates": _time_columns(labels),
                "original_headers": list(labels),
                "semantic_channels": sorted(semantic_channels),
            })
    finally:
        workbook.close()

    return {
        "schema_version": 1,
        "artifact_type": "controlled_private_multisource_mapping_workbench",
        "publication_prohibited": True,
        "repository_storage_prohibited": True,
        "source_workbook_sha256": _sha256(input_data),
        "sample_data_rows_structurally_inspected_in_memory": True,
        "measurement_values_persisted": False,
        "source_count": len(sources),
        "sources": sources,
        "claim_boundary": (
            "This private workbench is a custodian mapping aid only. It does not "
            "attest units, synchronization, state semantics, a validation event, "
            "model accuracy, safety, or publication readiness."
        ),
    }


def _mapping_template(workbench: dict[str, Any]) -> dict[str, Any]:
    return {
        "schema_version": 1,
        "sources": [
            {
                "worksheet": source["worksheet"],
                "header_row": source["header_row"],
                "time_column": "<select one time_column_candidate>",
                "time_format": "<optional strptime format or null>",
                "column_map": {
                    "<canonical_channel>": "<select an original_header>",
                },
            }
            for source in workbench["sources"]
        ],
        "alignment": {
            "method": "nearest_observation",
            "anchor_source": 0,
            "maximum_offset_s": "<custodian-approved tolerance>",
        },
        "mapping_instructions": (
            "Replace every placeholder. Each canonical output field except time_s must "
            "be assigned exactly once across sources. Do not infer or interpolate a "
            "missing source channel."
        ),
        "canonical_fields_required": list(OUTPUT_COLUMNS),
    }


def _attestation_template() -> dict[str, Any]:
    return {
        "schema_version": 1,
        "authorised_controlled_evaluation": False,
        "outcomes_accessed_before_protocol_freeze": "<true or false>",
        "units": {
            "vehicle_pressure_mpa_abs": "MPa_abs",
            "temperature_degC": "degC",
            "mass_flow_g_s": "g/s",
            "station_pressure_mpa_abs": "MPa_abs",
            "delivered_gas_temperature_degC": "degC",
            "cascade_source_pressure_mpa_abs": "MPa_abs",
        },
        "state_semantics": {
            field: "<custodian-approved meaning and valid states>"
            for field in OUTPUT_COLUMNS[7:]
        },
        "metadata": {
            field: "<custodian-approved declaration>" for field in _REQUIRED_METADATA
        },
        "source_synchronization": {
            "common_time_basis_confirmed": False,
            "alignment_method": "nearest_observation",
        },
        "attestation_instructions": (
            "Set authorization and common_time_basis_confirmed true only after a "
            "custodian reviews the event selection, units, calibration status, state "
            "meanings, and source clocks."
        ),
    }


def prepare_workbench(input_data: Path, output_directory: Path) -> dict[str, Any]:
    """Write private mapping and attestation templates outside the repository."""

    input_data = _outside_repository(input_data, label="input data")
    output_directory = _outside_repository(output_directory, label="output directory")
    if input_data.suffix.casefold() not in {".xlsx", ".xlsm"}:
        raise ValueError("controlled mapping workbench requires XLSX or XLSM input")
    if output_directory.exists() and any(output_directory.iterdir()):
        raise ValueError("output directory must be empty")

    workbench = _workbench(input_data)
    output_directory.mkdir(parents=True, exist_ok=True)
    (output_directory / "private_source_catalog.json").write_text(
        json.dumps(workbench, ensure_ascii=False, indent=2) + "\n", encoding="utf-8"
    )
    (output_directory / "event-mapping.template.json").write_text(
        json.dumps(_mapping_template(workbench), ensure_ascii=False, indent=2) + "\n",
        encoding="utf-8",
    )
    (output_directory / "event-attestation.template.json").write_text(
        json.dumps(_attestation_template(), ensure_ascii=False, indent=2) + "\n",
        encoding="utf-8",
    )
    receipt = {
        "schema_version": 1,
        "artifact_type": "controlled_private_multisource_mapping_workbench_receipt",
        "created_at_utc": datetime.now(timezone.utc).isoformat(),
        "publication_prohibited": True,
        "repository_storage_prohibited": True,
        "measurement_values_persisted": False,
        "source_count": workbench["source_count"],
        "source_workbook_sha256": workbench["source_workbook_sha256"],
        "output_file_count": 3,
        "claim_boundary": workbench["claim_boundary"],
    }
    (output_directory / "receipt.json").write_text(
        json.dumps(receipt, ensure_ascii=False, indent=2) + "\n", encoding="utf-8"
    )
    return receipt


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--input", type=Path, required=True)
    parser.add_argument("--output-directory", type=Path, required=True)
    args = parser.parse_args()
    receipt = prepare_workbench(args.input, args.output_directory)
    print(json.dumps({
        "source_count": receipt["source_count"],
        "output_file_count": receipt["output_file_count"],
        "output_written": True,
    }))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
