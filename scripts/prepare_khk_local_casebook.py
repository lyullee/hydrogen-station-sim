"""Build a local-only KHK hydrogen-station incident casebook.

The KHK/METI workbook is intentionally not redistributed by this repository.
This command reads a separately downloaded local workbook and writes its
casebook and manifest below the gitignored ``data/`` directory.  Only the
pipeline, input digest and rights boundary are suitable for version control.
"""

from __future__ import annotations

import argparse
from datetime import date, datetime
from hashlib import sha256
import json
from pathlib import Path
from typing import Any, Iterable

from openpyxl import load_workbook


DEFAULT_SOURCE = Path("data/public_validation/raw/khk_hpg_incident_db_r7.xlsm")
DEFAULT_OUTPUT = Path("data/public_validation/results/khk_local_casebook/casebook.json")
DEFAULT_MANIFEST = Path("data/public_validation/results/khk_local_casebook/manifest.json")
SHEET = "高圧ガス保安法事故"

HEADERS = (
    "事故ｺｰﾄﾞ", "事故区分", "事故分類", "事故名称", "事故発生日", "県名",
    "死者", "重傷", "軽傷", "計", "物質名", "1次事象", "2次事象",
    "噴出・漏えいの程度", "噴出・漏えいの部位", "噴出・漏えい部位の寸法（径）",
    "噴出・漏えい部位の寸法（板厚）", "噴出・漏えい部位の寸法（呼び圧力）",
    "噴出・漏えいの分類", "業種", "設備区分", "取扱状態", "事故原因(主因)",
    "事故原因(副因)", "着火源", "事故概要", "事業所で講じた措置及び対策",
    "高圧ガス事故概要報告リンク",
)

# These terms deliberately select the HRS boundary rather than all industrial
# hydrogen incidents.  The source material remains local and this mapping is
# a candidate-family selector, not an expert judgement about causation.
HRS_TERMS = (
    "水素ステーション", "圧縮水素スタンド", "スタンド", "充填所",
    "ディスペンサ", "ディスペンサー", "充填ノズル", "充填ホース",
    "水素トレーラ", "水素カードル", "カードル", "FCV", "燃料電池車",
)


def _text(value: Any) -> str:
    return "" if value is None else str(value).strip()


def _date(value: Any) -> str | None:
    if isinstance(value, (datetime, date)):
        return value.date().isoformat() if isinstance(value, datetime) else value.isoformat()
    value = _text(value)
    return value or None


def sha256_file(path: Path) -> str:
    return sha256(path.read_bytes()).hexdigest()


def is_station_case(row: dict[str, Any]) -> bool:
    """Return whether a local row belongs to the HRS boundary candidate set."""

    material = _text(row.get("物質名"))
    if "水素" not in material:
        return False
    searchable = " ".join(
        _text(row.get(key))
        for key in ("事故名称", "業種", "設備区分", "事故概要", "取扱状態")
    )
    return any(term in searchable for term in HRS_TERMS)


def response_plan_candidates(row: dict[str, Any]) -> list[str]:
    """Map structured incident signals to existing response families.

    Multiple families are retained because a real incident can combine a
    release, a failed isolation and a fueling or trailer boundary condition.
    """

    title = _text(row.get("事故名称"))
    event = _text(row.get("1次事象")) + " " + _text(row.get("2次事象"))
    equipment = _text(row.get("設備区分"))
    text = " ".join((title, event, equipment, _text(row.get("事故概要"))))
    plans: list[str] = []

    if any(token in text for token in ("火災", "爆発", "着火")):
        plans.append("hydrogen_fire")
    if any(token in text for token in ("漏洩", "漏えい", "噴出", "ガス漏")):
        plans.append("gas_release")
    if "安全弁" in text or "リリーフ" in text:
        plans.append("relief_discharge")
    if any(token in text for token in ("圧縮機", "コンプレッサ")):
        plans.append("compressor_thermal")
    if any(token in text for token in ("ホース", "ノズル", "継手", "配管", "バルブ")):
        plans.append("hose_connection")
    if any(token in text for token in ("トレーラ", "カードル", "供給", "移動")):
        plans.append("supply_connection")
    if any(token in text for token in ("充填", "ディスペンサ", "ディスペンサー", "FCV", "燃料電池車")):
        plans.append("fueling_fault")
    if any(token in text for token in ("圧力上昇", "過圧", "閉め切り")):
        plans.append("overpressure")
    # Preserve deterministic order while removing duplicate families.
    return list(dict.fromkeys(plans or ["gas_release"]))


def _rows(source: Path) -> Iterable[dict[str, Any]]:
    workbook = load_workbook(source, read_only=True, data_only=True, keep_vba=True)
    if SHEET not in workbook.sheetnames:
        raise ValueError(f"KHK workbook does not contain {SHEET!r}")
    worksheet = workbook[SHEET]
    header_row = next(worksheet.iter_rows(values_only=True))
    headers = [_text(value) for value in header_row]
    if headers[: len(HEADERS)] != list(HEADERS):
        raise ValueError("KHK header changed; stop rather than silently selecting wrong fields")
    for values in worksheet.iter_rows(min_row=2, values_only=True):
        row = {
            headers[index]: values[index] if index < len(values) else None
            for index in range(min(len(headers), len(values)))
            if headers[index]
        }
        if is_station_case(row):
            yield row


def build_casebook(source: Path, *, include_descriptions: bool = True) -> dict[str, Any]:
    cases: list[dict[str, Any]] = []
    for row in _rows(source):
        case = {
            "event_id": _text(row.get("事故ｺｰﾄﾞ")),
            "event_date": _date(row.get("事故発生日")),
            "event_class": _text(row.get("事故分類")),
            "title": _text(row.get("事故名称")),
            "primary_event": _text(row.get("1次事象")),
            "secondary_event": _text(row.get("2次事象")),
            "material": _text(row.get("物質名")),
            "industry": _text(row.get("業種")),
            "equipment": _text(row.get("設備区分")),
            "handling_state": _text(row.get("取扱状態")),
            "response_plan_candidates": response_plan_candidates(row),
        }
        if include_descriptions:
            case["incident_description"] = _text(row.get("事故概要"))
            case["operator_measures"] = _text(row.get("事業所で講じた措置及び対策"))
        cases.append(case)
    return {
        "schema_version": 1,
        "source_rights": "KHK/METI local-only; do not transfer, distribute, publish, mirror or commit this casebook",
        "selection": {
            "sheet": SHEET,
            "material_contains": "水素",
            "station_boundary_terms": list(HRS_TERMS),
            "response_text_included": include_descriptions,
            "coordinator_approval": False,
        },
        "cases": cases,
        "claim_boundary": "Candidate scenario-family mapping from local KHK accident records. It is not an expert-approved holdout, probability estimate, or proof of response effectiveness.",
    }


def build_manifest(source: Path, output: Path, *, include_descriptions: bool, casebook: dict[str, Any]) -> dict[str, Any]:
    return {
        "schema_version": 1,
        "generated_at": datetime.now().astimezone().isoformat(),
        "access_status": "LOCAL_ARCHIVE_READABLE",
        "source": {
            "path": source.as_posix(),
            "sha256": sha256_file(source),
            "not_committed": True,
        },
        "output": {
            "path": output.as_posix(),
            "not_committed": True,
            "response_text_included": include_descriptions,
        },
        "casebook_sha256": sha256_file(output),
        "case_count": len(casebook["cases"]),
        "rights_and_restrictions": {
            "public_posting_or_public_download_mirror_prohibited": True,
            "commercial_permission_required": True,
            "raw_or_derived_casebook_may_be_committed": False,
        },
        "validation_classification": {
            "actual_accident_data_grounding_candidate": True,
            "expert_holdout_ready": False,
            "full_loop_station_vehicle_holdout_eligible": False,
            "goal_completion_permitted": False,
        },
        "next_action": "Obtain written permission before any de-identified aggregate or casebook is shared; obtain coordinator and expert-review approval before using responses as an effectiveness holdout.",
    }


def main() -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument("--source", type=Path, default=DEFAULT_SOURCE)
    parser.add_argument("--output", type=Path, default=DEFAULT_OUTPUT)
    parser.add_argument("--manifest", type=Path, default=DEFAULT_MANIFEST)
    parser.add_argument("--omit-descriptions", action="store_true")
    args = parser.parse_args()
    if not args.source.exists():
        raise SystemExit(f"local KHK workbook not found: {args.source}")
    args.output.parent.mkdir(parents=True, exist_ok=True)
    args.manifest.parent.mkdir(parents=True, exist_ok=True)
    include_descriptions = not args.omit_descriptions
    casebook = build_casebook(args.source, include_descriptions=include_descriptions)
    args.output.write_text(json.dumps(casebook, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    manifest = build_manifest(args.source, args.output, include_descriptions=include_descriptions, casebook=casebook)
    args.manifest.write_text(json.dumps(manifest, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    print(json.dumps({"case_count": manifest["case_count"], "casebook_sha256": manifest["casebook_sha256"]}, ensure_ascii=False))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
