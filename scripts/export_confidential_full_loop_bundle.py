"""Export one authorised HRS event as a de-identified full-loop bundle.

This is a *local, controlled-access* bridge for owner-held CSV, XLSX, or XLSM
logs.  It never writes the original header names, filenames, source identity,
or absolute time to its output.  The export directory is deliberately required
to be outside the Git worktree, so a numerical trace cannot be added to the
public project by mistake.

The exporter is not a model evaluator.  It only makes an authorised trace
ready for the existing hash, quality, and full-loop channel screens.  A
successful export therefore does not establish accuracy, safety, or IJHE
readiness; it makes those later evaluations reproducible when a genuinely
eligible event is available.
"""

from __future__ import annotations

import argparse
import csv
from datetime import datetime, timezone
import hashlib
import json
import math
from pathlib import Path
from typing import Any

from intake_external_hrs_bundle import build_manifest
from validate_external_hrs_full_loop import validate_full_loop_trace
from validate_external_hrs_manifest import validate as validate_manifest


ROOT = Path(__file__).resolve().parents[1]

NUMERIC_COLUMNS = (
    "vehicle_pressure_mpa_abs",
    "temperature_degC",
    "mass_flow_g_s",
    "station_pressure_mpa_abs",
    "delivered_gas_temperature_degC",
    "cascade_source_pressure_mpa_abs",
)
STATE_COLUMNS = (
    "cascade_selected_bank",
    "compressor_state",
    "precooler_state",
    "leak_check_state",
    "vent_state",
    "fault_state",
    "esd_state",
)
OUTPUT_COLUMNS = ("time_s",) + NUMERIC_COLUMNS + STATE_COLUMNS
REQUIRED_METADATA = (
    "initial_conditions",
    "tank_capacity_or_geometry",
    "protocol_mode_or_pressure_ramp",
    "units_and_sampling_interval",
    "source_cascade_state",
    "compressor_state",
    "precooler_state",
    "stop_abort_fault_markers",
    "quality_and_calibration_metadata",
    "reuse_terms",
    "source_identity",
    "custodian_or_archive",
    "acquired_at_utc",
    "license_or_reuse_reference",
)

VEHICLE_TEMPERATURE_OBSERVATION_OPERATORS = frozenset({
    "gas_temperature",
    "liner_temperature",
    "shell_temperature",
    "sensor_weighted_tank_temperature",
})
DELIVERED_TEMPERATURE_OBSERVATION_OPERATOR = "delivered_gas_temperature"


def _json(path: Path) -> dict[str, Any]:
    value = json.loads(path.read_text(encoding="utf-8"))
    if not isinstance(value, dict):
        raise ValueError(f"{path} must contain a JSON object")
    return value


def _sha256(path: Path) -> str:
    """Hash a controlled file without retaining any source content."""

    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for block in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(block)
    return digest.hexdigest()


def _within(path: Path, root: Path) -> bool:
    try:
        path.resolve().relative_to(root.resolve())
    except ValueError:
        return False
    return True


def _outside_repository(path: Path, *, field: str) -> Path:
    resolved = path.resolve()
    if _within(resolved, ROOT):
        raise ValueError(f"{field} must be outside the repository worktree")
    return resolved


def _finite(value: object | None, *, column: str, row_number: int) -> float:
    try:
        number = float(value.strip()) if isinstance(value, str) else float(value)
    except (TypeError, ValueError) as exc:
        raise ValueError(f"row {row_number}: {column} is not numeric") from exc
    if not math.isfinite(number):
        raise ValueError(f"row {row_number}: {column} is not finite")
    return number


def _relative_time(value: object | None, mapping: dict[str, Any], *, row_number: int) -> float:
    if isinstance(value, datetime):
        timestamp = value
        if timestamp.tzinfo is None:
            timestamp = timestamp.replace(tzinfo=timezone.utc)
        return timestamp.timestamp()
    raw = value.strip() if isinstance(value, str) else "" if value is None else str(value).strip()
    if not raw:
        raise ValueError(f"row {row_number}: time value is empty")
    time_format = mapping.get("time_format")
    if time_format:
        try:
            return datetime.strptime(raw, str(time_format)).replace(tzinfo=timezone.utc).timestamp()
        except ValueError as exc:
            raise ValueError(f"row {row_number}: time does not match time_format") from exc
    try:
        return _finite(raw, column="time", row_number=row_number)
    except ValueError as numeric_error:
        try:
            return datetime.fromisoformat(raw.replace("Z", "+00:00")).timestamp()
        except ValueError:
            raise numeric_error


def _temperature_observation_semantics(attestation: dict[str, Any]) -> dict[str, str]:
    """Validate private thermal-measurement meaning and return safe categories.

    The free-text location and calibration declarations remain in the private
    attestation.  Only a coarse, non-identifying observation operator is
    eligible for the de-identified evaluator declaration.
    """

    semantics = attestation.get("temperature_observation")
    if not isinstance(semantics, dict):
        raise ValueError("attestation must declare temperature observation semantics")
    vehicle = semantics.get("vehicle_temperature_degC")
    delivered = semantics.get("delivered_gas_temperature_degC")
    if not isinstance(vehicle, dict) or not isinstance(delivered, dict):
        raise ValueError("attestation must declare both vehicle and delivered temperature observations")
    vehicle_operator = vehicle.get("observation_operator")
    delivered_operator = delivered.get("observation_operator")
    if vehicle_operator not in VEHICLE_TEMPERATURE_OBSERVATION_OPERATORS:
        raise ValueError("vehicle temperature observation operator is unsupported or unknown")
    if delivered_operator != DELIVERED_TEMPERATURE_OBSERVATION_OPERATOR:
        raise ValueError("delivered temperature must be declared as delivered_gas_temperature")
    for declaration in (vehicle, delivered):
        if declaration.get("sensor_location_verified") is not True:
            raise ValueError("temperature observation requires a custodian-verified sensor location")
        if not isinstance(declaration.get("measurement_method"), str) or not declaration["measurement_method"].strip():
            raise ValueError("temperature observation requires a nonempty measurement method")
        if not isinstance(declaration.get("calibration_or_traceability"), str) or not declaration["calibration_or_traceability"].strip():
            raise ValueError("temperature observation requires a nonempty calibration or traceability declaration")
    return {
        "vehicle_temperature_degC": vehicle_operator,
        "delivered_gas_temperature_degC": delivered_operator,
    }


def _attestation(mapping: dict[str, Any], attestation: dict[str, Any]) -> dict[str, str]:
    if mapping.get("schema_version") != 1:
        raise ValueError("mapping schema_version must be 1")
    if attestation.get("schema_version") != 1:
        raise ValueError("attestation schema_version must be 1")
    if attestation.get("authorised_controlled_evaluation") is not True:
        raise ValueError("attestation must explicitly authorize controlled evaluation")
    if attestation.get("outcomes_accessed_before_protocol_freeze") is not False:
        raise ValueError("attestation must confirm no outcome access before protocol freeze")
    units = attestation.get("units")
    if not isinstance(units, dict):
        raise ValueError("attestation units must be an object")
    expected_units = {
        "vehicle_pressure_mpa_abs": "MPa_abs",
        "temperature_degC": "degC",
        "mass_flow_g_s": "g/s",
        "station_pressure_mpa_abs": "MPa_abs",
        "delivered_gas_temperature_degC": "degC",
        "cascade_source_pressure_mpa_abs": "MPa_abs",
    }
    for column, expected in expected_units.items():
        if units.get(column) != expected:
            raise ValueError(f"attestation unit for {column} must be {expected}")
    state_semantics = attestation.get("state_semantics")
    if not isinstance(state_semantics, dict) or not all(
        isinstance(state_semantics.get(column), str) and state_semantics[column].strip()
        for column in STATE_COLUMNS
    ):
        raise ValueError("attestation must provide nonempty semantics for every state column")
    metadata = attestation.get("metadata")
    if not isinstance(metadata, dict) or not all(
        isinstance(metadata.get(key), str) and metadata[key].strip()
        for key in REQUIRED_METADATA
    ):
        raise ValueError("attestation must provide all required generic metadata declarations")
    return _temperature_observation_semantics(attestation)


def _source_column(mapping: dict[str, Any], canonical: str) -> str:
    columns = mapping.get("column_map")
    if not isinstance(columns, dict) or not isinstance(columns.get(canonical), str):
        raise ValueError(f"mapping requires column_map.{canonical}")
    return columns[canonical]


def _input_format(path: Path) -> str:
    suffix = path.suffix.lower()
    if suffix == ".csv":
        return "csv"
    if suffix in {".xlsx", ".xlsm"}:
        return "xlsx"
    raise ValueError("input must be a CSV, XLSX, or XLSM file")


def _source_header_and_rows(
    input_path: Path,
    mapping: dict[str, Any],
    *,
    include_rows: bool,
) -> tuple[tuple[str, ...], list[dict[str, object]] | None, str]:
    """Read a controlled source without persisting source labels or values.

    ``include_rows=False`` is deliberately limited to the header contract.  It
    lets a custodian freeze the semantic mapping before model outcomes are
    read, while returning only canonical coverage in the preflight receipt.
    """

    source_format = _input_format(input_path)
    if source_format == "csv":
        delimiter = str(mapping.get("delimiter") or ",")
        encoding = str(mapping.get("encoding") or "utf-8-sig")
        with input_path.open("r", encoding=encoding, newline="") as handle:
            reader = csv.DictReader(handle, delimiter=delimiter)
            header = tuple(reader.fieldnames or ())
            rows = list(reader) if include_rows else None
        return header, rows, source_format

    worksheet_name = mapping.get("worksheet")
    if not isinstance(worksheet_name, str) or not worksheet_name.strip():
        raise ValueError("Excel mapping requires a nonempty worksheet name")
    try:
        from openpyxl import load_workbook
    except ImportError as exc:
        raise RuntimeError("openpyxl is required for controlled Excel intake") from exc
    workbook = load_workbook(input_path, read_only=True, data_only=True)
    try:
        if worksheet_name not in workbook.sheetnames:
            raise ValueError("specified worksheet is not present in the Excel input")
        worksheet = workbook[worksheet_name]
        iterator = worksheet.iter_rows(values_only=True)
        first_row = next(iterator, None)
        if first_row is None:
            return (), [] if include_rows else None, source_format
        header = tuple(str(value).strip() if value is not None else "" for value in first_row)
        if not include_rows:
            return header, None, source_format
        rows: list[dict[str, object]] = []
        for values in iterator:
            rows.append({
                header[index]: value
                for index, value in enumerate(values)
                if index < len(header)
            })
        return header, rows, source_format
    finally:
        workbook.close()


def _mapped_source_columns(mapping: dict[str, Any]) -> tuple[str, dict[str, str]]:
    return (
        _source_column(mapping, "time_s"),
        {column: _source_column(mapping, column) for column in OUTPUT_COLUMNS[1:]},
    )


def _missing_canonical_channels(
    header: tuple[str, ...], original_time: str, source_columns: dict[str, str],
) -> list[str]:
    available = set(header)
    missing = ["time_s"] if original_time not in available else []
    missing.extend(
        canonical for canonical, source in source_columns.items()
        if source not in available
    )
    return missing


def preflight_bundle(
    input_data: Path,
    mapping_path: Path,
    attestation_path: Path,
    protocol_path: Path,
) -> dict[str, Any]:
    """Check a private source schema without reading or retaining row values.

    The returned record intentionally contains canonical names only.  It is
    safe to retain with a controlled-access protocol because it never emits an
    original header, worksheet, time value, filename, identity, or row count.
    """

    input_data = _outside_repository(input_data, field="input data")
    mapping_path = _outside_repository(mapping_path, field="mapping")
    attestation_path = _outside_repository(attestation_path, field="attestation")
    if not input_data.is_file():
        raise FileNotFoundError(input_data)
    mapping = _json(mapping_path)
    attestation = _json(attestation_path)
    _attestation(mapping, attestation)
    protocol = _json(protocol_path)
    if protocol.get("status") != "prospective_intake_contract":
        raise ValueError("intake protocol must be prospective")
    original_time, source_columns = _mapped_source_columns(mapping)
    header, _, source_format = _source_header_and_rows(
        input_data, mapping, include_rows=False,
    )
    missing = _missing_canonical_channels(header, original_time, source_columns)
    return {
        "schema_version": 1,
        "artifact_type": "controlled_deidentified_hrs_full_loop_preflight",
        "ready_for_controlled_export": not missing,
        "source_format": source_format,
        "worksheet_mapping_required": source_format == "xlsx",
        "required_canonical_channels": list(OUTPUT_COLUMNS),
        "missing_canonical_channels": missing,
        "source_headers_exposed": False,
        "source_rows_read": False,
        "mapping_sha256": _sha256(mapping_path),
        "attestation_sha256": _sha256(attestation_path),
        "protocol_sha256": _sha256(protocol_path),
        "claim_boundary": (
            "This is a schema-and-attestation preflight only. It does not read "
            "outcome rows, establish trace quality, validate the model, or "
            "establish safety or IJHE readiness."
        ),
    }


def export_bundle(
    input_data: Path,
    mapping_path: Path,
    attestation_path: Path,
    output_directory: Path,
    protocol_path: Path,
) -> dict[str, Any]:
    """Create and locally verify one privacy-bounded full-loop event bundle."""

    input_data = _outside_repository(input_data, field="input data")
    mapping_path = _outside_repository(mapping_path, field="mapping")
    attestation_path = _outside_repository(attestation_path, field="attestation")
    output_directory = _outside_repository(output_directory, field="output directory")
    if not input_data.is_file():
        raise FileNotFoundError(input_data)
    mapping = _json(mapping_path)
    attestation = _json(attestation_path)
    temperature_observation = _attestation(mapping, attestation)
    protocol = _json(protocol_path)
    if protocol.get("status") != "prospective_intake_contract":
        raise ValueError("intake protocol must be prospective")

    original_time, source_columns = _mapped_source_columns(mapping)
    header, input_rows, source_format = _source_header_and_rows(
        input_data, mapping, include_rows=True,
    )
    missing = _missing_canonical_channels(header, original_time, source_columns)
    if missing:
        raise ValueError(
            "input source is missing mapped canonical channels: " + ", ".join(missing)
        )
    assert input_rows is not None
    if len(input_rows) < 20:
        raise ValueError("input source has fewer than 20 rows")

    normalized: list[dict[str, str]] = []
    origin: float | None = None
    previous_time: float | None = None
    for index, row in enumerate(input_rows, start=2):
        absolute_or_elapsed = _relative_time(row.get(original_time), mapping, row_number=index)
        if origin is None:
            origin = absolute_or_elapsed
        time_s = absolute_or_elapsed - origin
        if previous_time is not None and time_s <= previous_time:
            raise ValueError(f"row {index}: time must be strictly increasing")
        previous_time = time_s
        exported = {"time_s": f"{time_s:.9g}"}
        for column in NUMERIC_COLUMNS:
            exported[column] = f"{_finite(row.get(source_columns[column]), column=column, row_number=index):.9g}"
        for column in STATE_COLUMNS:
            raw_value = row.get(source_columns[column])
            value = raw_value.strip() if isinstance(raw_value, str) else str(raw_value or "").strip()
            if not value:
                raise ValueError(f"row {index}: {column} is empty")
            exported[column] = value
        normalized.append(exported)

    if output_directory.exists() and any(output_directory.iterdir()):
        raise ValueError("output directory must be empty to avoid mixing controlled data bundles")
    output_directory.mkdir(parents=True, exist_ok=True)
    trace_path = output_directory / "full_loop_event.csv"
    with trace_path.open("w", encoding="utf-8", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=OUTPUT_COLUMNS)
        writer.writeheader()
        writer.writerows(normalized)

    declaration = {
        "schema_version": 1,
        # Keep the declared type compatible with the frozen intake validator.
        # The controlled/de-identified properties below distinguish it from a
        # public archive without weakening the common channel contract.
        "declaration_type": "external_hrs_bundle_metadata",
        "outcomes_accessed_before_freeze": False,
        "controlled_access_only": True,
        "source_identifiers_published": False,
        "absolute_timestamps_published": False,
        "original_column_names_published": False,
        "channels": {
            "common_time_base": {"present": True},
            "vehicle_or_receptacle_pressure": {"present": True, "unit": "MPa_abs"},
            "gas_or_tank_temperature": {
                "present": True,
                "unit": "degC",
                "observation_operator": temperature_observation["vehicle_temperature_degC"],
            },
            "delivered_gas_temperature": {
                "present": True,
                "unit": "degC",
                "observation_operator": temperature_observation["delivered_gas_temperature_degC"],
            },
            "mass_flow_or_transferred_mass": {"present": True, "unit": "g/s"},
        },
        "metadata": {key: "declared" for key in REQUIRED_METADATA},
        "claim_boundary": (
            "A controlled, de-identified export that has passed channel and time-base "
            "screens is an evaluator input only. It is not an independently public "
            "holdout, safety certification, or an IJHE completion result."
        ),
    }
    declaration_path = output_directory / "declaration.json"
    declaration_path.write_text(json.dumps(declaration, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    manifest = build_manifest(output_directory, protocol_path)
    manifest_path = output_directory / "manifest.json"
    manifest_path.write_text(json.dumps(manifest, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    intake = validate_manifest(manifest_path, declaration_path, protocol_path, bundle_root=output_directory)
    screen = validate_full_loop_trace(
        trace_path, manifest_path, declaration_path, protocol_path, bundle_root=output_directory
    )
    if not screen.get("full_loop_trace_ready"):
        raise RuntimeError("export did not pass full-loop quality screen: " + "; ".join(screen.get("reasons") or []))

    receipt = {
        "schema_version": 1,
        "artifact_type": "controlled_deidentified_hrs_full_loop_export_receipt",
        "created_at_utc": datetime.now(timezone.utc).isoformat(),
        "controlled_access_only": True,
        "source_identifiers_published": False,
        "absolute_timestamps_published": False,
        "original_column_names_published": False,
        "raw_rows_persisted_in_repository": False,
        "temperature_observation_semantics_attested": True,
        "output_file_count": 4,
        "input_format": source_format,
        "rows": len(normalized),
        "time_duration_s": float(normalized[-1]["time_s"]),
        "generic_output_columns": list(OUTPUT_COLUMNS),
        "source_input_sha256": _sha256(input_data),
        "mapping_sha256": _sha256(mapping_path),
        "attestation_sha256": _sha256(attestation_path),
        "protocol_sha256": _sha256(protocol_path),
        "trace_sha256": _sha256(trace_path),
        "manifest_sha256": _sha256(manifest_path),
        "declaration_sha256": _sha256(declaration_path),
        "intake_decision": intake.get("decision"),
        "quality_decision": screen.get("decision"),
        "full_loop_trace_ready": screen.get("full_loop_trace_ready"),
        "claim_boundary": declaration["claim_boundary"],
    }
    receipt_path = output_directory / "receipt.json"
    receipt_path.write_text(json.dumps(receipt, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    return receipt


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--input", type=Path, required=True, help="Authorised local CSV, XLSX, or XLSM event")
    parser.add_argument("--mapping", type=Path, required=True, help="Local semantic mapping JSON outside this repository")
    parser.add_argument("--attestation", type=Path, required=True, help="Local unit/state attestation JSON outside this repository")
    parser.add_argument("--output-directory", type=Path, help="Controlled export directory outside this repository")
    parser.add_argument("--protocol", type=Path, default=ROOT / "research/external_hrs_intake_protocol.json")
    parser.add_argument("--preflight", action="store_true", help="Check header coverage without reading source rows")
    args = parser.parse_args()
    if args.preflight:
        report = preflight_bundle(args.input, args.mapping, args.attestation, args.protocol)
        print(json.dumps(report, ensure_ascii=False, indent=2))
        return 0 if report["ready_for_controlled_export"] else 2
    if args.output_directory is None:
        parser.error("--output-directory is required unless --preflight is used")
    receipt = export_bundle(args.input, args.mapping, args.attestation, args.output_directory, args.protocol)
    print(json.dumps(receipt, ensure_ascii=False, indent=2))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
