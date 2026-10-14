#!/usr/bin/env python3
"""Audit the public Grüne/PS H2 dispersion release without copying raw rows.

The dataset is a useful actual-hydrogen spatial benchmark, but it is being
opened after the current detector analysis.  This script therefore creates a
post-access intake record only; it cannot close the independent H2SAFE gate or
promote a detector-placement model.
"""

from __future__ import annotations

import argparse
from datetime import datetime, timezone
import hashlib
import json
from pathlib import Path
import re
import zipfile
from typing import Any

import openpyxl


DOI = "10.5281/zenodo.4668554"
RECORD_URL = "https://zenodo.org/records/4668554"
LICENSE = "CC BY 4.0"
EXPECTED = {
    "HTE2440PS000MIXED_300321.pdf": (706086, "d4b59189dd4017c695cc652f4f572128"),
    "HTE2440PS001MIXED_300321.xlsx": (52610, "da176772755b1a2449a76df9e12b9cd3"),
    "HTE2440PS002MIXED_300321.xlsx": (65696, "20f996be57fc3d8cf7d091d1f982bd75"),
    "HTE2440PS003MIXED_300321.xlsx": (50590, "953b6a033553288831b1ee02b881dbec"),
    "HTE2440PS004MIXED_300321.xlsx": (81649, "093d3d8e6ef030615704ee9096d3d84f"),
    "HTE2440PS005FLMT00300321.xlsx": (26724, "c4b9195a2bb89213eda9aa2ac9e9ec07"),
    "HTE2440PS006MIXED_300321.zip": (11108400, "ade3911c0df50ee70486153455f4bd12"),
}


def md5(path: Path) -> str:
    digest = hashlib.md5()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def _text(value: Any) -> str:
    return str(value).strip() if value is not None else ""


def _sheet_summary(ws: Any) -> dict[str, Any]:
    rows = list(ws.iter_rows(values_only=True))
    header_row = None
    x_start = None
    for row_index, row in enumerate(rows):
        for column_index, value in enumerate(row):
            if _text(value).lower() == "x/ mm":
                header_row = row_index
                x_start = column_index + 1
                break
        if header_row is not None:
            break
    if header_row is None or x_start is None:
        return {
            "title": ws.title,
            "rows": ws.max_row,
            "columns": ws.max_column,
            "has_coordinate_grid": False,
            "numeric_measurement_cells": 0,
            "hydrogen_sheet": "H2" in ws.title.upper(),
            "helium_sheet": "HE" in ws.title.upper(),
            "ventilation_condition_declared": bool(re.search(r"W\s*=", ws.title, re.IGNORECASE)),
        }

    x_coordinates = [
        value for value in rows[header_row][x_start:]
        if isinstance(value, (int, float))
    ]
    measurement_cells = 0
    numeric_measurement_cells = 0
    for row in rows[header_row + 1 :]:
        values = row[x_start : x_start + len(x_coordinates)]
        for value in values:
            if value is None or _text(value) == "":
                continue
            measurement_cells += 1
            if isinstance(value, (int, float)):
                numeric_measurement_cells += 1

    metadata = " ".join(_text(value) for row in rows[: header_row + 1] for value in row[:x_start])
    nozzle = re.search(r"Nozzle\s*/\s*mm", metadata, re.IGNORECASE)
    return {
        "title": ws.title,
        "rows": ws.max_row,
        "columns": ws.max_column,
        "has_coordinate_grid": bool(x_coordinates),
        "x_coordinate_count": len(x_coordinates),
        "measurement_row_count": max(0, len(rows) - header_row - 1),
        "measurement_cells": measurement_cells,
        "numeric_measurement_cells": numeric_measurement_cells,
        "numeric_measurement_fraction": (
            numeric_measurement_cells / measurement_cells if measurement_cells else 0.0
        ),
        "header_found_at_zero_based_row": header_row,
        "nozzle_label_present": bool(nozzle),
        "hydrogen_sheet": "H2" in ws.title.upper(),
        "helium_sheet": "HE" in ws.title.upper(),
        "ventilation_condition_declared": bool(re.search(r"W\s*=", ws.title, re.IGNORECASE)),
    }


def build(directory: Path) -> dict[str, Any]:
    files = []
    for name, (size, checksum) in EXPECTED.items():
        path = directory / name
        if not path.exists():
            raise FileNotFoundError(path)
        actual = md5(path)
        if path.stat().st_size != size or actual != checksum:
            raise ValueError(f"hash/size mismatch for {name}")
        files.append({"name": name, "bytes": size, "md5": actual, "raw_rows_committed": False})

    workbooks = []
    for path in sorted(directory.glob("*.xlsx")):
        workbook = openpyxl.load_workbook(path, read_only=True, data_only=True)
        sheets = [_sheet_summary(sheet) for sheet in workbook.worksheets]
        workbook.close()
        workbooks.append({
            "name": path.name,
            "sheet_count": len(sheets),
            "hydrogen_sheet_count": sum(item["hydrogen_sheet"] for item in sheets),
            "helium_sheet_count": sum(item["helium_sheet"] for item in sheets),
            "coordinate_grid_sheet_count": sum(item["has_coordinate_grid"] for item in sheets),
            "sheets": sheets,
        })

    zip_path = directory / "HTE2440PS006MIXED_300321.zip"
    with zipfile.ZipFile(zip_path) as archive:
        members = archive.namelist()
    return {
        "schema_version": 1,
        "artifact_type": "public_actual_hydrogen_spatial_intake",
        "generated_at": datetime.now(timezone.utc).isoformat(),
        "evidence_role": "post_access_development_diagnostic",
        "source": {
            "title": "Efficiency of mechanical ventilation on H2 dispersion (PS)",
            "doi": DOI,
            "record_url": RECORD_URL,
            "license": LICENSE,
            "access_right": "open",
            "publisher": "Zenodo",
            "authors": ["Joachim Grüne", "Karsten Sempert"],
        },
        "integrity": {
            "file_count": len(files),
            "files": files,
            "cad_archive_member_count": len(members),
            "cad_archive_member_examples": members[:10],
        },
        "schema_summary": {
            "workbook_count": len(workbooks),
            "worksheet_count": sum(item["sheet_count"] for item in workbooks),
            "hydrogen_worksheet_count": sum(item["hydrogen_sheet_count"] for item in workbooks),
            "helium_worksheet_count": sum(item["helium_sheet_count"] for item in workbooks),
            "coordinate_grid_worksheet_count": sum(item["coordinate_grid_sheet_count"] for item in workbooks),
            "workbooks": workbooks,
        },
        "eligibility": {
            "actual_hydrogen_spatial_concentration_data": True,
            "ventilation_condition_metadata": True,
            "sensor_coordinate_grid_context": True,
            "time_resolved_logger_trace": False,
            "station_controller_or_vehicle_channels": False,
            "independent_h2safe_holdout": False,
            "reason_not_independent": "The archive was opened after the current detector analysis; it is retained as post-access development evidence only.",
            "candidate_use": "Qualitative and development-stage spatial detector/HVAC model comparison after an explicit protocol freeze.",
        },
        "claim_boundary": "This artifact verifies the public CC BY 4.0 archive, actual-hydrogen concentration sheets, ventilation conditions and coordinate-grid structure. It does not close the H2SAFE independent spatial holdout, calibrate station alarm thresholds, validate station-to-vehicle operation, or authorize runtime detector routing.",
        "runtime_parameter_updated": False,
    }


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--directory", type=Path, required=True)
    parser.add_argument("--output", type=Path, default=Path("research/grune_ps_h2_dispersion_intake_2026_10_14.json"))
    args = parser.parse_args()
    result = build(args.directory.resolve())
    args.output.parent.mkdir(parents=True, exist_ok=True)
    args.output.write_text(json.dumps(result, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    print(json.dumps({
        "workbooks": result["schema_summary"]["workbook_count"],
        "worksheets": result["schema_summary"]["worksheet_count"],
        "hydrogen_worksheets": result["schema_summary"]["hydrogen_worksheet_count"],
        "coordinate_grid_worksheets": result["schema_summary"]["coordinate_grid_worksheet_count"],
        "claim_boundary": result["claim_boundary"],
    }, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
