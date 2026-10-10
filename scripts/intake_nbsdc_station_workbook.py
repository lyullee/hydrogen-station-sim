"""Privacy-bounded intake for an authorized NBSDC HRS workbook.

The NBSDC public record exposes the entity file only after registration/login.
This command is designed for a custodian-approved download.  It hashes the
input and records only aggregate schema evidence; it never copies the workbook,
serializes original headers, sheet names, timestamps, rows, or source paths.
The result is an intake receipt, not a validation result.
"""
from __future__ import annotations

import argparse
import csv
from dataclasses import dataclass
from hashlib import sha256
import json
from pathlib import Path
import re
from typing import Iterable, Iterator
from zipfile import BadZipFile, ZipFile

from full_loop_schema_contract import CONTRACT_REQUIRED, contract_coverage

ROOT = Path(__file__).resolve().parents[1]
SUPPORTED = {".csv", ".xlsx", ".xlsm", ".zip"}
MAX_MEMBER_BYTES = 250 * 1024 * 1024
HEADER_SCAN_ROWS = 40

SEMANTIC_TERMS: dict[str, tuple[str, ...]] = {
    "time": ("time", "timestamp", "datetime", "localtime", "zeit", "时间", "시간", "일시", "시각"),
    "pressure": ("pressure", "press", "pt", "pi", "압력"),
    "temperature": ("temperature", "temp", "tt", "온도"),
    "mass_flow": ("flow", "mass", "rate", "mfm", "fqi", "fwg", "유량", "질량"),
    "vehicle": ("vehicle", "car", "automobile", "veh", "차량", "자동차", "수소차"),
    "cascade_or_storage": ("cascade", "bank", "storage", "vessel", "저장", "뱅크", "용기"),
    "controller_state": ("state", "status", "valve", "esd", "mode", "control", "alarm", "interlock", "밸브", "상태", "제어", "경보", "차단"),
    "compressor": ("compress", "comp", "압축", "부스터"),
    "dispenser_or_nozzle": ("dispenser", "nozzle", "hose", "refuel", "충전기", "노즐", "호스", "충전"),
}
FULL_LOOP_REQUIRED = frozenset({"time", "pressure", "temperature", "mass_flow", "vehicle", "cascade_or_storage", "controller_state"})

@dataclass(frozen=True)
class TableSchema:
    coverage: frozenset[str]
    contract_coverage: frozenset[str]
    row_count: int


def _outside_repo(path: Path) -> Path:
    resolved = path.resolve()
    try:
        resolved.relative_to(ROOT.resolve())
    except ValueError:
        return resolved
    raise ValueError("input must be outside the repository worktree; raw files are never committed")


def _classify(labels: Iterable[object]) -> frozenset[str]:
    text = " ".join(str(x or "").strip().lower() for x in labels if str(x or "").strip())
    compact = re.sub(r"[^a-z0-9가-힣]+", " ", text)
    out: set[str] = set()
    for group, terms in SEMANTIC_TERMS.items():
        if any(term in text or term in compact for term in terms):
            out.add(group)
    return frozenset(out)


def _rows_csv(path: Path) -> Iterator[list[str]]:
    with path.open("r", encoding="utf-8-sig", errors="replace", newline="") as handle:
        yield from csv.reader(handle)


def _rows_xlsx(path: Path) -> Iterator[list[object]]:
    try:
        import openpyxl
    except ImportError as exc:  # pragma: no cover - environment-specific
        raise RuntimeError("openpyxl is required for XLSX intake") from exc
    workbook = openpyxl.load_workbook(path, read_only=True, data_only=True)
    try:
        for worksheet in workbook.worksheets:
            for row in worksheet.iter_rows(values_only=True):
                yield list(row)
    finally:
        workbook.close()


def _table_schema(rows: Iterator[list[object]]) -> TableSchema:
    buffered: list[list[object]] = []
    for row in rows:
        buffered.append(row)
        if len(buffered) >= HEADER_SCAN_ROWS:
            break
    if not buffered:
        return TableSchema(frozenset(), frozenset(), 0)
    best = max(buffered, key=lambda row: (len(_classify(row)), sum(bool(str(v or '').strip()) for v in row)))
    coverage = _classify(best)
    # Count only the rows we can safely inspect in streaming mode; the count is
    # intentionally a lower bound and is not used as validation evidence.
    return TableSchema(coverage, contract_coverage(best), max(0, len(buffered) - 1))


def _sha256(path: Path) -> str:
    digest = sha256()
    with path.open("rb") as handle:
        for block in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(block)
    return digest.hexdigest()


def _inspect(path: Path) -> tuple[list[TableSchema], int]:
    suffix = path.suffix.lower()
    if suffix == ".csv":
        return [_table_schema(_rows_csv(path))], 1
    if suffix in {".xlsx", ".xlsm"}:
        # Keep worksheet boundaries private while still reporting aggregate
        # coverage.  Each worksheet is a possible logical table.
        try:
            import openpyxl
        except ImportError as exc:  # pragma: no cover
            raise RuntimeError("openpyxl is required for XLSX intake") from exc
        workbook = openpyxl.load_workbook(path, read_only=True, data_only=True)
        try:
            schemas = [_table_schema((list(row) for row in ws.iter_rows(values_only=True))) for ws in workbook.worksheets]
            return schemas, len(workbook.worksheets)
        finally:
            workbook.close()
    if suffix == ".zip":
        schemas: list[TableSchema] = []
        with ZipFile(path) as archive:
            for info in archive.infolist():
                if info.is_dir() or info.file_size > MAX_MEMBER_BYTES:
                    continue
                member_suffix = Path(info.filename).suffix.lower()
                if member_suffix == ".csv":
                    with archive.open(info) as raw:
                        text = (line.decode("utf-8-sig", errors="replace") for line in raw)
                        schemas.append(_table_schema(csv.reader(text)))
                elif member_suffix in {".xlsx", ".xlsm"}:
                    # Do not unpack an untrusted archive into the worktree. A
                    # future custodian workflow should supply the extracted file.
                    continue
        return schemas, len(schemas)
    raise ValueError(f"unsupported input format: {suffix}")


def build_receipt(
    path: Path,
    *,
    source_url: str,
    persistent_identifier: str,
    rights: str,
    access_attested: bool,
    output_path: Path | None = None,
) -> dict[str, object]:
    source = _outside_repo(path)
    if source.suffix.lower() not in SUPPORTED:
        raise ValueError(f"unsupported input format: {source.suffix}")
    if not source.is_file():
        raise FileNotFoundError(source)
    schemas, table_count = _inspect(source)
    observed = set().union(*(set(schema.coverage) for schema in schemas)) if schemas else set()
    co_located = any(FULL_LOOP_REQUIRED <= schema.coverage for schema in schemas)
    strict_co_located = any(CONTRACT_REQUIRED <= schema.contract_coverage for schema in schemas)
    missing = sorted(FULL_LOOP_REQUIRED - observed)
    contract_observed = set().union(*(set(schema.contract_coverage) for schema in schemas)) if schemas else set()
    receipt: dict[str, object] = {
        "schema_version": 1,
        "artifact_type": "nbsdc_authorized_hrs_schema_intake",
        "status": "schema_intake_complete",
        "source": {
            "persistent_identifier": persistent_identifier,
            "source_url": source_url,
            "rights": rights,
            "access_attested": bool(access_attested),
            "file_format": source.suffix.lower().lstrip("."),
            "sha256": _sha256(source),
            "bytes": source.stat().st_size,
        },
        "aggregate_schema": {
            "table_count": table_count,
            "tables_with_time": sum("time" in s.coverage for s in schemas),
            "tables_with_pressure": sum("pressure" in s.coverage for s in schemas),
            "tables_with_temperature": sum("temperature" in s.coverage for s in schemas),
            "tables_with_mass_flow": sum("mass_flow" in s.coverage for s in schemas),
            "tables_with_vehicle": sum("vehicle" in s.coverage for s in schemas),
            "tables_with_controller_state": sum("controller_state" in s.coverage for s in schemas),
            "observed_semantic_groups": sorted(observed),
            "missing_full_loop_semantic_groups": missing,
            "full_loop_channels_co_located_in_one_table": co_located,
            "full_loop_contract_channels_co_located_in_one_table": strict_co_located,
            "observed_full_loop_contract_groups": sorted(contract_observed),
            "missing_full_loop_contract_groups": sorted(CONTRACT_REQUIRED - contract_observed),
        },
        "admission": {
            "full_loop_external_holdout_eligible": False,
            "reason": "Schema evidence alone cannot attest units, calibration, synchronization, vehicle-event linkage, controller/bank semantics, rights for reviewer access, or a frozen outcome-free split.",
            "next_checks": [
                "custodian attestation for channel names and engineering units",
                "common-clock and station/dispenser/vehicle event linkage",
                "calibration and sampling uncertainty metadata",
                "protocol/model/split hash freeze before numerical outcome review",
                "independent reviewer access under the declared rights",
            ],
        },
        "privacy": {
            "raw_rows_persisted": False,
            "original_headers_persisted": False,
            "worksheet_names_persisted": False,
            "absolute_timestamps_persisted": False,
            "source_path_persisted": False,
            "site_company_manufacturer_published": False,
        },
    }
    if output_path is not None:
        output_path.parent.mkdir(parents=True, exist_ok=True)
        output_path.write_text(json.dumps(receipt, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    return receipt


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--input", type=Path, required=True)
    parser.add_argument("--output", type=Path, required=True)
    parser.add_argument("--source-url", default="https://www.nbsdc.cn/general/dataDetail?id=64f08267bb16e06dfdc78c57&type=1")
    parser.add_argument("--persistent-identifier", default="CSTR 16666.11.nbsdc.8c2MpcJl")
    parser.add_argument("--rights", default="CC0 1.0 (as displayed on the official NBSDC record; retain custodian attestation)")
    parser.add_argument("--access-attested", action="store_true")
    args = parser.parse_args()
    receipt = build_receipt(args.input, source_url=args.source_url, persistent_identifier=args.persistent_identifier, rights=args.rights, access_attested=args.access_attested, output_path=args.output)
    print(json.dumps({"status": receipt["status"], "sha256": receipt["source"]["sha256"], "full_loop_external_holdout_eligible": receipt["admission"]["full_loop_external_holdout_eligible"], "missing_full_loop_semantic_groups": receipt["aggregate_schema"]["missing_full_loop_semantic_groups"]}, ensure_ascii=False))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
