"""Screen the frozen Type-IV tank model on NREL's H2FillS HDVS sample.

The workbook is an independent physical test distributed with the H2FillS
package.  It is intentionally not committed: the package licence permits
internal use but does not grant redistribution rights for the raw file.  This
runner therefore takes a local path, records the workbook digest and writes
only derived metrics to the ignored validation-results directory.

This is a tank-submodel validation.  The sample has no hose, nozzle,
receptacle or station-controller traces, so its result cannot be presented as
full HRS closed-loop validation.
"""

from __future__ import annotations

import argparse
from dataclasses import dataclass
from datetime import datetime, timezone
import hashlib
import json
from pathlib import Path
import subprocess
import sys
from typing import Any

import numpy as np
from openpyxl import load_workbook
from scipy.integrate import solve_ivp

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "src"))

from h2station.scenario import build_vehicle_tank  # noqa: E402
from h2station.tabulated import PropsSI  # noqa: E402
from h2station.vehicle import (  # noqa: E402
    CompositeTankFitParameters,
    CompositeTankState,
    TankBoundaryFlow,
)


TANK_IDS = (1, 2, 3, 5, 7, 8, 9)
SCREENING_LIMITS = {
    "pressure_rmse_mpa_max": 5.0,
    "temperature_rmse_c_max": 10.0,
    "mass_final_abs_error_kg_max": 0.5,
}
DEFAULT_WORKBOOK = ROOT / "data/public_validation/raw/nrel_h2fills_2022_hdvs_typeiv.xlsx"
DEFAULT_OUTPUT = ROOT / (
    "data/public_validation/results/nrel_h2fills_hdvs_typeiv/validation.json"
)


@dataclass(frozen=True)
class TankTrace:
    tank_id: int
    time_s: np.ndarray
    inlet_pressure_mpa: np.ndarray
    inlet_temperature_c: np.ndarray
    mass_kg: np.ndarray
    pressure_mpa: np.ndarray
    temperature_c: np.ndarray


@dataclass(frozen=True)
class NrelDataset:
    path: Path
    description: dict[str, str]
    traces: tuple[TankTrace, ...]

    @property
    def sample_count(self) -> int:
        return int(len(self.traces[0].time_s))

    @property
    def time_s(self) -> np.ndarray:
        return self.traces[0].time_s

    def summary(self) -> dict[str, Any]:
        total_mass = np.sum([trace.mass_kg for trace in self.traces], axis=0)
        return {
            "tank_ids": list(TANK_IDS),
            "tank_count": len(self.traces),
            "sample_count": self.sample_count,
            "sample_period_median_s": float(np.median(np.diff(self.time_s))),
            "duration_s": float(self.time_s[-1] - self.time_s[0]),
            "capacity_per_tank_kg": 9.8,
            "capacity_total_kg": 68.6,
            "initial_total_mass_kg": float(total_mass[0]),
            "final_total_mass_kg": float(total_mass[-1]),
            "mass_added_kg": float(total_mass[-1] - total_mass[0]),
            "reported_mass_added_kg": 61.5,
            "mass_closure_ratio_to_reported": float(
                (total_mass[-1] - total_mass[0]) / 61.5
            ),
            "initial_pressure_mean_mpa": float(
                np.mean([trace.pressure_mpa[0] for trace in self.traces])
            ),
            "final_pressure_mean_mpa": float(
                np.mean([trace.pressure_mpa[-1] for trace in self.traces])
            ),
            "peak_temperature_mean_c": float(
                np.max(np.mean([trace.temperature_c for trace in self.traces], axis=0))
            ),
            "ambient_temperature_c": 15.0,
        }


def _implied_volume_m3(trace: TankTrace) -> np.ndarray:
    """Infer the EOS-equivalent gas volume from measured tank states.

    This is a diagnostic only.  It does not change the frozen model, fit any
    parameter, or establish vessel geometry: a time-varying value can also
    reflect measurement bias, non-equilibrium temperature, or an unmodelled
    boundary volume.  The calculation is retained to distinguish a geometry
    mismatch from an integrator failure when an independent tank test is
    screened.
    """

    volumes: list[float] = []
    for pressure_mpa, temperature_c, mass_kg in zip(
        trace.pressure_mpa, trace.temperature_c, trace.mass_kg
    ):
        if mass_kg <= 0.0:
            volumes.append(float("nan"))
            continue
        density = float(
            PropsSI(
                "Dmass",
                "P",
                max(2.0e5, float(pressure_mpa) * 1.0e6),
                "T",
                float(temperature_c) + 273.15,
                "Hydrogen",
            )
        )
        volumes.append(float(mass_kg) / density)
    return np.asarray(volumes, dtype=float)


def _geometry_diagnostic(
    traces: tuple[TankTrace, ...],
    *,
    capacity_kg: float,
    effective_volume_multiplier: float,
) -> dict[str, Any]:
    """Report observed EOS-implied volume against the frozen geometry assumption."""

    nominal_volume_m3 = 0.122 * capacity_kg / 4.7
    frozen_volume_m3 = nominal_volume_m3 * effective_volume_multiplier
    rows: list[dict[str, Any]] = []
    for trace in traces:
        implied = _implied_volume_m3(trace)
        finite = implied[np.isfinite(implied)]
        if finite.size == 0:
            raise ValueError(f"Tank {trace.tank_id} has no positive-mass state")
        median = float(np.median(finite))
        rows.append(
            {
                "tank_id": trace.tank_id,
                "implied_volume_m3_median": median,
                "implied_volume_m3_final": float(implied[-1]),
                "implied_volume_m3_min": float(np.min(finite)),
                "implied_volume_m3_max": float(np.max(finite)),
                "ratio_to_frozen_effective_volume": median / frozen_volume_m3,
                "difference_from_frozen_effective_volume_m3": median - frozen_volume_m3,
            }
        )
    medians = np.asarray([row["implied_volume_m3_median"] for row in rows])
    return {
        "status": "diagnostic_only_no_parameter_update",
        "method": (
            "EOS-equivalent volume = measured hydrogen mass / CoolProp-tabulated "
            "density at measured internal pressure and temperature"
        ),
        "nominal_volume_assumption_m3": nominal_volume_m3,
        "frozen_effective_volume_m3": frozen_volume_m3,
        "implied_volume_m3_median_across_tanks": float(np.median(medians)),
        "ratio_to_frozen_effective_volume_median": float(
            np.median(medians) / frozen_volume_m3
        ),
        "rows": rows,
        "interpretation": (
            "A systematic ratio below one is consistent with a geometry or state-"
            "definition mismatch, but is not sufficient to identify vessel volume. "
            "It must not be fitted on this external screen or presented as an "
            "independent validation pass."
        ),
    }


def _sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def _description(workbook: Any) -> dict[str, str]:
    sheet = workbook["Description"]
    result: dict[str, str] = {}
    for row in sheet.iter_rows(min_row=1, values_only=True):
        if len(row) >= 3 and row[1] is not None:
            name, value = row[1], row[2]
        elif len(row) >= 2:
            name, value = row[0], row[1]
        else:
            continue
        if name is None or value is None:
            continue
        result[str(name).strip()] = str(value).strip()
    return result


def read_nrel_workbook(path: Path) -> NrelDataset:
    """Read the official supplemental workbook without changing its values."""

    if not path.is_file():
        raise FileNotFoundError(
            f"NREL workbook not found: {path}. Download it from the official NLR page."
        )
    workbook = load_workbook(path, data_only=True, read_only=True)
    if "Data" not in workbook.sheetnames or "Description" not in workbook.sheetnames:
        raise ValueError("NREL workbook must contain Description and Data sheets")
    description = _description(workbook)
    rows = workbook["Data"].iter_rows(values_only=True)
    try:
        header = tuple(next(rows))
    except StopIteration as exc:
        raise ValueError("NREL Data sheet is empty") from exc
    index = {str(value): i for i, value in enumerate(header) if value is not None}
    required = {
        "Time [s]",
        *(
            f"HDVS_ tank#{tank_id}_{field}"
            for tank_id in TANK_IDS
            for field in (
                "inlet_press [MPa]",
                "inlet_temp [degC]",
                "mass [kg]",
                "internal_press [MPa]",
                "internal_temp [degC]",
            )
        ),
    }
    missing = sorted(required - set(index))
    if missing:
        raise ValueError(f"NREL Data sheet is missing columns: {missing}")
    raw_rows = [row for row in rows if row[0] is not None]
    if len(raw_rows) < 3:
        raise ValueError("NREL Data sheet must contain at least three timed rows")
    time_s = np.asarray([float(row[index["Time [s]"]]) for row in raw_rows])
    if np.any(~np.isfinite(time_s)) or np.any(np.diff(time_s) <= 0.0):
        raise ValueError("NREL time samples must be finite and strictly increasing")

    traces: list[TankTrace] = []
    for tank_id in TANK_IDS:
        def values(field: str) -> np.ndarray:
            return np.asarray(
                [float(row[index[f"HDVS_ tank#{tank_id}_{field}"]]) for row in raw_rows],
                dtype=float,
            )

        trace = TankTrace(
            tank_id=tank_id,
            time_s=time_s.copy(),
            inlet_pressure_mpa=values("inlet_press [MPa]"),
            inlet_temperature_c=values("inlet_temp [degC]"),
            mass_kg=values("mass [kg]"),
            pressure_mpa=values("internal_press [MPa]"),
            temperature_c=values("internal_temp [degC]"),
        )
        for array in (
            trace.inlet_pressure_mpa,
            trace.inlet_temperature_c,
            trace.mass_kg,
            trace.pressure_mpa,
            trace.temperature_c,
        ):
            if np.any(~np.isfinite(array)):
                raise ValueError(f"Non-finite value in tank {tank_id}")
        if np.any(trace.mass_kg < 0.0):
            raise ValueError(f"Tank {tank_id} mass contains a negative value")
        traces.append(trace)
    return NrelDataset(path=path, description=description, traces=tuple(traces))


def _simulate_tank(
    trace: TankTrace,
    fit: CompositeTankFitParameters,
    *,
    capacity_kg: float = 9.8,
    ambient_temperature_c: float = 15.0,
) -> dict[str, Any]:
    """Run the frozen tank model using measured mass and inlet conditions."""

    tank = build_vehicle_tank(0.122 * capacity_kg / 4.7, fit)
    initial = tank.initial_state(
        trace.pressure_mpa[0] * 1e6,
        trace.temperature_c[0] + 273.15,
    )
    # The workbook is sampled every second.  Interpolation preserves the
    # measured boundary without tuning the model to this independent test.
    measured_flow = np.gradient(trace.mass_kg, trace.time_s)

    def rhs(time_s: float, values: np.ndarray) -> np.ndarray:
        state = CompositeTankState.from_vector(values)
        inlet_pressure_pa = float(np.interp(time_s, trace.time_s, trace.inlet_pressure_mpa)) * 1e6
        inlet_temperature_k = float(
            np.interp(time_s, trace.time_s, trace.inlet_temperature_c)
        ) + 273.15
        inlet_enthalpy = float(
            PropsSI(
                "Hmass",
                "P",
                max(2.0e5, inlet_pressure_pa),
                "T",
                inlet_temperature_k,
                "Hydrogen",
            )
        )
        return tank.derivative(
            state,
            TankBoundaryFlow(
                inlet_mass_flow_kg_s=max(
                    0.0, float(np.interp(time_s, trace.time_s, measured_flow))
                ),
                inlet_specific_enthalpy_j_kg=inlet_enthalpy,
                ambient_temperature_k=ambient_temperature_c + 273.15,
            ),
        ).as_vector()

    solution = solve_ivp(
        rhs,
        (float(trace.time_s[0]), float(trace.time_s[-1])),
        initial.as_vector(),
        t_eval=trace.time_s,
        method="BDF",
        rtol=5.0e-6,
        atol=5.0e-8,
    )
    if not solution.success:
        raise RuntimeError(f"Tank {trace.tank_id} integration failed: {solution.message}")
    states = [CompositeTankState.from_vector(row) for row in solution.y.T]
    gas = [tank.gas_state(state) for state in states]
    predicted_pressure = np.asarray([state.pressure_pa / 1e6 for state in gas])
    predicted_temperature = np.asarray([state.temperature_k - 273.15 for state in gas])
    predicted_mass = np.asarray([state.hydrogen_mass_kg for state in states])
    implied_volume = _implied_volume_m3(trace)
    finite_implied_volume = implied_volume[np.isfinite(implied_volume)]
    return {
        "tank_id": trace.tank_id,
        "sample_count": int(len(trace.time_s)),
        "pressure_rmse_mpa": float(np.sqrt(np.mean((predicted_pressure - trace.pressure_mpa) ** 2))),
        "temperature_rmse_c": float(np.sqrt(np.mean((predicted_temperature - trace.temperature_c) ** 2))),
        "mass_rmse_kg": float(np.sqrt(np.mean((predicted_mass - trace.mass_kg) ** 2))),
        "pressure_final_error_mpa": float(predicted_pressure[-1] - trace.pressure_mpa[-1]),
        "temperature_peak_error_c": float(
            np.max(predicted_temperature) - np.max(trace.temperature_c)
        ),
        "mass_final_error_kg": float(predicted_mass[-1] - trace.mass_kg[-1]),
        "screening_pass": bool(
            np.sqrt(np.mean((predicted_pressure - trace.pressure_mpa) ** 2))
            <= SCREENING_LIMITS["pressure_rmse_mpa_max"]
            and np.sqrt(np.mean((predicted_temperature - trace.temperature_c) ** 2))
            <= SCREENING_LIMITS["temperature_rmse_c_max"]
            and abs(predicted_mass[-1] - trace.mass_kg[-1])
            <= SCREENING_LIMITS["mass_final_abs_error_kg_max"]
        ),
        "geometry_diagnostic": {
            "nominal_volume_m3": 0.122 * capacity_kg / 4.7,
            "frozen_effective_volume_m3": (
                0.122 * capacity_kg / 4.7 * fit.effective_volume_multiplier
            ),
            "implied_volume_m3_median": float(np.median(finite_implied_volume)),
        },
    }


def _git_commit() -> str | None:
    try:
        return subprocess.check_output(
            ["git", "rev-parse", "HEAD"], cwd=ROOT, text=True
        ).strip()
    except (OSError, subprocess.CalledProcessError):
        return None


def _git_dirty() -> bool | None:
    try:
        return bool(
            subprocess.check_output(
                ["git", "status", "--porcelain"], cwd=ROOT, text=True
            ).strip()
        )
    except (OSError, subprocess.CalledProcessError):
        return None


def run(workbook_path: Path, output_path: Path, fit_path: Path) -> dict[str, Any]:
    dataset = read_nrel_workbook(workbook_path)
    fit_json = json.loads(fit_path.read_text(encoding="utf-8"))
    fit_data = fit_json.get("fit") or {}
    fit = CompositeTankFitParameters(
        effective_volume_multiplier=float(fit_data["effective_volume_multiplier"]),
        gas_liner_ua_multiplier=float(fit_data["gas_liner_ua_multiplier"]),
    )
    rows = [
        _simulate_tank(trace, fit)
        for trace in dataset.traces
    ]
    geometry_diagnostic = _geometry_diagnostic(
        dataset.traces,
        capacity_kg=9.8,
        effective_volume_multiplier=fit.effective_volume_multiplier,
    )
    aggregate = {
        "tank_count": len(rows),
        "sample_count_per_tank": dataset.sample_count,
        "screening_pass_count": sum(row["screening_pass"] for row in rows),
        "screening_pass_fraction": float(
            sum(row["screening_pass"] for row in rows) / len(rows)
        ),
        **{
            name: float(np.mean([row[name] for row in rows]))
            for name in (
                "pressure_rmse_mpa",
                "temperature_rmse_c",
                "mass_rmse_kg",
                "pressure_final_error_mpa",
                "temperature_peak_error_c",
                "mass_final_error_kg",
            )
        },
    }
    report: dict[str, Any] = {
        "schema_version": 1,
        "generated_at": datetime.now(timezone.utc).isoformat(),
        "evidence_role": "independent_tank_thermal_external_validation",
        "claim_boundary": (
            "Tank pressure/temperature response only. The workbook has no hose, "
            "nozzle, receptacle or station-controller trace and cannot establish "
            "full HRS closed-loop or field safety performance."
        ),
        "source": {
            "title": "NREL H2FillS 2022 HDVS Type IV sample test result",
            "publisher": "National Renewable Energy Laboratory / National Laboratory of the Rockies",
            "official_download_page": "https://www.nlr.gov/hydrogen/h2fills-download",
            "package_url": "https://www.nlr.gov/docs/libraries/hydrogen/h2fills-package.zip?sfvrsn=4ee4550e_0",
            "workbook_name": workbook_path.name,
            "workbook_path": str(workbook_path),
            "workbook_sha256": _sha256(workbook_path),
            "package_sha256": "f44a079ffbd995275dd7b6e1c55b557fa1736c512b419f2a6666057b99c43ac0",
            "raw_file_policy": "local ignored file only; do not redistribute without permission",
            "dataset_summary": dataset.summary(),
        },
        "frozen_model": {
            "fit_source": str(fit_path),
            "fit_source_sha256": _sha256(fit_path),
            "post_access_parameter_tuning": False,
            "boundary_conditions": "measured per-tank inlet pressure, inlet temperature and mass derivative",
            "ambient_temperature_c": 15.0,
        },
        "screening_limits": SCREENING_LIMITS,
        "aggregate": aggregate,
        "geometry_diagnostic": geometry_diagnostic,
        "tanks": rows,
        "source_commit": _git_commit(),
        "source_worktree_dirty": _git_dirty(),
        "reproducibility": {
            "command": "python scripts/run_nrel_h2fills_hdvs_validation.py",
            "raw_workbook_is_not_tracked": True,
            "interpretation": (
                "A failed screen is retained as a model-boundary diagnostic; it is not "
                "used to fit new parameters or to claim validation success."
            ),
        },
    }
    output_path.parent.mkdir(parents=True, exist_ok=True)
    output_path.write_text(
        json.dumps(report, ensure_ascii=False, indent=2) + "\n", encoding="utf-8"
    )
    report_path = output_path.with_suffix(".md")
    lines = [
        "# NREL H2FillS HDVS Type-IV tank screening",
        "",
        "This is an independent tank-submodel screen, not a full HRS closed-loop validation.",
        "",
        f"- Tanks: {aggregate['tank_count']} (7-tank, 68.6 kg HDVS)",
        f"- Samples: {aggregate['sample_count_per_tank']} per tank at 1 s cadence",
        f"- Pressure RMSE: {aggregate['pressure_rmse_mpa']:.3f} MPa",
        f"- Temperature RMSE: {aggregate['temperature_rmse_c']:.3f} °C",
        f"- Mean final mass error: {aggregate['mass_final_error_kg']:.3f} kg",
        f"- Joint tank screen: {aggregate['screening_pass_count']}/{aggregate['tank_count']}",
        f"- EOS-implied volume median: {geometry_diagnostic['implied_volume_m3_median_across_tanks']:.5f} m³",
        f"- Ratio to frozen effective volume: {geometry_diagnostic['ratio_to_frozen_effective_volume_median']:.3f}",
        "",
        "The EOS-implied volume comparison is a diagnostic only. It does not fit or replace the frozen geometry and cannot be used as a validation pass.",
        "",
        "The raw workbook remains in the ignored data directory because the H2FillS package licence does not grant redistribution rights.",
    ]
    report_path.write_text("\n".join(lines) + "\n", encoding="utf-8")
    return report


def main() -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument("--workbook", type=Path, default=DEFAULT_WORKBOOK)
    parser.add_argument("--fit", type=Path, default=ROOT / "research/tank_model_validation_v2.json")
    parser.add_argument("--output", type=Path, default=DEFAULT_OUTPUT)
    args = parser.parse_args()
    report = run(args.workbook, args.output, args.fit)
    print(json.dumps(report["aggregate"], ensure_ascii=False, indent=2))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
