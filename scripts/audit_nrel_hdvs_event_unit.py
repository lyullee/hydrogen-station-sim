"""Audit the independent-event unit of the local NREL HDVS workbook.

The workbook has seven receiving tanks on one common clock.  The tanks are
parallel channels in one fill and must not be counted as seven independent
charging events.  This audit retains only structural metadata and never copies
the internal-use raw workbook into the repository.
"""

from __future__ import annotations

from datetime import datetime, timezone
import hashlib
import json
from pathlib import Path
import sys

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
RAW = ROOT / "data/public_validation/raw/nrel_h2fills_2022_hdvs_typeiv.xlsx"
SOURCE_RECORD = ROOT / "research/nrel_hdvs_raw_trace_boundary_2026_10_05.json"
OUTPUT = ROOT / "research/nrel_hdvs_event_unit_audit_2026_10_11.json"


def _sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def run() -> dict:
    source = json.loads(SOURCE_RECORD.read_text(encoding="utf-8"))
    expected_hash = (source.get("source") or {}).get("workbook_sha256")
    if not RAW.is_file():
        raise FileNotFoundError(RAW)
    actual_hash = _sha256(RAW)
    if expected_hash and actual_hash != expected_hash:
        raise ValueError("NREL workbook hash does not match the retained boundary record")

    workbook = load_workbook(RAW, read_only=True, data_only=True)
    if "Data" not in workbook.sheetnames:
        raise ValueError("NREL workbook is missing the Data sheet")
    rows = workbook["Data"].iter_rows(values_only=True)
    header = tuple(next(rows))
    names = {str(value) for value in header if value is not None}
    timed_rows = [row for row in rows if row and row[0] is not None]
    tank_ids = sorted({
        int(name.split("#", 1)[1].split("_", 1)[0])
        for name in names
        if name.startswith("HDVS_ tank#") and "_mass [kg]" in name
    })
    time_axes = [float(row[0]) for row in timed_rows]
    if len(time_axes) < 2 or any(b <= a for a, b in zip(time_axes, time_axes[1:])):
        raise ValueError("NREL time axis is not strictly increasing")

    report = {
        "schema_version": 1,
        "generated_at": datetime.now(timezone.utc).isoformat(),
        "artifact_type": "nrel_hdvs_independent_event_unit_audit",
        "evidence_role": "privacy_bounded_structural_audit",
        "source": {
            "boundary_record": str(SOURCE_RECORD.relative_to(ROOT)).replace("\\", "/"),
            "workbook_sha256": actual_hash,
            "raw_workbook_retained_locally_only": True,
        },
        "experiment": {
            "independent_event_count": 1,
            "analysis_unit": "one common-clock multi-tank fill",
            "shared_clock_group_count": 1,
            "parallel_tank_channel_count": len(tank_ids),
            "tank_ids": tank_ids,
            "timed_row_count": len(timed_rows),
            "time_range_s": [time_axes[0], time_axes[-1]],
            "independence_rule": (
                "Parallel tanks sharing one time axis and one fill boundary are not "
                "independent events; they remain correlated channels of one event."
            ),
        },
        "claim_boundary": (
            "This audit prevents seven parallel tank channels from being reported as "
            "seven independent charging events. It does not validate the station "
            "controller, safety response, full-loop accuracy, or runtime parameters."
        ),
    }
    OUTPUT.write_text(json.dumps(report, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    return report


if __name__ == "__main__":
    print(json.dumps(run(), ensure_ascii=False, indent=2))

